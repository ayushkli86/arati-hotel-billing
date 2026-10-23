"""Read an uploaded Excel workbook and build a self-contained analysis."""

from __future__ import annotations

import io
import re
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import PurePath

from openpyxl import load_workbook

from .services import local_now

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 15000
MAX_COLUMNS = 80
PREVIEW_ROWS = 20
PREVIEW_COLUMNS = 12

_MONEY = re.compile(
    r"(amount|total|price|rate|revenue|sales?|paid|balance|discount|vat|tax|"
    r"charge|cost|fee|salary|income|expense|payment|npr|rs\.?|₹)",
    re.IGNORECASE,
)
_IDENTIFIER = re.compile(
    r"^(s\.?\s*n\.?|id|serial|no\.?|number|code|phone|mobile|contact|qty|"
    r"quantity|count|year|month|day|nights?|pax|units?)$",
    re.IGNORECASE,
)
_DATE_HEADER = re.compile(r"(date|day|month|period|time)", re.IGNORECASE)
_DATE_TEXT = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
_CURRENCY_PREFIX = re.compile(r"^(npr|rs\.?|₹)\s*", re.IGNORECASE)
_DATE_FORMATS = (
    "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d",
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
    "%m/%d/%Y", "%m-%d-%Y",
    "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S",
    "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S",
)


def _to_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    if isinstance(value, str):
        text = _CURRENCY_PREFIX.sub("", value.strip()).replace(",", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _to_datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, str):
        text = value.strip()
        if not _DATE_TEXT.match(text):
            return None
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
    return None


def _cell(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip()
    return text or None


def _header(rows: list[tuple]) -> int | None:
    for index, values in enumerate(rows[:10]):
        non_empty = [value for value in values if value is not None]
        if len(non_empty) < 2:
            continue
        textual = [
            value for value in non_empty
            if isinstance(value, str)
            and _to_number(value) is None
            and _to_datetime(value) is None
        ]
        if len(textual) >= max(1, len(non_empty) // 2):
            return index
    return None


def _normalise_headers(row: tuple, width: int) -> list[str]:
    headers: list[str] = []
    counts: Counter[str] = Counter()
    for index in range(width):
        raw = row[index] if index < len(row) else None
        name = re.sub(r"\s+", " ", str(raw).strip())[:120] if raw is not None else ""
        if not name:
            name = f"Column {index + 1}"
        counts[name] += 1
        if counts[name] > 1:
            name = f"{name} ({counts[name]})"
        headers.append(name)
    return headers


def _classify(values: list[object]) -> str:
    present = [value for value in values if value is not None]
    if not present:
        return "empty"
    numbers = sum(1 for value in present if _to_number(value) is not None)
    if numbers / len(present) >= 0.8:
        return "number"
    dates = sum(1 for value in present if _to_datetime(value) is not None)
    if dates / len(present) >= 0.8:
        return "date"
    return "text"


def _format_number(value: float, money: bool = False) -> str:
    if abs(value - round(value)) < 1e-9:
        return f"{int(round(value)):,}"
    return f"{value:,.2f}" if money else f"{value:,.2f}".rstrip("0").rstrip(".")


def _format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.strftime("%H:%M") if value.time() != time.min else value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return _format_number(value)
    return str(value)


def _pick_primary_amount(headers: list[str], numeric: list[int]) -> int | None:
    money_columns = [index for index in numeric if _MONEY.search(headers[index])]
    eligible = [
        index for index in (money_columns or numeric)
        if not _IDENTIFIER.match(headers[index].strip())
    ]
    if eligible:
        return eligible[0]
    return numeric[0] if numeric else None


def _load_rows(workbook) -> tuple[str, list[tuple]]:
    for sheet in workbook.worksheets:
        rows = []
        truncated = False
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_DATA_ROWS + 20:
                truncated = True
                break
            rows.append(tuple(row[:MAX_COLUMNS]))
        if any(any(value is not None for value in row) for row in rows):
            if truncated:
                rows.append(("__truncated__",))
            return sheet.title, rows
    raise ValueError("The workbook has no data. Add a sheet with rows and try again.")


def analyze(content: bytes, filename: str) -> dict:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("That file is larger than 10 MB. Upload a smaller workbook.")
    suffix = PurePath(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Choose an Excel workbook (.xlsx or .xlsm).")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("That file could not be read as an Excel workbook.") from exc

    try:
        sheet_name, raw_rows = _load_rows(workbook)
        sheets = workbook.sheetnames
    finally:
        workbook.close()

    truncated = bool(raw_rows) and raw_rows[-1] == ("__truncated__",)
    if truncated:
        raw_rows = raw_rows[:-1]

    header_index = _header(raw_rows)
    header_row = raw_rows[header_index] if header_index is not None else ()
    data_rows = raw_rows[header_index + 1:] if header_index is not None else raw_rows
    width = min(
        MAX_COLUMNS,
        max((len(row) for row in raw_rows), default=0),
    )
    headers = _normalise_headers(header_row, width)

    empty_rows = 0
    records: list[list[object]] = []
    for row in data_rows:
        values = [_cell(row[index]) if index < len(row) else None for index in range(width)]
        if not any(value is not None for value in values):
            empty_rows += 1
            continue
        records.append(values)

    columns = [[record[index] for record in records] for index in range(width)]
    kinds = [_classify(values) for values in columns]
    numeric = [index for index, kind in enumerate(kinds) if kind == "number"]
    dates = [index for index, kind in enumerate(kinds) if kind == "date"]
    texts = [index for index, kind in enumerate(kinds) if kind == "text"]
    primary = _pick_primary_amount(headers, numeric)

    numeric_columns = []
    for index in numeric:
        values = [number for value in columns[index] if (number := _to_number(value)) is not None]
        if not values:
            continue
        money = bool(_MONEY.search(headers[index]))
        numeric_columns.append({
            "name": headers[index],
            "money": money,
            "count": len(values),
            "sum": _format_number(sum(values), money),
            "mean": _format_number(sum(values) / len(values), money),
            "min": _format_number(min(values), money),
            "max": _format_number(max(values), money),
            "negative": sum(1 for value in values if value < 0),
            "non_numeric": sum(
                1 for value in columns[index]
                if value is not None and _to_number(value) is None
            ),
        })

    primary_total = None
    if primary is not None:
        values = [number for value in columns[primary] if (number := _to_number(value)) is not None]
        if values:
            primary_total = {
                "name": headers[primary],
                "total": _format_number(sum(values), bool(_MONEY.search(headers[primary]))),
            }

    date_total = None
    date_values = [
        parsed for index in dates for value in columns[index]
        if (parsed := _to_datetime(value)) is not None
    ]
    if date_values:
        first, last = min(date_values), max(date_values)
        date_total = (
            first.strftime("%Y-%m-%d")
            if first.date() == last.date()
            else f"{first:%Y-%m-%d} → {last:%Y-%m-%d}"
        )

    breakdowns = []
    amount_index = primary
    for index in texts:
        labels = [str(value).strip() for value in columns[index] if value is not None]
        if not labels:
            continue
        counts = Counter(labels)
        if len(counts) > 40 or len(counts) == len(labels):
            continue
        totals: dict[str, float] = {}
        if amount_index is not None:
            for row in records:
                label = row[index]
                amount = _to_number(row[amount_index])
                if label is not None and amount is not None:
                    key = str(label).strip()
                    totals[key] = totals.get(key, 0.0) + amount
        money = bool(amount_index is not None and _MONEY.search(headers[amount_index]))
        grand = sum(totals.values())
        rows = []
        for label, count in sorted(counts.items(), key=lambda item: (-totals.get(item[0], 0.0), -item[1], item[0])):
            total = totals.get(label, 0.0)
            rows.append({
                "label": label[:80],
                "count": count,
                "total": _format_number(total, money) if amount_index is not None else "",
                "share": round(total / grand * 100, 1) if grand else None,
            })
            if len(rows) >= 10:
                break
        if rows:
            breakdowns.append({
                "name": headers[index],
                "money": money,
                "has_total": amount_index is not None,
                "rows": rows,
            })
        if len(breakdowns) >= 3:
            break

    trend = None
    if amount_index is not None and dates:
        date_index = next(
            (index for index in dates if _DATE_HEADER.search(headers[index])),
            dates[0],
        )
        money = bool(_MONEY.search(headers[amount_index]))
        buckets: dict[str, float] = {}
        for row in records:
            parsed = _to_datetime(row[date_index])
            amount = _to_number(row[amount_index])
            if parsed is None or amount is None:
                continue
            key = parsed.strftime("%Y-%m-%d")
            buckets[key] = buckets.get(key, 0.0) + amount
        if len(buckets) > 60:
            monthly: dict[str, float] = {}
            for key, value in buckets.items():
                monthly[key[:7]] = monthly.get(key[:7], 0.0) + value
            buckets = monthly
        points = [
            {"label": key, "value": value}
            for key, value in sorted(buckets.items())
        ][-60:]
        if points:
            peak = max(point["value"] for point in points)
            for point in points:
                point["display"] = _format_number(point["value"], money)
                point["width"] = max(2, round(point["value"] / peak * 100)) if peak > 0 else 2
            trend = {
                "name": f"{headers[amount_index]} by {headers[date_index]}",
                "money": money,
                "points": points,
            }

    preview_columns = min(len(headers), PREVIEW_COLUMNS)
    preview = {
        "headers": headers[:preview_columns],
        "rows": [
            [_format_cell(value) for value in record[:preview_columns]]
            for record in records[:PREVIEW_ROWS]
        ],
        "columns_shown": preview_columns,
        "columns_total": len(headers),
        "rows_shown": min(len(records), PREVIEW_ROWS),
        "rows_total": len(records),
    }

    findings = []
    if truncated:
        findings.append({
            "level": "high",
            "text": f"Only the first {MAX_DATA_ROWS:,} rows were analysed. Split the file for a complete review.",
        })
    missing = [
        (headers[index], sum(1 for value in columns[index] if value is None))
        for index in range(width)
    ]
    missing = sorted((item for item in missing if item[1]), key=lambda item: -item[1])
    for name, count in missing[:5]:
        findings.append({"level": "review", "text": f"“{name}” has {count:,} blank cell(s)."})
    if empty_rows:
        findings.append({"level": "info", "text": f"{empty_rows:,} empty row(s) were skipped."})
    duplicate_count = len(records) - len({tuple(record) for record in records})
    if duplicate_count:
        findings.append({
            "level": "medium",
            "text": f"{duplicate_count:,} duplicate row(s) were found. Check for repeated entries before relying on totals.",
        })
    for column in numeric_columns:
        if column["negative"]:
            findings.append({
                "level": "medium",
                "text": f"“{column['name']}” contains {column['negative']:,} negative value(s).",
            })
        if column["non_numeric"]:
            findings.append({
                "level": "review",
                "text": f"“{column['name']}” has {column['non_numeric']:,} value(s) that are not numbers and were excluded from totals.",
            })
    if not numeric_columns:
        findings.append({
            "level": "info",
            "text": "No numeric columns were detected, so totals and averages are unavailable.",
        })
    if not findings:
        findings.append({"level": "info", "text": "No structural issues were detected in the imported data."})
    order = {"high": 0, "medium": 1, "review": 2, "info": 3}
    findings.sort(key=lambda item: order[item["level"]])

    return {
        "filename": PurePath(filename).name,
        "sheet": sheet_name,
        "sheets": sheets,
        "generated_at": local_now().strftime("%Y-%m-%d %H:%M"),
        "row_count": len(records),
        "column_count": len(headers),
        "numeric_count": len(numeric_columns),
        "date_total": date_total,
        "primary_total": primary_total,
        "truncated": truncated,
        "numeric_columns": numeric_columns,
        "breakdowns": breakdowns,
        "trend": trend,
        "preview": preview,
        "findings": findings,
    }

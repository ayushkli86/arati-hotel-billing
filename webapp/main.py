"""FastAPI entry point for the Arati Hotel front-office dashboard."""

from __future__ import annotations

import logging
import os
import re
import json
import asyncio
import secrets
import csv
import io
import threading
import time as _time
from contextlib import asynccontextmanager
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from types import SimpleNamespace

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from markupsafe import escape as html_escape
from fastapi.responses import (
    FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, Response, StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
import httpx
from sqlalchemy import func, inspect, or_, select, text, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload
from starlette.exceptions import HTTPException as StarletteHTTPException
from starlette.middleware.sessions import SessionMiddleware

from .defaults import DEFAULT_SETTINGS
from .money import (
    BILL_TYPE_LABELS, Line, compute_totals, discount_from_input, payment_status,
)
from .nepali import amount_in_words as amount_in_words_ne, label as source_invoice_label, to_bs

from . import analytics as web_analytics, auth, daily_closing, excel_analysis, models, reporting, whatsapp
from .database import Base, SessionLocal, engine, get_db
from .seed_menu import MENU as MENU_ITEMS
from .models import (
    AuditLog, CatalogItem, DailyReport, Guest, HotelSetting, Invoice, InvoiceItem,
    Kot, Order, Payment, Refund, Reservation, Room, RoomCharge, Stay, User, UserSession,
    WhatsAppConversation, WhatsAppMessage, StaffAttendance,
)
from .services import guest_service, invoice_service, payment_service, room_service, stay_service
from zoneinfo import ZoneInfo
BUSINESS_TZ_LOCAL = ZoneInfo("Asia/Kathmandu")

_log = logging.getLogger("webapp")
from .orders import (
    FULFILMENT_LABELS, KIND_LABELS, OPEN_STATUSES, ORDER_STATUSES, STATUS_BADGES,
    STATUS_LABELS, can_transition, item_summary, next_status,
)
from .services import BUSINESS_TZ, amount_in_words, collect_payment, local_now, money
from .ui import toast_event, toast_from_request, toast_url


ROOT = Path(__file__).resolve().parent
templates = Jinja2Templates(directory=str(ROOT / "templates"))

# The development database can contain rows created by the automated test
# suites (generated room numbers, "Analytics test guest", etc.). They are real
# database rows, but they must never leak into the front-office UI. These
# helpers keep operational views clean without deleting any accounting data.
_TEST_ROOM_NUMBER = re.compile(r"^[A-Za-z][0-9a-fA-F]{6}$")
_TEST_ROOM_TYPE = re.compile(r"\b(test|updated)\b", re.IGNORECASE)
_TEST_GUEST_NAME = re.compile(
    r"(\btest\b|analytics|print test|draft lifecycle|booked guest|cancellation guest|"
    r"ledger guest|draft owner|new unique guest|protected guest|updated guest|"
    r"other guest|reserved guest|payment methods|receipt guest|room ops|void me|"
    r"next guest)",
    re.IGNORECASE,
)
_TEST_HEX_SUFFIX = re.compile(r"\s[0-9a-f]{4,6}$", re.IGNORECASE)
_TEST_INVOICE_PREFIX = re.compile(
    r"^(TEST|RCPT|LEDGE|REPOR|ANALY|OTHER|METHO|GUEST|DRAFT)", re.IGNORECASE
)


def _is_demo_room(room: Room | None) -> bool:
    if room is None:
        return False
    number = (room.number or "").strip()
    kind = (room.type or "").strip()
    return bool(_TEST_ROOM_NUMBER.match(number) or _TEST_ROOM_TYPE.search(kind))


def _is_demo_guest(guest: Guest | None) -> bool:
    if guest is None:
        return False
    name = (guest.name or "").strip()
    if not name:
        return False
    if _TEST_GUEST_NAME.search(name) or _TEST_HEX_SUFFIX.search(name):
        return True
    return False


def _is_demo_invoice(invoice: Invoice | None) -> bool:
    if invoice is None:
        return False
    if _TEST_INVOICE_PREFIX.match((invoice.invoice_number or "").strip()):
        return True
    return _is_demo_guest(invoice.guest) or _is_demo_room(invoice.room)


def _receipt_money(value: object) -> str:
    """The Tkinter receipt prints bare amounts; the currency symbol goes on TOTAL."""
    try:
        return f"{float(value or 0):,.2f}"
    except (TypeError, ValueError):
        return "0.00"


def _document_directory(kind: str) -> Path:
    base = Path(os.getenv("ARATI_WEB_DOCUMENT_DIRECTORY", str(ROOT / "data"))).resolve()
    directory = base / kind
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _save_pdf_copy(pdf: bytes, kind: str, filename: str) -> Path:
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "-", filename).strip("-._") or "document.pdf"
    target = _document_directory(kind) / safe_name
    temporary = target.with_name(f".{target.name}.{secrets.token_hex(4)}.tmp")
    temporary.write_bytes(pdf)
    temporary.replace(target)
    return target


templates.env.globals["receipt_money"] = _receipt_money
templates.env.globals["money"] = money
templates.env.globals["amount_in_words"] = amount_in_words


def billing_label(key: str, language: str) -> str:
    language_code = "ne" if language.lower() in {"ne", "nepali"} else "en"
    return source_invoice_label(key, language_code)


templates.env.globals["billing_label"] = billing_label
templates.env.globals["amount_in_words_ne"] = amount_in_words_ne


def vat_taxable_base(vat_amount, vat_rate) -> Decimal:
    """Reverse the VAT amount back to its taxable base, rounded to paisa."""
    try:
        rate = Decimal(str(vat_rate or 0))
        vat = Decimal(str(vat_amount or 0))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.00")
    if rate <= 0:
        return Decimal("0.00")
    return (vat * Decimal("100") / rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def pad_money(value) -> str:
    """Summary totals with thousands separators and always two paisa digits."""
    try:
        amount = Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        amount = Decimal("0.00")
    return f"{amount:,.2f}"


def rs_ps(value) -> tuple[str, str]:
    """Split an amount into the (rupees, paisa) columns of the print layout."""
    try:
        paisa_total = int(
            (Decimal(str(value or 0)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        )
    except (InvalidOperation, ValueError, TypeError):
        paisa_total = 0
    sign = "-" if paisa_total < 0 else ""
    rupees, paisa = divmod(abs(paisa_total), 100)
    return f"{sign}{rupees:,}", f"{paisa:02d}"


def bs_date_slashes(value) -> str:
    """Bikram Sambat date as YYYY/MM/DD, or blank when unavailable."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        iso = value.isoformat()
    else:
        iso = str(value)[:10]
    converted = to_bs(iso)
    if not converted:
        return ""
    return converted[0].replace("-", "/")


templates.env.globals["vat_taxable_base"] = vat_taxable_base
templates.env.globals["pad_money"] = pad_money
templates.env.globals["rs_ps"] = rs_ps
templates.env.globals["bs_date_slashes"] = bs_date_slashes


def local_datetime(value: datetime | None, fmt: str = "%d %b · %H:%M") -> str:
    if value is None:
        return "—"
    value = value.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ)
    return value.strftime(fmt)


templates.env.globals["local_datetime"] = local_datetime
templates.env.globals["local_now"] = local_now

PAYMENT_METHODS = ("Cash", "Card", "eSewa", "Khalti", "Fonepay", "Bank Transfer", "Room Credit")
REFUND_REASONS = (
    "Service not delivered",
    "Guest overcharged",
    "Booking cancelled",
    "Courtesy / goodwill",
    "Other (see note)",
)
ADJUSTMENT_REASONS = (
    "Correct earlier entry",
    "Goodwill credit",
    "Service charge waiver",
    "Other (see note)",
)


def _is_production() -> bool:
    return os.getenv("ARATI_ENV", "development").strip().lower() == "production"


def _session_secret() -> str:
    """Never fall back to a shared, committed secret.

    A public default key lets anyone forge a session cookie holding
    ``{"user_id": 1}`` and become the administrator. In production a real
    secret is mandatory; in development a per-process random secret is used,
    which only means sessions do not survive a restart.
    """
    configured = os.getenv("ARATI_WEB_SECRET", "").strip()
    if configured and configured != "change-this-development-secret-before-deploying":
        return configured
    if _is_production():
        raise RuntimeError(
            "Set a strong ARATI_WEB_SECRET (32+ characters) before production startup."
        )
    return secrets.token_hex(32)


@asynccontextmanager
async def lifespan(_app: FastAPI):
    if _is_production() and not os.getenv("ARATI_WEB_SECRET", "").strip():
        raise RuntimeError("Set a strong ARATI_WEB_SECRET before production startup.")
    upgrade_database()
    scheduler = None
    if os.getenv("ARATI_WEB_DAILY_CLOSE_ENABLED", "1").strip().lower() not in {
        "0", "false", "no", "off",
    }:
        scheduler = daily_closing.DailyCloseScheduler(SessionLocal)
        scheduler.start()
    try:
        yield
    finally:
        if scheduler is not None:
            scheduler.stop()


def upgrade_database(bind=engine) -> None:
    """Create Web tables and add Tinker-compatible fields without dropping data."""
    Base.metadata.create_all(bind=bind)
    additions = {
        "users": {
            "must_change_password": "BOOLEAN NOT NULL DEFAULT 0",
            "session_version": "INTEGER NOT NULL DEFAULT 0",
            "created_at": "DATETIME",
        },
        "guests": {
            "address": "VARCHAR(240) NOT NULL DEFAULT ''",
            "pan_vat_number": "VARCHAR(80) NOT NULL DEFAULT ''",
            "customer_type": "VARCHAR(30) NOT NULL DEFAULT 'individual'",
            "active": "BOOLEAN NOT NULL DEFAULT 1",
            "created_at": "DATETIME",
        },
        "rooms": {
            "active": "BOOLEAN NOT NULL DEFAULT 1",
            "room_type_item_id": "INTEGER REFERENCES catalog_items(id)",
            "nightly_rate_override": "BOOLEAN NOT NULL DEFAULT 1",
        },
        "reservations": {
            "room_number": "VARCHAR(20) NOT NULL DEFAULT ''",
            "guest_name": "VARCHAR(160) NOT NULL DEFAULT ''",
            "guest_phone": "VARCHAR(40) NOT NULL DEFAULT ''",
            "guest_email": "VARCHAR(180) NOT NULL DEFAULT ''",
            "guest_address": "VARCHAR(240) NOT NULL DEFAULT ''",
            "guest_pan": "VARCHAR(80) NOT NULL DEFAULT ''",
            "nightly_rate": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "deposit": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "source": "VARCHAR(80) NOT NULL DEFAULT 'front desk'",
            "notes": "TEXT NOT NULL DEFAULT ''",
            "cancel_reason": "TEXT NOT NULL DEFAULT ''",
            "business_date": "VARCHAR(20) NOT NULL DEFAULT ''",
            "created_at": "DATETIME",
            "user_id": "INTEGER REFERENCES users(id)",
        },
        "stays": {
            "reservation_id": "INTEGER REFERENCES reservations(id)",
            "room_number": "VARCHAR(20) NOT NULL DEFAULT ''",
            "invoice_number": "VARCHAR(40) NOT NULL DEFAULT ''",
            "guest_name": "VARCHAR(160) NOT NULL DEFAULT ''",
            "guest_phone": "VARCHAR(40) NOT NULL DEFAULT ''",
            "guest_email": "VARCHAR(180) NOT NULL DEFAULT ''",
            "guest_address": "VARCHAR(240) NOT NULL DEFAULT ''",
            "guest_pan": "VARCHAR(80) NOT NULL DEFAULT ''",
            "expected_checkout": "DATETIME",
            "nights": "INTEGER NOT NULL DEFAULT 1",
            "nightly_rate": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "business_date": "VARCHAR(20) NOT NULL DEFAULT ''",
            "created_at": "DATETIME",
            "user_id": "INTEGER REFERENCES users(id)",
        },
        "invoices": {
            "invoice_type": "VARCHAR(30) NOT NULL DEFAULT 'Tax Invoice'",
            "language": "VARCHAR(20) NOT NULL DEFAULT 'English'",
            "copy_type": "VARCHAR(30) NOT NULL DEFAULT 'Customer Copy'",
            "payment_terms": "VARCHAR(60) NOT NULL DEFAULT 'Due on Check-out'",
            "notes": "TEXT NOT NULL DEFAULT ''", "vat_amount": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "subtotal": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "check_in": "DATETIME", "check_out": "DATETIME", "nights": "INTEGER NOT NULL DEFAULT 1",
            "created_at": "DATETIME",
            "bill_date": "VARCHAR(20) NOT NULL DEFAULT ''",
            "guest_name": "VARCHAR(160) NOT NULL DEFAULT ''",
            "guest_phone": "VARCHAR(40) NOT NULL DEFAULT ''",
            "guest_address": "VARCHAR(240) NOT NULL DEFAULT ''",
            "guest_pan": "VARCHAR(80) NOT NULL DEFAULT ''",
            "room_number": "VARCHAR(20) NOT NULL DEFAULT ''",
            "item_discount": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "net_subtotal": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "taxable_net": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "service_charge_rate": "NUMERIC(7, 3) NOT NULL DEFAULT 0",
            "service_charge": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "vat_rate": "NUMERIC(7, 3) NOT NULL DEFAULT 0",
            "payment_method": "VARCHAR(40) NOT NULL DEFAULT 'Cash'",
            "prepared_by": "VARCHAR(80) NOT NULL DEFAULT ''",
            "received_by": "VARCHAR(80) NOT NULL DEFAULT ''",
            "customer_type": "VARCHAR(30) NOT NULL DEFAULT 'individual'",
            "is_draft": "BOOLEAN NOT NULL DEFAULT 0",
            "bill_type": "VARCHAR(30) NOT NULL DEFAULT 'tax'",
            "copy_label": "VARCHAR(80) NOT NULL DEFAULT ''",
            "user_id": "INTEGER REFERENCES users(id)",
        },
        "invoice_items": {
            "discount_amount": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "vat_amount": "NUMERIC(12, 2) NOT NULL DEFAULT 0",
            "item_code": "VARCHAR(80) NOT NULL DEFAULT ''",
            "category": "VARCHAR(30) NOT NULL DEFAULT 'food'",
            "unit": "VARCHAR(30) NOT NULL DEFAULT 'pcs'",
            "taxable": "BOOLEAN NOT NULL DEFAULT 1",
            "kot_id": "INTEGER REFERENCES kots(id)",
        },
        "catalog_items": {
            "item_code": "VARCHAR(80) NOT NULL DEFAULT ''",
            "unit": "VARCHAR(30) NOT NULL DEFAULT 'pcs'",
            "taxable": "BOOLEAN NOT NULL DEFAULT 1",
            "sort_order": "INTEGER NOT NULL DEFAULT 0",
            "menu_group": "VARCHAR(80) NOT NULL DEFAULT ''",
        },
        "kots": {
            "room_id": "INTEGER REFERENCES rooms(id)",
            "room_number": "VARCHAR(20) NOT NULL DEFAULT ''",
            "issued_by_name": "VARCHAR(160) NOT NULL DEFAULT ''",
            "cut_at": "DATETIME",
            "cut_by_id": "INTEGER REFERENCES users(id)",
        },
        "room_charges": {
            "kot_id": "INTEGER REFERENCES kots(id)",
        },
        "payments": {
            "invoice_number": "VARCHAR(40) NOT NULL DEFAULT ''",
            "business_date": "VARCHAR(20) NOT NULL DEFAULT ''",
            "created_at": "DATETIME",
            "user_id": "INTEGER REFERENCES users(id)",
        },
        "orders": {
            "business_date": "VARCHAR(20) NOT NULL DEFAULT ''",
        },
        "user_sessions": {
            "session_id": "VARCHAR(128) NOT NULL",
            "user_id": "INTEGER NOT NULL REFERENCES users(id)",
            "ip_address": "VARCHAR(45) NOT NULL DEFAULT ''",
            "user_agent": "VARCHAR(512) NOT NULL DEFAULT ''",
            "created_at": "DATETIME NOT NULL",
            "last_seen_at": "DATETIME NOT NULL",
            "revoked_at": "DATETIME",
            "revoked_by_id": "INTEGER REFERENCES users(id)",
            "revoked_reason": "TEXT NOT NULL DEFAULT ''",
        },
    }
    with bind.begin() as connection:
        for table, columns in additions.items():
            present = {column["name"] for column in inspect(connection).get_columns(table)}
            for name, declaration in columns.items():
                if name not in present:
                    connection.execute(text(f'ALTER TABLE "{table}" ADD COLUMN "{name}" {declaration}'))

    try:
        with Session(bind=bind) as db:
            db.execute(
                text(
                    'CREATE UNIQUE INDEX IF NOT EXISTS ix_user_sessions_sessionid_userid '
                    'ON "user_sessions" ("session_id", "user_id")'
                )
            )
            db.commit()
    except Exception:
        # The index only exists when the schema was created by this app's
        # migrations; older installs may have conflicts — startup must not
        # die for a best-effort index, but the reason still needs a trace.
        _log.warning("user_sessions index creation skipped", exc_info=True)

    # Relaxed constraints repairs pre-1.0 databases; business-day backfill
    # gives every row a local calendar day for the day-wise calendar.
    _relax_stale_not_null(bind)
    _backfill_business_dates(bind)
    # Day-wise lookups (calendar, reports) run on every refresh; keep them
    # indexed even on databases upgraded in place, where create_all never
    # adds the index for ALTER-added columns.
    with bind.begin() as connection:
        for _table, _column in (
            ("reservations", "business_date"), ("stays", "business_date"),
            ("orders", "business_date"), ("invoices", "bill_date"),
            ("kots", "room_id"), ("kots", "room_number"), ("kots", "cut_by_id"),
            ("room_charges", "kot_id"),
        ):
            if _column in {info["name"] for info in inspect(connection).get_columns(_table)}:
                connection.execute(text(
                    f'CREATE INDEX IF NOT EXISTS "ix_{_table}_{_column}" '
                    f'ON "{_table}" ("{_column}")'
                ))


def _relax_stale_not_null(bind=engine) -> None:
    """Rebuild pre-1.0 tables whose columns are still declared NOT NULL.

    Old installs were created when ``stays.check_out`` (checked-in rooms have
    no departure yet), ``reservations.room_id`` / ``reservations.guest_id``
    (a walk-in booking without a room) and ``payments.invoice_id`` (a direct
    ledger payment) were required. The current models make them nullable, so
    every new check-in failed with ``NOT NULL constraint failed: stays.check_out``
    and no room bill could ever be issued. SQLite cannot ALTER a constraint,
    so the affected table is rebuilt in place, copying every row and index.
    """
    targets = {
        "stays": {"check_out"},
        "reservations": {"room_id", "guest_id"},
        "payments": {"invoice_id"},
    }
    for table, columns in targets.items():
        try:
            rows = []
            with bind.begin() as connection:
                rows = connection.execute(
                    text(f'PRAGMA table_info("{table}")')
                ).fetchall()
                names = [row[1] for row in rows]
                if not names or not set(columns) <= set(names):
                    continue
                stale = [row for row in rows if row[1] in columns and row[3]]
                if not stale:
                    continue
                # Column SQL is regenerated from PRAGMA info; unchanged
                # columns keep type, nullability, default and PK flags.
                definitions = []
                for _cid, name, type_, notnull, default, pk in rows:
                    if name in columns:
                        notnull = 0
                    if pk:
                        definitions.append(f'"{name}" {type_} NOT NULL PRIMARY KEY'.strip())
                    else:
                        line = f'"{name}" {type_}'.strip()
                        if notnull:
                            line += " NOT NULL"
                        if default is not None:
                            line += f" DEFAULT {default}"
                        definitions.append(line)
                temp_table = f"{table}_nullable_rebuild"
                connection.execute(text(f'DROP TABLE IF EXISTS "{temp_table}"'))
                connection.execute(
                    text(f'CREATE TABLE "{temp_table}" ({", ".join(definitions)})')
                )
                connection.execute(
                    text(f'INSERT INTO "{temp_table}" SELECT * FROM "{table}"')
                )
            # Secondary indexes are tied to the old table; rebuild them after
            # the swap. The swap needs its own connection so the PRAGMA is
            # not swallowed by an open transaction.
            with bind.begin() as connection:
                indexes = connection.execute(
                    text(
                        "SELECT name, sql FROM sqlite_master "
                        "WHERE type = 'index' AND tbl_name = :table AND sql IS NOT NULL"
                    ),
                    {"table": table},
                ).fetchall()
                connection.execute(text('PRAGMA foreign_keys=OFF'))
                connection.execute(text(f'DROP TABLE "{table}"'))
                connection.execute(text(f'ALTER TABLE "{temp_table}" RENAME TO "{table}"'))
                for _name, sql in indexes:
                    connection.execute(text(sql))
                connection.execute(text('PRAGMA foreign_keys=ON'))
            _log.info("Relaxed schema constraint on %s.%s", table, columns)
        except Exception:
            _log.exception("Could not relax NOT NULL constraint on %s — check-in may keep failing", table)


def _backfill_business_dates(bind=engine) -> None:
    """Give every row a local (Asia/Kathmandu) business day so the day-wise
    calendar can always find its data. Missing values are derived from the
    nearest timestamp column each record carries."""
    with Session(bind=bind) as db:
        changed = False
        for reservation in db.scalars(
            select(Reservation).where(Reservation.business_date == "")
        ).all():
            stamp = reservation.created_at or reservation.check_in
            if stamp is not None:
                reservation.business_date = _business_date(stamp).isoformat()
                changed = True
        for stay in db.scalars(select(Stay).where(Stay.business_date == "")).all():
            stamp = stay.check_in
            if stamp is not None:
                stay.business_date = _business_date(stamp).isoformat()
                changed = True
        for order in db.scalars(select(Order).where(Order.business_date == "")).all():
            stamp = order.created_at or order.closed_at
            if stamp is not None:
                order.business_date = _business_date(stamp).isoformat()
                changed = True
        for invoice in db.scalars(select(Invoice).where(Invoice.bill_date == "")).all():
            stamp = invoice.issue_date
            if stamp is not None:
                invoice.bill_date = _business_date(stamp).isoformat()
                changed = True
        for payment in db.scalars(select(Payment).where(Payment.business_date == "")).all():
            stamp = payment.created_at or payment.paid_at
            if stamp is not None:
                payment.business_date = _business_date(stamp).isoformat()
                changed = True
        for refund in db.scalars(select(Refund).where(Refund.business_date == "")).all():
            stamp = refund.created_at or refund.refunded_at
            if stamp is not None:
                refund.business_date = _business_date(stamp).isoformat()
                changed = True
        for charge in db.scalars(select(RoomCharge).where(RoomCharge.charge_date == "")).all():
            stamp = charge.created_at
            if stamp is not None:
                charge.charge_date = _business_date(stamp).isoformat()
                changed = True
        if changed:
            db.commit()

    with Session(bind=bind) as db:
        for room in db.scalars(select(Room).where(Room.room_type_item_id.is_(None))).all():
            room_type = room.type.strip() or "Unspecified"
            item = db.scalar(
                select(CatalogItem)
                .where(
                    CatalogItem.category == "room",
                    func.lower(CatalogItem.name) == room_type.lower(),
                )
                .order_by(CatalogItem.active.desc(), CatalogItem.id)
            )
            if item is None:
                item = CatalogItem(
                    name=room_type,
                    category="room",
                    rate=room.nightly_rate,
                    unit="night",
                    taxable=True,
                )
                db.add(item)
                db.flush()
            room.type = item.name
            room.room_type_item = item
            room.nightly_rate_override = room.nightly_rate != item.rate
        db.commit()


_docs_enabled = not _is_production() or os.getenv("ARATI_WEB_DOCS", "0").strip().lower() in {"1", "true", "yes", "on"}
app = FastAPI(
    title="Arati Hotel · Khandbari",
    version="1.0.0",
    lifespan=lifespan,
    docs_url="/docs" if _docs_enabled else None,
    redoc_url="/redoc" if _docs_enabled else None,
    openapi_url="/openapi.json" if _docs_enabled else None,
)
app.add_middleware(
    SessionMiddleware,
    secret_key=_session_secret(),
    session_cookie="arati_session",
    max_age=60 * 60 * 10,
    same_site="lax",
    https_only=os.getenv("ARATI_WEB_HTTPS", "1" if _is_production() else "0") == "1",
)
@app.middleware("http")
async def security_headers(request: Request, call_next):
    """Browser-enforced hardening for every response.

    CSP keeps 'unsafe-inline'/'unsafe-eval' because the UI ships inline
    handlers and the Alpine/Tailwind CDN builds compile expressions at
    runtime; self-hosting those assets is the follow-up that lets us
    tighten it. Iframes are refused outright, which removes clickjacking
    as a way around CSRF.
    """
    response = await call_next(request)
    headers = response.headers
    headers.setdefault("X-Content-Type-Options", "nosniff")
    headers.setdefault("X-Frame-Options", "DENY")
    headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    headers.setdefault("Cross-Origin-Opener-Policy", "same-origin")
    headers.setdefault(
        "Content-Security-Policy",
        "default-src 'self'; "
        "script-src 'self' 'unsafe-inline' 'unsafe-eval' https://cdn.tailwindcss.com https://unpkg.com; "
        "style-src 'self' 'unsafe-inline' https://cdn.tailwindcss.com https://fonts.googleapis.com; "
        "img-src 'self' data: blob:; "
        "font-src 'self' data: https://fonts.gstatic.com; "
        "connect-src 'self'; "
        "object-src 'none'; "
        "base-uri 'self'; "
        "form-action 'self'; "
        "frame-ancestors 'none'",
    )
    if request.url.scheme == "https" or os.getenv("ARATI_WEB_HTTPS") == "1":
        headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    if "text/html" in headers.get("content-type", ""):
        # Guest data must not be restorable from the browser cache after logout.
        headers.setdefault("Cache-Control", "no-store")
    return response


app.mount("/static", StaticFiles(directory=str(ROOT / "static")), name="static")


@app.exception_handler(HTTPException)
@app.exception_handler(StarletteHTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    is_htmx = request.headers.get("HX-Request") == "true"
    accepts_html = "text/html" in request.headers.get("Accept", "")

    if exc.status_code == 303 and exc.headers and "Location" in exc.headers:
        loc = exc.headers["Location"]
        if is_htmx:
            return Response(status_code=200, headers={"HX-Redirect": loc})
        return RedirectResponse(url=loc, status_code=303)

    detail = exc.detail if isinstance(exc.detail, str) else str(exc.detail)
    safe_detail = html_escape(detail)

    if is_htmx:
        html = f"""<div role="alert" class="rounded-xl border border-red-200 bg-red-50 p-4 text-sm font-semibold text-red-800 shadow-sm mb-3">
            <div class="flex items-center gap-2"><i class="ph ph-warning-circle"></i><span>{safe_detail}</span></div>
        </div>"""
        return HTMLResponse(content=html, status_code=exc.status_code)

    if accepts_html:
        html = _error_page(exc.status_code, safe_detail)
        return HTMLResponse(content=html, status_code=exc.status_code)

    from fastapi.responses import JSONResponse
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})


def _error_page(status_code: int, safe_detail: str) -> str:
    """Solo-brand friendly error card, shared by every non-HTMX error path."""
    glyph = {404: "map-triangle-x", 403: "shield-warning", 422: "list-magnifying-glass"}.get(
        status_code, "lifebuoy",
    )
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Error {status_code}</title>
<script src="https://cdn.tailwindcss.com"></script><link rel="stylesheet" href="/static/phosphor/style.css">
<style>body{{font-family:Inter,'Noto Sans Devanagari',ui-sans-serif,system-ui,sans-serif}}</style>
</head><body class="min-h-screen grid place-items-center p-6" style="background:linear-gradient(160deg,#F7F1DE 0%,#EFE7CF 100%);">
<div class="max-w-md w-full rounded-2xl p-8 text-center" style="background:var(--color-surface);border:1px solid rgba(212,201,168,.6);box-shadow:0 12px 32px rgba(27,58,42,.10);">
<div class="mx-auto mb-4 grid h-14 w-14 place-items-center rounded-2xl" style="background:rgba(125,155,122,.18);color:#1B3A2A;">
<i class="ph {glyph}" style="font-size:1.75rem;"></i></div>
<h1 class="text-3xl font-extrabold tracking-tight" style="color:var(--color-primary);">{status_code}</h1>
<p class="mt-2 text-sm" style="color:var(--color-text-secondary);">{safe_detail}</p>
<a href="/dashboard" class="mt-6 inline-flex items-center gap-2 rounded-xl px-4 py-2 text-sm font-bold no-underline" style="background:var(--color-primary);color:#fff;">Continue<i class="ph ph-arrow-right"></i></a>
</div></body></html>"""


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    """Path/query/type conversion failures (``/invoices/abc/pdf``) answered
    politely for HTML, HTMX, and API consumers alike instead of a raw 422."""
    _log.warning("Validation rejected: %s %s — %s", request.method, request.url.path, exc.errors()[:3])
    return await http_exception_handler(
        request,
        HTTPException(422, "That record reference is not valid. Go back and pick it again."),
    )


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception):
    """Last line of defence: nothing from this app ever reaches the guest.

    The full traceback goes to the server log with the request path; the user
    gets the same calm brand treatment the HTTPException path uses, with a
    message that never leaks internals.
    """
    _log.exception("Unhandled error on %s %s", request.method, request.url.path)
    return await http_exception_handler(
        request,
        HTTPException(500, "Something went wrong on our side. The details were logged; please try again in a moment."),
    )


def csrf(request: Request) -> str:
    value = request.session.get("csrf")
    if not value:
        value = secrets.token_urlsafe(32)
        request.session["csrf"] = value
    return value


def check_csrf(request: Request, submitted: str | None) -> None:
    expected = request.session.get("csrf", "")
    if not expected or not isinstance(submitted, str) or not secrets.compare_digest(expected, submitted):
        raise HTTPException(status_code=403, detail="Your form expired. Refresh the page and try again.")


STAFF_ALLOWED_SECTIONS = (
    "/dashboard", "/orders", "/billing", "/invoices", "/payments",
    "/menu", "/rooms", "/reservations", "/customers", "/kots", "/mobile", "/kitchen",
    "/password/change",
)


def user_or_redirect(request: Request, db: Session) -> User:
    user = auth.current_user(request, db)
    if user.must_change_password and request.url.path != "/password/change":
        raise HTTPException(303, headers={"Location": "/password/change"})
    if user.role != "admin":
        path = request.url.path
        if not any(path == section or path.startswith(f"{section}/") for section in STAFF_ALLOWED_SECTIONS):
            raise HTTPException(
                403,
                "Staff accounts can only access the Operations and Finance sections.",
            )
    return user


def header_dates() -> dict:
    """English + Bikram Sambat dates for the sticky header banner."""
    now_local = local_now()
    english = now_local.strftime("%A, %d %B %Y")
    bs = to_bs(now_local.date().isoformat())
    return {"english": english, "nepali": bs[1] if bs else "", "short": bs[0] if bs else ""}


def render(request: Request, template: str, db: Session, context: dict | None = None,
           status_code: int = 200, headers: dict | None = None) -> HTMLResponse:
    values = {"request": request, "current_user": None, "csrf_token": csrf(request),
              "header_dates": header_dates(), "server_toast": toast_from_request(request)}
    if request.session.get("user_id"):
        values["current_user"] = db.get(User, request.session["user_id"])
    if context:
        values.update(context)
    return templates.TemplateResponse(request=request, name=template, context=values,
                                      status_code=status_code, headers=headers)


def business_window(day):
    start_local = datetime.combine(day, time.min, BUSINESS_TZ)
    end_local = start_local + timedelta(days=1)
    return (start_local.astimezone(timezone.utc).replace(tzinfo=None),
            end_local.astimezone(timezone.utc).replace(tzinfo=None))


def dashboard_data(db: Session, room_filter: str = "attention") -> dict:
    now_local = local_now()
    start, end = business_window(now_local.date())
    sales_invoices = [
        invoice for invoice in db.scalars(select(Invoice).options(
            selectinload(Invoice.guest), selectinload(Invoice.room)
        ).where(
            Invoice.issue_date >= start, Invoice.issue_date < end,
            Invoice.status != "void",
        )).all()
        if not _is_demo_invoice(invoice)
    ]
    gross = sum(float(invoice.total + invoice.discount_amount) for invoice in sales_invoices)
    net = sum(float(invoice.total) for invoice in sales_invoices)
    cash = db.scalar(select(func.coalesce(func.sum(models.Payment.amount), 0)).where(
        models.Payment.paid_at >= start, models.Payment.paid_at < end,
        models.Payment.method == "Cash",
    )) or 0
    rooms = [room for room in db.scalars(
        select(Room).where(Room.active.is_(True)).order_by(Room.number)
    ).all() if not _is_demo_room(room)]
    active_stays = db.scalars(
        select(Stay).options(selectinload(Stay.guest)).where(Stay.status == "checked_in")
    ).all()
    stay_by_room = {
        stay.room_id: stay for stay in active_stays
        if not _is_demo_guest(stay.guest)
    }
    room_rows = []
    balances = dict(db.execute(
        select(Invoice.room_id, func.sum(Invoice.balance))
        .where(Invoice.balance > 0, Invoice.status == "issued", Invoice.room_id.is_not(None))
        .group_by(Invoice.room_id)
    ).all())
    today = now_local.date()
    for room in rooms:
        stay = stay_by_room.get(room.id)
        planned_checkout = (stay.expected_checkout or stay.check_out) if stay else None
        checkout_today = bool(planned_checkout and planned_checkout.date() <= today)
        room_rows.append({
            "room": room, "stay": stay,
            "guest": stay.guest.name if stay else "",
            "checkout": (stay.expected_checkout or stay.check_out) if stay else None,
            "checkout_today": checkout_today,
            "balance": float(balances.get(room.id, 0) or 0),
        })
    if room_filter == "occupied":
        room_rows = [row for row in room_rows if row["room"].status == "occupied"]
    elif room_filter == "available":
        room_rows = [row for row in room_rows if row["room"].status == "available"]
    elif room_filter == "dirty":
        room_rows = [row for row in room_rows if row["room"].housekeeping_status == "dirty"]
    elif room_filter == "all":
        pass
    else:
        room_rows = [
            row for row in room_rows
            if row["room"].status == "occupied"
            or row["room"].housekeeping_status == "dirty"
            or row["checkout_today"]
        ]

    outstanding_candidates = db.scalars(
        select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room))
        .where(Invoice.balance > 0, Invoice.status == "issued")
        .order_by(Invoice.balance.desc(), Invoice.issue_date.desc())
    ).all()
    outstanding_candidates = [invoice for invoice in outstanding_candidates if not _is_demo_invoice(invoice)]
    outstanding = outstanding_candidates[:5]
    outstanding_total = sum(float(invoice.balance) for invoice in outstanding_candidates)
    recent_candidates = db.scalars(
        select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room))
        .where(Invoice.status != "void").order_by(Invoice.issue_date.desc(), Invoice.id.desc())
        .limit(60)
    ).all()
    recent = [invoice for invoice in recent_candidates if not _is_demo_invoice(invoice)][:10]
    occupied = sum(room.status == "occupied" for room in rooms)
    capacity = len(rooms)
    rate = round(occupied / capacity * 100) if capacity else 0
    room_revenue = sum(float(inv.total) for inv in sales_invoices if inv.room_id is not None)
    checkouts_today = sum(1 for row in room_rows if row["checkout_today"])
    available = sum(room.status == "available" for room in rooms)
    return {
        "business_date": now_local.strftime("%A, %d %B %Y"), "rooms": room_rows,
        "room_filter": room_filter, "gross": gross, "net": net,
        "outstanding": outstanding, "outstanding_total": float(outstanding_total),
        "outstanding_count": len(outstanding_candidates),
        "recent": recent, "cash_today": float(cash), "invoice_count": len(sales_invoices),
        "room_revenue": room_revenue, "occupied": occupied, "capacity": capacity,
        "occupancy_rate": rate, "available": available,
        "checkouts_today": checkouts_today,
    }



# --- login throttling -------------------------------------------------------
# Per (client IP, username) sliding window. In-process only: with several
# workers each keeps its own counters, so a deployment with >1 worker should
# move this to the database or a shared cache.
LOGIN_MAX_FAILURES = max(1, int(os.getenv("ARATI_LOGIN_MAX_FAILURES", "5")))
LOGIN_WINDOW_SECONDS = max(30, int(os.getenv("ARATI_LOGIN_WINDOW_SECONDS", "900")))
_login_failures: dict[tuple[str, str], list[float]] = {}
_login_lock = threading.Lock()
# Paid once at import so a missing user costs the same PBKDF2 work as a real
# one: without it, response timing reveals which usernames exist.
_DUMMY_PASSWORD_HASH = auth.hash_password(secrets.token_hex(16))


def _login_identity(request: Request, username: str) -> tuple[str, str]:
    client = request.client.host if request.client else "unknown"
    return (client, username.strip().lower())


def _login_retry_after(key: tuple[str, str]) -> int:
    now = _time.monotonic()
    with _login_lock:
        attempts = [t for t in _login_failures.get(key, []) if now - t < LOGIN_WINDOW_SECONDS]
        if attempts:
            _login_failures[key] = attempts
        else:
            _login_failures.pop(key, None)
        if len(attempts) < LOGIN_MAX_FAILURES:
            return 0
        return max(1, int(LOGIN_WINDOW_SECONDS - (now - attempts[0])) + 1)


def _record_login_failure(key: tuple[str, str]) -> None:
    now = _time.monotonic()
    with _login_lock:
        attempts = [t for t in _login_failures.get(key, []) if now - t < LOGIN_WINDOW_SECONDS]
        attempts.append(now)
        _login_failures[key] = attempts


def _clear_login_failures(key: tuple[str, str]) -> None:
    with _login_lock:
        _login_failures.pop(key, None)


@app.get("/", include_in_schema=False)
def home(request: Request):
    return RedirectResponse("/dashboard" if request.session.get("user_id") else "/login", status_code=303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, db: Session = Depends(get_db)):
    return render(request, "login.html", db)


@app.post("/login")
def login(
    request: Request, username: str = Form(""), password: str = Form(""),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    try:
        check_csrf(request, csrf_token)
    except HTTPException:
        return render(request, "login.html", db, {"error": "Your session or form expired. Please try again."}, 403)
    username = username.strip().lower()
    if not username or not password:
        return render(request, "login.html", db, {"error": "Please enter both username and password."}, 401)
    key = _login_identity(request, username)
    retry_after = _login_retry_after(key)
    if retry_after:
        return render(request, "login.html", db, {
            "error": "Too many failed sign-in attempts. Please wait a few minutes and try again.",
        }, 429, headers={"Retry-After": str(retry_after)})
    user = db.scalar(select(User).where(User.username == username, User.active.is_(True)))
    password_ok = auth.verify_password(
        password, user.password_hash if user is not None else _DUMMY_PASSWORD_HASH
    )
    if user is None or not password_ok:
        _record_login_failure(key)
        db.add(AuditLog(user_id=user.id if user is not None else None, username=username,
                        action="login.failed", entity="user",
                        entity_id=str(user.id) if user is not None else "",
                        detail=f"client={key[0]}"))
        db.commit()
        return render(request, "login.html", db, {"error": "Username or password is incorrect."}, 401)
    _clear_login_failures(key)
    request.session.clear()
    request.session["user_id"] = user.id
    request.session["sv"] = user.session_version or 0
    csrf(request)
    destination = "/password/change" if user.must_change_password else "/dashboard"
    return RedirectResponse(destination, status_code=303)


@app.get("/password/change", response_class=HTMLResponse)
def required_password_change_page(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if not user.must_change_password:
        return RedirectResponse("/dashboard", status_code=303)
    return render(request, "password_change.html", db)


@app.post("/password/change", response_class=HTMLResponse)
def required_password_change(
    request: Request, new_password: str = Form(""), confirm_password: str = Form(""),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if not user.must_change_password:
        return RedirectResponse("/dashboard", status_code=303)
    try:
        check_csrf(request, csrf_token)
    except HTTPException:
        return render(request, "password_change.html", db, {
            "error": "Your form expired. Refresh the page and try again.",
        }, status_code=403)
    if len(new_password) < 12:
        return render(request, "password_change.html", db, {
            "error": "Use a password with at least 12 characters.",
        }, status_code=422)
    if new_password != confirm_password:
        return render(request, "password_change.html", db, {
            "error": "The passwords do not match.",
        }, status_code=422)
    user.password_hash = auth.hash_password(new_password)
    user.must_change_password = False
    user.session_version = (user.session_version or 0) + 1
    db.commit()
    # Keep this browser signed in; every other cookie for this account is now stale.
    request.session["sv"] = user.session_version
    return RedirectResponse(toast_url("/dashboard", "Password updated.", "success"), status_code=303)


@app.post("/logout")
def logout(request: Request, csrf_token: str = Form("")):
    try:
        check_csrf(request, csrf_token)
    except HTTPException:
        pass
    user_id = request.session.get("user_id")
    record_id = request.session.get("_session_record_id")
    if record_id and user_id:
        db = SessionLocal()
        try:
            row = db.get(UserSession, record_id)
            if row is not None and row.user_id == int(user_id) and row.revoked_at is None:
                row.revoked_at = models.utcnow()
                row.revoked_reason = "Signed out"
                db.commit()
        except Exception:
            db.rollback()
        finally:
            db.close()
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request, room_filter: str = Query("attention"),
              db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    return render(request, "dashboard.html", db, dashboard_data(db, room_filter))


@app.get("/dashboard/sales", response_class=HTMLResponse)
def dashboard_sales(request: Request, period: str = Query("30"),
                    db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    try:
        data = web_analytics.sales_details(db, period)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return render(request, "dashboard_sales.html", db, data)


@app.get("/dashboard/rooms", response_class=HTMLResponse)
def dashboard_rooms(request: Request, room_filter: str = Query("attention"),
                    db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    data = dashboard_data(db, room_filter)
    return render(request, "partials/room_rows.html", db, data)


ROOM_STATUS_PRIORITY = {"occupied": 0, "dirty": 1, "reserved": 2, "available": 3, "maintenance": 4}


def _filter_room_rows(rows: list[dict], filters: str = "attention", q: str = "") -> list[dict]:
    """Filter by status/search and put rooms that need attention first."""
    selected = {part.strip() for part in filters.split(",")}
    today = local_now().date()
    if "attention" in selected or "all" not in selected:
        if "attention" in selected:
            rows = [
                row for row in rows
                if row["status"] in {"occupied", "dirty"}
                or (row["stay"] and (row["stay"].expected_checkout or row["stay"].check_out)
                    and (row["stay"].expected_checkout or row["stay"].check_out).date() <= today)
            ]
        else:
            valid = selected & set(ROOM_STATUS_PRIORITY)
            if valid:
                rows = [row for row in rows if row["status"] in valid]
    if q.strip():
        term = q.strip().casefold()
        rows = [row for row in rows if term in row["room"].number.casefold()
                or term in row["room"].type.casefold()
                or (row["guest"] and term in row["guest"].name.casefold())]
    return sorted(rows, key=lambda row: (ROOM_STATUS_PRIORITY.get(row["status"], 9), row["room"].number))


@app.get("/rooms", response_class=HTMLResponse)
def rooms_page(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    data = _room_board_data(db)
    data["room_rows"] = _filter_room_rows(data["room_rows"], "all")
    data["view"] = "grid"
    return render(request, "rooms.html", db, data)


def _room_board_data(db: Session) -> dict:
    rooms = [room for room in db.scalars(
        select(Room).where(Room.active.is_(True)).order_by(Room.number)
    ).all() if not _is_demo_room(room)]
    stays = db.scalars(select(Stay).options(selectinload(Stay.guest)).where(Stay.status == "checked_in")).all()
    now = local_now()
    day_start, day_end = business_window(now.date())
    reservations = db.scalars(
        select(Reservation).options(selectinload(Reservation.guest))
        .where(Reservation.status == "confirmed")
        .order_by(Reservation.check_in)
    ).all()
    stay_by_room = {stay.room_id: stay for stay in stays}
    reservations_by_room = {}
    for reservation in reservations:
        reservations_by_room.setdefault(reservation.room_id, []).append(reservation)
    balances = dict(db.execute(
        select(Invoice.room_id, func.sum(Invoice.balance))
        .where(Invoice.room_id.is_not(None), Invoice.balance > 0, Invoice.status == "issued")
        .group_by(Invoice.room_id)
    ).all())
    rows = []
    for room in rooms:
        stay = stay_by_room.get(room.id)
        room_reservations = reservations_by_room.get(room.id, [])
        reservation = next(
            (item for item in room_reservations
             if item.check_in < day_end and item.check_out > day_start),
            room_reservations[0] if room_reservations else None,
        )
        if room.housekeeping_status == "maintenance" or room.status == "maintenance":
            status = "maintenance"
        elif stay:
            status = "occupied"
        elif reservation and reservation.check_in < day_end and reservation.check_out > day_start:
            status = "reserved"
        elif room.housekeeping_status in {"dirty", "needs_cleaning"} or room.status == "dirty":
            status = "dirty"
        else:
            status = "available"
        
        # Calculate stay duration for occupied rooms
        stay_duration = None
        if stay and stay.check_in:
            check_in_local = stay.check_in.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ)
            delta = now - check_in_local
            total_minutes = int(delta.total_seconds() // 60)
            hours = total_minutes // 60
            minutes = total_minutes % 60
            if hours > 0:
                stay_duration = f"{hours}h {minutes}m"
            else:
                stay_duration = f"{minutes}m"

        rows.append({
            "room": room, "status": status, "stay": stay, "reservation": reservation,
            "guest": stay.guest if stay else reservation.guest if reservation else None,
            "balance": float(balances.get(room.id, 0) or 0),
            "stay_duration": stay_duration,
        })
    counts = {name: sum(row["status"] == name for row in rows) for name in
              ("available", "occupied", "dirty", "reserved", "maintenance")}
    total = len(rows)
    percentages = {name: round(count / total * 100) if total else 0 for name, count in counts.items()}
    return {"room_rows": rows, "room_counts": counts, "room_percentages": percentages,
            "room_total": total, "business_date": local_now().strftime("%d %b %Y")}


@app.get("/rooms/board", response_class=HTMLResponse)
def room_board_partial(request: Request, filters: str = "all", q: str = "",
                       view: str = "grid", db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    rows = _filter_room_rows(_room_board_data(db)["room_rows"], filters, q)
    return render(request, "partials/room_board_content.html", db,
                   {"room_rows": rows, "view": view if view in {"grid", "list"} else "grid"})


@app.get("/rooms/board/summary", response_class=HTMLResponse)
def room_board_summary(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    return render(request, "partials/room_summary.html", db, _room_board_data(db))


def _room_action_response(message: str | None = None, kind: str = "success") -> Response:
    triggers = {"room-board-refresh": {}, "room-modal-close": {}}
    if message:
        triggers["showToast"] = toast_event(message, kind)
    return Response(status_code=204, headers={
        "HX-Trigger": json.dumps(triggers)
    })


def _guest_from_form(db: Session, value: str) -> Guest | None:
    if not value.strip():
        return None
    try:
        return db.get(Guest, int(value))
    except ValueError as exc:
        raise HTTPException(422, "Select a valid guest.") from exc


def _parse_local_datetime(value: str, field_name: str) -> datetime:
    try:
        return datetime.fromisoformat(value).replace(tzinfo=BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    except ValueError as exc:
        raise HTTPException(422, f"Enter a valid {field_name} date and time.") from exc


def _business_date(value: datetime):
    return value.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ).date()


def _audit(db: Session, user: User, action: str, entity: str, entity_id: str, detail: str) -> None:
    db.add(AuditLog(
        user_id=user.id, username=user.username, action=action,
        entity=entity, entity_id=entity_id, detail=detail,
    ))


MAX_MONEY_INPUT = Decimal("100000000")  # 10 crore ceiling per field
MAX_QUANTITY_INPUT = Decimal("10000")


def _finite(value: Decimal, label: str) -> Decimal:
    """Reject NaN/Infinity before any comparison or arithmetic touches it."""
    if not value.is_finite():
        raise ValueError(f"{label} must be a finite number.")
    return value


def _bounded_money(value: Decimal, label: str, maximum: Decimal = MAX_MONEY_INPUT) -> Decimal:
    _finite(value, label)
    if value < 0 or value > maximum:
        raise ValueError(f"{label} must be between 0 and {maximum}.")
    return value


def _draft_access(invoice: Invoice, user: User) -> None:
    """A draft belongs to the person who created it; admins may open any."""
    if invoice.user_id is not None and invoice.user_id != user.id and user.role != "admin":
        raise HTTPException(403, "This draft invoice belongs to another staff account.")


def _stay_folio(db: Session, stay: Stay) -> dict:
    room = db.get(Room, stay.room_id)
    if room is None:
        raise HTTPException(404, "The stay's room no longer exists.")
    charges = db.scalars(
        select(RoomCharge).where(RoomCharge.stay_id == stay.id)
        .order_by(RoomCharge.charge_date, RoomCharge.id)
    ).all()
    departure = stay.expected_checkout or stay.check_out
    nights = 1
    if departure:
        nights = max(1, (_business_date(departure) - _business_date(stay.check_in)).days)
    nights += sum(
        1 for charge in charges
        if charge.category == "room" and charge.description.startswith("Extra night")
    )
    lines = [Line(
        description=f"Room {room.number} — {nights} night(s)",
        qty=float(nights), unit_price=float(stay.nightly_rate or room.nightly_rate),
        category="room",
    )]
    lines.extend(Line(
        description=charge.description,
        qty=float(charge.quantity),
        unit_price=float(charge.unit_price),
        category=charge.category,
    ) for charge in charges)
    settings = _billing_settings(db)
    try:
        service_charge_rate = float(settings.get("service_charge_rate", "10"))
        vat_rate = float(settings.get("vat_rate", "13"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(500, "Hotel billing rates are invalid. Update Hotel Settings.") from exc
    totals = compute_totals(
        lines, service_charge_rate=service_charge_rate,
        vat_rate=vat_rate, bill_type="tax",
    )
    return {"room": room, "charges": charges, "lines": lines, "nights": nights, "totals": totals}


@app.get("/rooms/{room_id}/checkin", response_class=HTMLResponse)
def room_checkin_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    room = db.get(Room, room_id)
    room_row = next(
        (row for row in _room_board_data(db)["room_rows"] if row["room"].id == room_id), None
    )
    if not room or not room_row or room_row["status"] not in {"available", "reserved"}:
        raise HTTPException(409, "This room cannot be checked in right now.")
    start, end = business_window(local_now().date())
    reservation = db.scalar(select(Reservation).options(selectinload(Reservation.guest)).where(
        Reservation.room_id == room.id, Reservation.status == "confirmed",
        Reservation.check_in < end, Reservation.check_out > start,
    ).order_by(Reservation.check_in))
    return render(request, "partials/room_checkin_modal.html", db, {
        "room": room, "reservation": reservation,
        "guests": db.scalars(select(Guest).order_by(Guest.name)).all(),
        "now_local": local_now().strftime("%Y-%m-%dT%H:%M"),
        "checkout_local": (
            local_datetime(reservation.check_out, "%Y-%m-%dT%H:%M") if reservation else ""
        ),
        "guest_name_val": (
            reservation.guest_name or (reservation.guest.name if reservation.guest else "")
            if reservation else ""
        ),
        "phone_val": (
            reservation.guest_phone or (reservation.guest.phone if reservation.guest else "")
            if reservation else ""
        ),
        "email_val": (
            reservation.guest_email or (reservation.guest.email if reservation.guest else "")
            if reservation else ""
        ),
        "address_val": (
            reservation.guest_address or (reservation.guest.address if reservation.guest else "")
            if reservation else ""
        ),
    })


@app.post("/rooms/{room_id}/checkin")
async def room_checkin(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    room = db.get(Room, room_id)
    if not room:
        raise HTTPException(404, "Room not found.")

    start, end = business_window(local_now().date())
    reservation = db.scalar(select(Reservation).options(selectinload(Reservation.guest)).where(
        Reservation.room_id == room.id, Reservation.status == "confirmed",
        Reservation.check_in < end, Reservation.check_out > start,
    ).order_by(Reservation.check_in))
    guests = db.scalars(select(Guest).order_by(Guest.name)).all()

    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        room_row = next(
            (row for row in _room_board_data(db)["room_rows"] if row["room"].id == room_id), None
        )
        if not room_row or room_row["status"] not in {"available", "reserved"}:
            raise ValueError("This room cannot be checked in right now (it may need cleaning or already be occupied).")

        guest_id = str(form.get("guest_id", "")).strip()
        guest = _guest_from_form(db, guest_id)
        if reservation and guest and reservation.guest_id and guest.id != reservation.guest_id:
            raise ValueError(f"Room {room.number} is reserved for {reservation.guest.name}.")
        if reservation and guest is None:
            guest = reservation.guest

        guest_name = str(form.get("guest_name", "")).strip()
        guest_phone = str(form.get("phone", "")).strip()
        guest_email = str(form.get("email", "")).strip()[:180]
        guest_address = str(form.get("address", "")).strip()
        guest_pan = str(form.get("pan_vat", "")).strip() or str(form.get("guest_pan", "")).strip()
        guest_citizenship = str(form.get("citizenship_no", "")).strip()
        if not guest_name and reservation:
            guest_name = reservation.guest_name or (guest.name if guest else "")
        if not guest_phone and reservation:
            guest_phone = reservation.guest_phone or (guest.phone if guest else "")
        if not guest_email and reservation:
            guest_email = reservation.guest_email or (guest.email if guest else "")
        if not guest_address and reservation:
            guest_address = reservation.guest_address or (guest.address if guest else "")
        if not guest_pan and reservation:
            guest_pan = reservation.guest_pan

        if not guest_name and not guest:
            raise ValueError("Enter a guest name for check-in.")

        guest = guest_service.get_or_create_guest(
            db,
            name=guest_name or (guest.full_name if guest else "") or (guest.name if guest else ""),
            phone=guest_phone or (guest.phone if guest else ""),
            email=guest_email or (guest.email if guest else ""),
            address=guest_address or (guest.address if guest else ""),
            pan_vat=guest_pan or (getattr(guest, "pan_vat_number", "") if guest else ""),
            id_type="Citizenship" if guest_citizenship else (guest.id_type if guest else ""),
            id_number=guest_citizenship or (guest.id_number if guest else ""),
            customer_type=(guest.customer_type if guest else "individual") or "individual",
        )

        check_in = _parse_local_datetime(str(form.get("check_in", "")), "check-in")
        checkout_raw = str(form.get("check_out", "")).strip()
        check_out = _parse_local_datetime(checkout_raw, "check-out") if checkout_raw else None
        if check_out is not None and check_out <= check_in:
            raise ValueError("Check-out date and time must be after check-in.")
        if reservation and not (reservation.check_in <= check_in < reservation.check_out):
            raise ValueError("Check-in must fall within the reservation dates.")
        if check_out is None and reservation:
            check_out = reservation.check_out
        if reservation:
            reservation.status = "checked_in"

        rate_raw = room.nightly_rate
        if reservation and reservation.nightly_rate > 0:
            rate_raw = reservation.nightly_rate
        try:
            nightly_rate = Decimal(str(rate_raw or 0))
        except Exception:
            nightly_rate = room.nightly_rate or Decimal("0")

        stay = Stay(
            room_id=room.id,
            guest_id=guest.id,
            check_in=check_in,
            check_out=None,
            expected_checkout=check_out,
            nights=(
                max(1, (_business_date(check_out) - _business_date(check_in)).days)
                if check_out else 1
            ),
            nightly_rate=nightly_rate,
            rate_per_night=nightly_rate,
            room_number=room.number,
            guest_name=guest.full_name or guest.name,
            guest_phone=guest.phone or "",
            guest_email=guest.email or "",
            guest_address=guest.address or "",
            guest_pan=guest.pan_vat_number or guest_pan or "",
            business_date=_business_date(check_in).isoformat(),
            created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            updated_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            user_id=user.id,
            status="checked_in",
        )
        db.add(stay)
        db.flush()

        room.status = "occupied"
        room.housekeeping_status = "clean"

        settings = _billing_settings(db)
        invoice_service.get_or_create_draft_invoice(
            db,
            guest_id=guest.id,
            stay_id=stay.id,
            room_id=room.id,
            user_id=user.id,
            settings=settings,
        )

        db.commit()
        return _room_action_response(f'Room {room.number} checked in to {(guest.full_name or guest.name)}.')
    except HTTPException as exc:
        db.rollback()
        if exc.status_code != 422:
            raise
        error_msg = exc.detail
        return render(request, "partials/room_checkin_modal.html", db, {
            "room": room, "reservation": reservation, "guests": guests,
            "now_local": str(form.get("check_in", "")) or local_now().strftime("%Y-%m-%dT%H:%M"),
            "checkout_local": str(form.get("check_out", "")),
            "error": error_msg,
            "guest_name_val": str(form.get("guest_name", "")),
            "phone_val": str(form.get("phone", "")),
            "email_val": str(form.get("email", "")),
            "address_val": str(form.get("address", "")),
            "citizenship_no_val": str(form.get("citizenship_no", "")),
            "selected_guest_id": str(form.get("guest_id", "")),
        }, status_code=200)
    except ValueError as exc:
        db.rollback()
        error_msg = str(exc)
        return render(request, "partials/room_checkin_modal.html", db, {
            "room": room, "reservation": reservation, "guests": guests,
            "now_local": str(form.get("check_in", "")) or local_now().strftime("%Y-%m-%dT%H:%M"),
            "checkout_local": str(form.get("check_out", "")),
            "error": error_msg,
            "guest_name_val": str(form.get("guest_name", "")),
            "phone_val": str(form.get("phone", "")),
            "email_val": str(form.get("email", "")),
            "address_val": str(form.get("address", "")),
            "citizenship_no_val": str(form.get("citizenship_no", "")),
            "selected_guest_id": str(form.get("guest_id", "")),
        }, status_code=200)


def _get_guest_active_stays(db: Session, guest_id: int) -> list[Stay]:
    """Get all checked-in stays for a guest, ordered by check_in (first room first)."""
    return db.scalars(
        select(Stay).where(
            Stay.guest_id == guest_id, Stay.status == "checked_in"
        ).order_by(Stay.check_in)
    ).all()


def _combine_guest_folios(db: Session, stays: list[Stay]) -> dict:
    """Combine multiple stays for the same guest into a single folio."""
    if not stays:
        return {"room": None, "charges": [], "lines": [], "nights": 0, "totals": None, "stays": []}
    
    all_charges = []
    all_lines = []
    total_nights = 0
    first_stay = stays[0]
    first_room = db.get(Room, first_stay.room_id)
    
    # Collect all charges from all stays
    for stay in stays:
        room = db.get(Room, stay.room_id)
        charges = db.scalars(
            select(RoomCharge).where(RoomCharge.stay_id == stay.id)
            .order_by(RoomCharge.charge_date, RoomCharge.id)
        ).all()
        all_charges.extend(charges)
        
        departure = stay.expected_checkout or stay.check_out
        nights = 1
        if departure:
            nights = max(1, (_business_date(departure) - _business_date(stay.check_in)).days)
        nights += sum(
            1 for charge in charges
            if charge.category == "room" and charge.description.startswith("Extra night")
        )
        total_nights += nights
        
        # Room line for this stay
        all_lines.append(Line(
            description=f"Room {room.number} — {nights} night(s)",
            qty=float(nights),
            unit_price=float(stay.nightly_rate or room.nightly_rate),
            category="room",
        ))
        
        # Additional charges for this stay
        for charge in charges:
            all_lines.append(Line(
                description=charge.description,
                qty=float(charge.quantity),
                unit_price=float(charge.unit_price),
                category=charge.category,
            ))
    
    settings = _billing_settings(db)
    try:
        service_charge_rate = float(settings.get("service_charge_rate", "10"))
        vat_rate = float(settings.get("vat_rate", "13"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(500, "Hotel billing rates are invalid. Update Hotel Settings.") from exc
    
    totals = compute_totals(
        all_lines, service_charge_rate=service_charge_rate,
        vat_rate=vat_rate, bill_type="tax",
    )
    
    return {
        "room": first_room,
        "charges": all_charges,
        "lines": all_lines,
        "nights": total_nights,
        "totals": totals,
        "stays": stays,
    }


@app.get("/rooms/{room_id}/checkout", response_class=HTMLResponse)
def room_checkout_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    row = next((row for row in _room_board_data(db)["room_rows"] if row["room"].id == room_id), None)
    if not row or not row["stay"]:
        raise HTTPException(404, "No active stay found for this room.")
    
    stay = row["stay"]
    
    # Find all active stays for this guest
    guest_stays = _get_guest_active_stays(db, stay.guest_id) if stay.guest_id else [stay]
    is_multi_room = len(guest_stays) > 1
    
    # Combine folios for all rooms
    combined_folio = _combine_guest_folios(db, guest_stays)
    
    # Check if any stay has an existing invoice
    invoiced = None
    if stay.invoice_number:
        invoiced = db.scalar(
            select(Invoice).where(Invoice.invoice_number == stay.invoice_number)
        )
    # Also check other stays
    if not invoiced:
        for s in guest_stays:
            if s.invoice_number and s.invoice_number != stay.invoice_number:
                invoiced = db.scalar(
                    select(Invoice).where(Invoice.invoice_number == s.invoice_number)
                )
                if invoiced:
                    break
    
    return render(request, "partials/room_checkout_modal.html", db, {
        "row": row,
        "folio": combined_folio,
        "payment_methods": _payment_methods(db),
        "invoiced_invoice": invoiced,
        "is_multi_room": is_multi_room,
        "guest_stays": guest_stays,
    })


@app.post("/rooms/{room_id}/checkout")
def room_checkout(
    request: Request, room_id: int, csrf_token: str = Form(""),
    amount_paid: str = Form("0"), payment_method: str = Form("Cash"),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    try:
        check_csrf(request, csrf_token)
        room = db.get(Room, room_id)
        stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
        if not room or not stay:
            raise ValueError("No active stay found for this room.")
        
        # Find all active stays for this guest
        guest_stays = _get_guest_active_stays(db, stay.guest_id) if stay.guest_id else [stay]
        is_multi_room = len(guest_stays) > 1
        
        # Check if any stay has an existing invoice
        invoiced_invoice = None
        for s in guest_stays:
            if s.invoice_number:
                invoiced_invoice = db.scalar(
                    select(Invoice).where(Invoice.invoice_number == s.invoice_number)
                )
                if invoiced_invoice:
                    break
            if invoiced_invoice:
                break
        
        if invoiced_invoice is not None:
            # If there's an existing invoice, we need to handle carefully
            # For now, just check out the single room that triggered this
            # (multi-room with existing invoices is complex)
            if is_multi_room:
                raise ValueError("Cannot check out multiple rooms when one has an existing invoice. Check out rooms individually.")
        
        try:
            paid_value = Decimal(amount_paid or "0")
        except InvalidOperation as exc:
            raise ValueError("Enter a valid amount received.") from exc
        if not paid_value.is_finite() or paid_value < 0:
            raise ValueError("Amount received cannot be negative.")
        if payment_method not in _payment_methods(db):
            raise ValueError("Choose a valid payment method.")
        
        # Combine folios for all stays
        combined_folio = _combine_guest_folios(db, guest_stays)
        
        if invoiced_invoice is not None:
            balance = (
                Decimal(str(invoiced_invoice.balance)) if invoiced_invoice.balance is not None
                else Decimal("0.00")
            ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if paid_value > balance:
                raise ValueError(
                    f"Amount received cannot exceed the outstanding invoice balance of {balance}."
                )
        else:
            totals = compute_totals(
                combined_folio["lines"],
                service_charge_rate=combined_folio["totals"].service_charge_rate,
                vat_rate=combined_folio["totals"].vat_rate,
                paid=float(paid_value),
                bill_type="tax",
            )
            if paid_value > Decimal(str(totals.total)):
                raise ValueError("Amount received cannot exceed the folio total.")
        
        if stay.guest_id is None or db.get(Guest, stay.guest_id) is None:
            raise ValueError("The active stay has no valid guest record.")
        guest = db.get(Guest, stay.guest_id)
        now = local_now()
        issue_date = now.astimezone(timezone.utc).replace(tzinfo=None)
        settings = _billing_settings(db)
        
        if invoiced_invoice is None:
            # Create single invoice for all rooms
            invoice = Invoice(
                is_draft=False,
                invoice_number=_next_billing_number(db), guest_id=guest.id, room_id=room.id,
                stay_id=stay.id, issue_date=issue_date,
                subtotal=Decimal(str(totals.subtotal)),
                item_discount=Decimal(str(totals.item_discount)),
                discount_amount=Decimal(str(totals.discount)),
                net_subtotal=Decimal(str(totals.net_subtotal)),
                taxable_net=Decimal(str(totals.taxable_net)),
                service_charge_rate=Decimal(str(totals.service_charge_rate)),
                service_charge=Decimal(str(totals.service_charge)),
                vat_rate=Decimal(str(totals.vat_rate)),
                vat_amount=Decimal(str(totals.vat)),
                total=Decimal(str(totals.total)),
                paid_amount=Decimal(str(totals.paid)),
                balance=Decimal(str(totals.balance)),
                status="paid" if totals.balance == 0 else "issued",
                invoice_type="Tax Invoice", bill_type="tax",
                language=settings.get("invoice_language", "en"),
                copy_type=settings.get("invoice_copy", "Customer Copy"),
                payment_terms=settings.get("payment_terms", "Payment due on presentation."),
                notes=f"Combined folio for {len(guest_stays)} room(s)" if is_multi_room else f"Room {room.number} folio",
                payment_method=payment_method,
                prepared_by=user.username, guest_name=stay.guest_name or guest.name,
                guest_phone=stay.guest_phone or guest.phone,
                guest_address=stay.guest_address or guest.address,
                guest_pan=stay.guest_pan or guest.pan_vat_number,
                room_number=room.number, bill_date=now.date().isoformat(),
                nights=combined_folio["nights"], check_in=stay.check_in,
                check_out=stay.expected_checkout or stay.check_out, user_id=user.id,
            )
            invoice.guest = guest
            for line in combined_folio["lines"]:
                line_vat = compute_totals(
                    [line], service_charge_rate=0,
                    vat_rate=totals.vat_rate, bill_type="tax",
                ).vat
                invoice.items.append(InvoiceItem(
                    description=line.description, quantity=line.qty, rate=line.unit_price,
                    amount=Decimal(str(line.total + line_vat)),
                    vat_amount=Decimal(str(line_vat)), category=line.category,
                    taxable=line.taxable, unit=line.unit,
                ))
                db.add(invoice)
                db.flush()
            
            # Check out ALL stays for this guest
            for s in guest_stays:
                s.status = "checked_out"
                s.check_out = issue_date
                s.nights = next(
                    (stay_nights for stay_nights in [combined_folio["nights"]] if s == guest_stays[0]),
                    0
                )
                # Calculate nights for this specific stay
                dep = s.expected_checkout or s.check_out
                s_nights = 1
                if dep:
                    s_nights = max(1, (_business_date(dep) - _business_date(s.check_in)).days)
                s_nights += sum(
                    1 for charge in combined_folio["charges"]
                    if charge.stay_id == s.id and charge.category == "room" and charge.description.startswith("Extra night")
                )
                s.nights = s_nights
                s.invoice_number = invoice.invoice_number
                s.room_number = db.get(Room, s.room_id).number if db.get(Room, s.room_id) else ""
                
                # Mark room as dirty
                r = db.get(Room, s.room_id)
                if r:
                    r.status = "dirty"
                    r.housekeeping_status = "dirty"
                
                # Update reservation if exists
                if s.reservation_id:
                    reservation = db.get(Reservation, s.reservation_id)
                    if reservation and reservation.status == "checked_in":
                        reservation.status = "checked_out"
            
            # Link all charges to the invoice
            for charge in combined_folio["charges"]:
                charge.invoice_number = invoice.invoice_number
            
            if totals.paid > 0:
                db.add(Payment(
                    invoice=invoice, amount=Decimal(str(totals.paid)),
                    method=payment_method, received_by=user.username,
                    note=f"Checkout {len(guest_stays)} room(s) for {guest.name}",
                    invoice_number=invoice.invoice_number,
                    business_date=now.date().isoformat(), user_id=user.id,
                ))
        else:
            # Existing invoice case - just check out this one room
            stay = guest_stays[0]
            stay.status = "checked_out"
            stay.check_out = issue_date
            stay.nights = combined_folio["nights"]
            room.status = "dirty"
            room.housekeeping_status = "dirty"
            if stay.reservation_id:
                reservation = db.get(Reservation, stay.reservation_id)
                if reservation and reservation.status == "checked_in":
                    reservation.status = "checked_out"
            if paid_value > 0:
                invoiced_invoice.paid_amount = (
                    Decimal(str(invoiced_invoice.paid_amount)) + paid_value
                ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                invoiced_invoice.balance = (
                    Decimal(str(invoiced_invoice.total)) - Decimal(str(invoiced_invoice.paid_amount))
                ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                invoiced_invoice.status = "paid" if invoiced_invoice.balance <= 0 else "issued"
                db.add(Payment(
                    invoice=invoiced_invoice, amount=paid_value,
                    method=payment_method, received_by=user.username,
                    note=f"Balance settled at Room {room.number} checkout",
                    invoice_number=invoiced_invoice.invoice_number,
                    business_date=now.date().isoformat(), user_id=user.id,
                ))
        
        _audit(db, user, "stay.check_out", "stay", str(stay.id),
               f"checked out {len(guest_stays)} room(s) for {guest.name}, invoiced as {invoice.invoice_number if invoiced_invoice is None else invoiced_invoice.invoice_number}")
        db.commit()
        
        if is_multi_room:
            return _room_action_response(f'Checked out {len(guest_stays)} rooms for {guest.name}. Combined invoice: {invoice.invoice_number if invoiced_invoice is None else invoiced_invoice.invoice_number}.')
        else:
            return _room_action_response(f'Room {room.number} checked out.' if not invoiced_invoice else f'Room {room.number} settled against {invoiced_invoice.invoice_number}.')
    
    except HTTPException as exc:
        db.rollback()
        if exc.status_code != 422:
            raise
        error_msg = exc.detail
        row = next((r for r in _room_board_data(db)["room_rows"] if r["room"].id == room_id), None)
        if row and row["stay"]:
            guest_stays = _get_guest_active_stays(db, row["stay"].guest_id) if row["stay"].guest_id else [row["stay"]]
            combined_folio = _combine_guest_folios(db, guest_stays)
            return render(request, "partials/room_checkout_modal.html", db, {
                "row": row, "folio": combined_folio, "error": error_msg,
                "payment_methods": _payment_methods(db),
                "payment_method_val": payment_method,
                "invoiced_invoice": invoiced_invoice,
                "is_multi_room": len(guest_stays) > 1,
                "guest_stays": guest_stays,
            }, status_code=200)
        raise
    except ValueError as exc:
        db.rollback()
        error_msg = str(exc)
        row = next((r for r in _room_board_data(db)["room_rows"] if r["room"].id == room_id), None)
        if row and row["stay"]:
            return render(request, "partials/room_checkout_modal.html", db, {
                "row": row, "folio": _stay_folio(db, row["stay"]), "error": error_msg,
                "payment_methods": _payment_methods(db),
                "payment_method_val": payment_method,
                "invoiced_invoice": invoiced_invoice,
            }, status_code=200)
        raise HTTPException(404, error_msg) from exc


@app.get("/rooms/{room_id}/extend", response_class=HTMLResponse)
def room_extend_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    room = db.get(Room, room_id)
    stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
    if room is None or stay is None:
        raise HTTPException(404, "No active stay found for this room.")
    return render(request, "partials/room_extend_modal.html", db, {
        "room": room,
        "stay": stay,
        "checkout_local": local_datetime(stay.expected_checkout or stay.check_out, "%Y-%m-%dT%H:%M"),
    })


@app.post("/rooms/{room_id}/extend")
async def room_extend(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
        if stay is None:
            raise ValueError("Only an active stay can be extended.")
        new_checkout = _parse_local_datetime(str(form.get("check_out", "")), "check-out")
        current_checkout = stay.expected_checkout or stay.check_out
        if current_checkout and new_checkout <= current_checkout:
            raise ValueError("The new check-out must be later than the current expected check-out.")
        overlap = db.scalar(select(Reservation.id).where(
            Reservation.room_id == room_id, Reservation.status == "confirmed",
            Reservation.check_in < new_checkout,
            Reservation.check_out > (stay.check_in),
        ).limit(1))
        if overlap is not None:
            raise ValueError("This extension conflicts with a confirmed reservation.")
        stay.expected_checkout = new_checkout
        stay.nights = max(1, (_business_date(new_checkout) - _business_date(stay.check_in)).days)
        _audit(db, user, "stay.extend", "stay", str(stay.id),
               f"room {room_id} expected check-out moved to {new_checkout.isoformat()}")
        db.commit()
        return _room_action_response(f'Room {room.number} stay extended.')
    except HTTPException as exc:
        db.rollback()
        if exc.status_code != 422:
            raise
        error_msg = exc.detail
        room = db.get(Room, room_id)
        stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
        if room is not None and stay is not None:
            return render(request, "partials/room_extend_modal.html", db, {
                "room": room, "stay": stay, "error": error_msg,
                "checkout_local": str(form.get("check_out", "")),
            }, status_code=200)
        raise
    except ValueError as exc:
        db.rollback()
        error_msg = str(exc)
        room = db.get(Room, room_id)
        stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
        if room is not None and stay is not None:
            return render(request, "partials/room_extend_modal.html", db, {
                "room": room, "stay": stay, "error": error_msg,
                "checkout_local": str(form.get("check_out", "")),
            }, status_code=200)
        raise HTTPException(404, error_msg)


@app.get("/rooms/{room_id}/charge", response_class=HTMLResponse)
def room_charge_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    room = db.get(Room, room_id)
    stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
    if room is None or stay is None:
        raise HTTPException(404, "Only an active stay can take charges.")
    return render(request, "partials/room_charge_modal.html", db, {"room": room, "stay": stay})


@app.post("/rooms/{room_id}/charge")
async def room_add_charge(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    room = db.get(Room, room_id)
    stay = db.scalar(select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in"))
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        if room is None or stay is None:
            raise ValueError("Only an active stay can take charges.")
        description = str(form.get("description", "")).strip()[:240]
        category = str(form.get("category", "service"))
        if category not in {"food", "beverage", "laundry", "room", "service"}:
            raise ValueError("Choose a valid charge category.")
        try:
            quantity = Decimal(str(form.get("quantity", "1")))
            unit_price = Decimal(str(form.get("unit_price", "0")))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("Quantity and rate must be valid numbers.") from exc
        if (
            not description or not quantity.is_finite() or not unit_price.is_finite()
            or quantity <= 0 or quantity > 365 or unit_price <= 0
            or unit_price > Decimal("9999999999.99")
        ):
            raise ValueError("Enter a description, positive quantity, and positive rate.")
        amount = (quantity * unit_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        db.add(RoomCharge(
            stay_id=stay.id, room_number=room.number, description=description,
            category=category, quantity=quantity, unit_price=unit_price,
            amount=amount, charge_date=local_now().date().isoformat(),
            created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            user_id=user.id,
        ))
        _audit(db, user, "room_charge.create", "stay", str(stay.id),
               f"{description} {amount:.2f}")
        db.commit()
        return _room_action_response(f'Charge added to room {room.number}.')
    except HTTPException as exc:
        db.rollback()
        if exc.status_code != 422:
            raise
        if room is None or stay is None:
            raise HTTPException(404, "Only an active stay can take charges.")
        error_msg = exc.detail
        return render(request, "partials/room_charge_modal.html", db, {
            "room": room, "stay": stay, "error": error_msg,
            "description_val": str(form.get("description", "")),
            "quantity_val": str(form.get("quantity", "1")),
            "rate_val": str(form.get("unit_price", "")),
            "selected_category": str(form.get("category", "service")),
        }, status_code=200)
    except (InvalidOperation, ValueError) as exc:
        db.rollback()
        if room is None or stay is None:
            raise HTTPException(404, "Only an active stay can take charges.")
        error_msg = str(exc)
        return render(request, "partials/room_charge_modal.html", db, {
            "room": room, "stay": stay, "error": error_msg,
            "description_val": str(form.get("description", "")),
            "quantity_val": str(form.get("quantity", "1")),
            "rate_val": str(form.get("unit_price", "")),
            "selected_category": str(form.get("category", "service")),
        }, status_code=200)


@app.get("/rooms/{room_id}/history", response_class=HTMLResponse)
def room_history_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    stays = db.scalars(
        select(Stay).options(selectinload(Stay.guest))
        .where(Stay.room_id == room_id)
        .order_by(Stay.check_in.desc(), Stay.id.desc())
    ).all()
    return render(request, "partials/room_history_modal.html", db, {"room": room, "stays": stays})


@app.get("/rooms/{room_id}/rate", response_class=HTMLResponse)
def room_rate_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change room rates.")
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    return render(request, "partials/room_rate_modal.html", db, {"room": room})


@app.post("/rooms/{room_id}/rate")
async def room_set_rate(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change room rates.")
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        room = db.get(Room, room_id)
        rate = Decimal(str(form.get("nightly_rate", "")))
        if room is None:
            raise HTTPException(404, "Room not found.")
        if not rate.is_finite() or rate < 0 or rate > Decimal("9999999999.99"):
            raise ValueError("Room rate must be a valid non-negative amount.")
        if str(form.get("use_type_rate", "")).lower() in {"1", "on", "true"}:
            if room.room_type_item is None:
                raise ValueError("This room has no linked room type rate.")
            room.nightly_rate = room.room_type_item.rate
            room.nightly_rate_override = False
        else:
            room.nightly_rate = rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            room.nightly_rate_override = (
                room.room_type_item is None or room.nightly_rate != room.room_type_item.rate
            )
        rate_source = "custom override" if room.nightly_rate_override else "room type"
        _audit(
            db, user, "room.rate", "room", room.number,
            f"nightly rate set to {room.nightly_rate:.2f} ({rate_source})",
        )
        db.commit()
        return _room_action_response(f'Rate updated for room {room.number}.')
    except HTTPException:
        db.rollback()
        raise
    except (InvalidOperation, ValueError) as exc:
        db.rollback()
        room = db.get(Room, room_id)
        if room is None:
            raise HTTPException(404, "Room not found.")
        return render(request, "partials/room_rate_modal.html", db, {
            "room": room, "error": exc.detail if isinstance(exc, HTTPException) else str(exc),
            "rate_val": str(form.get("nightly_rate", "")),
            "use_type_rate_val": str(form.get("use_type_rate", "")).lower()
            in {"1", "on", "true"},
        }, status_code=200)


@app.post("/rooms/{room_id}/maintenance")
def room_set_maintenance(
    request: Request, room_id: int, state: str = Form("maintenance"),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change maintenance status.")
    check_csrf(request, csrf_token)
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    if state not in {"maintenance", "clean"}:
        raise HTTPException(422, "Choose a valid maintenance state.")
    if room.status == "occupied" or db.scalar(
        select(Stay.id).where(Stay.room_id == room_id, Stay.status == "checked_in")
    ) is not None:
        raise HTTPException(409, "An occupied room cannot be marked as maintenance.")
    if state == "maintenance" and db.scalar(select(Reservation.id).where(
        Reservation.room_id == room_id, Reservation.status == "confirmed",
        Reservation.check_out > datetime.now(timezone.utc).replace(tzinfo=None),
    )) is not None:
        raise HTTPException(409, "A room with a confirmed reservation cannot be marked as maintenance.")
    room.housekeeping_status = state
    room.status = "maintenance" if state == "maintenance" else "available"
    _audit(db, user, "room.maintenance", "room", room.number, f"state set to {state}")
    db.commit()
    return _room_action_response(f'Room {room.number} marked {state}.')


@app.post("/rooms/{room_id}/clean")
def room_mark_clean(request: Request, room_id: int, csrf_token: str = Form(""), db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    if db.scalar(select(Stay.id).where(Stay.room_id == room_id, Stay.status == "checked_in")):
        raise HTTPException(409, "An occupied room cannot be marked as clean.")
    if room.status == "dirty" or room.housekeeping_status in {"dirty", "needs_cleaning"}:
        room.status = "available"
        room.housekeeping_status = "clean"
        _audit(db, user, "room.housekeeping", "room", room.number, "housekeeping set to clean")
        db.commit()
        return _room_action_response(f'Room {room.number} marked clean.')
    return _room_action_response("Room already clean.", "info")


@app.get("/rooms/inventory", response_class=HTMLResponse)
def room_inventory_page(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can manage room inventory.")
    rooms = db.scalars(select(Room).order_by(Room.active.desc(), Room.number)).all()
    room_types = db.scalars(
        select(CatalogItem).where(
            CatalogItem.category == "room", CatalogItem.active.is_(True),
        ).order_by(CatalogItem.sort_order, CatalogItem.name)
    ).all()
    return render(request, "room_inventory.html", db, {
        "rooms": rooms, "room_types": room_types,
    })


@app.post("/rooms/inventory")
async def room_inventory_create(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can manage room inventory.")
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        number = str(form.get("number", "")).strip()
        room_type_id = str(form.get("room_type_id", "")).strip()
        floor = str(form.get("floor", "1")).strip()
        if not number or len(number) > 20:
            raise ValueError("Enter a room number within the allowed length.")
        if not floor or len(floor) > 30:
            raise ValueError("Enter a valid floor.")
        room_type = db.get(CatalogItem, int(room_type_id)) if room_type_id.isdigit() else None
        if room_type is None or room_type.category != "room" or not room_type.active:
            raise ValueError("Select an active room type from the catalog.")
        rate_input = str(form.get("nightly_rate", "")).strip()
        rate = Decimal(rate_input) if rate_input else room_type.rate
        if not rate.is_finite() or rate < 0 or rate > Decimal("9999999999.99"):
            raise ValueError("Nightly rate must be a valid non-negative amount.")
        rate = rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if db.scalar(select(Room.id).where(Room.number == number)) is not None:
            raise ValueError("That room number already exists.")
        room = Room(
            number=number, type=room_type.name, floor=floor, status="available",
            housekeeping_status="clean",
            nightly_rate=rate, nightly_rate_override=rate != room_type.rate,
            room_type_item=room_type, active=True,
        )
        db.add(room)
        db.flush()
        _audit(db, user, "room.create", "room", number, f"{room_type.name}, floor {floor}")
        db.commit()
        return RedirectResponse(toast_url("/rooms/inventory", f"Room {number} added to the inventory.", "success"), status_code=303)
    except (InvalidOperation, ValueError, HTTPException) as exc:
        db.rollback()
        if isinstance(exc, HTTPException):
            raise
        rooms = db.scalars(select(Room).order_by(Room.active.desc(), Room.number)).all()
        room_types = db.scalars(
            select(CatalogItem).where(
                CatalogItem.category == "room", CatalogItem.active.is_(True),
            ).order_by(CatalogItem.sort_order, CatalogItem.name)
        ).all()
        return render(request, "room_inventory.html", db, {
            "rooms": rooms, "room_types": room_types,
            "error": str(exc),
            "form_values": {
                "number": str(form.get("number", "")),
                "room_type_id": room_type_id,
                "floor": str(form.get("floor", "1")),
                "nightly_rate": str(form.get("nightly_rate", "")),
            },
        }, status_code=200)


@app.get("/rooms/{room_id}/details", response_class=HTMLResponse)
def room_details_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can edit room details.")
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    room_types = db.scalars(
        select(CatalogItem).where(
            CatalogItem.category == "room",
            (CatalogItem.active.is_(True)) | (CatalogItem.id == room.room_type_item_id),
        ).order_by(CatalogItem.sort_order, CatalogItem.name)
    ).all()
    return render(request, "partials/room_details_modal.html", db, {
        "room": room, "room_types": room_types,
    })


@app.post("/rooms/{room_id}/details")
async def room_details_update(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can edit room details.")
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        room = db.get(Room, room_id)
        if room is None:
            raise HTTPException(404, "Room not found.")
        room_type_id = str(form.get("room_type_id", "")).strip()
        room_type = db.get(CatalogItem, int(room_type_id)) if room_type_id.isdigit() else None
        if room_type is None or room_type.category != "room" or not room_type.active:
            raise ValueError("Select an active room type from the catalog.")
        floor = str(form.get("floor", "")).strip()
        if not floor or len(floor) > 30:
            raise ValueError("Enter a floor within 30 characters.")
        room.type = room_type.name
        room.room_type_item = room_type
        room.floor = floor
        if not room.nightly_rate_override:
            room.nightly_rate = room_type.rate
        _audit(
            db, user, "room.update", "room", room.number,
            f"room type {room_type.name}, floor {floor}",
        )
        db.commit()
        return Response(status_code=204, headers={"HX-Redirect": "/rooms/inventory"})
    except HTTPException:
        db.rollback()
        raise
    except ValueError as exc:
        db.rollback()
        room = db.get(Room, room_id)
        if room is None:
            raise HTTPException(404, "Room not found.")
        room_types = db.scalars(
            select(CatalogItem).where(
                CatalogItem.category == "room",
                CatalogItem.active.is_(True),
            ).order_by(CatalogItem.sort_order, CatalogItem.name)
        ).all()
        return render(request, "partials/room_details_modal.html", db, {
            "room": room, "room_types": room_types, "error": str(exc),
            "room_type_id": str(form.get("room_type_id", "")),
            "floor_val": str(form.get("floor", "")),
        }, status_code=200)


@app.post("/rooms/{room_id}/archive")
def room_inventory_archive(
    request: Request, room_id: int, csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can manage room inventory.")
    check_csrf(request, csrf_token)
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    if db.scalar(select(Stay.id).where(Stay.room_id == room_id, Stay.status == "checked_in")):
        raise HTTPException(409, "A room with an active stay cannot be archived.")
    if db.scalar(select(Reservation.id).where(
        Reservation.room_id == room_id, Reservation.status == "confirmed",
        Reservation.check_out > datetime.now(timezone.utc).replace(tzinfo=None),
    )):
        raise HTTPException(409, "A room with a future reservation cannot be archived.")
    room.active = False
    _audit(db, user, "room.archive", "room", room.number, "room removed from active inventory")
    db.commit()
    return RedirectResponse(toast_url("/rooms/inventory", f"Room {room.number} archived.", "warning"), status_code=303)


@app.post("/rooms/{room_id}/restore")
def room_inventory_restore(
    request: Request, room_id: int, csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can manage room inventory.")
    check_csrf(request, csrf_token)
    room = db.get(Room, room_id)
    if room is None:
        raise HTTPException(404, "Room not found.")
    if room.room_type_item is not None and not room.room_type_item.active:
        raise HTTPException(409, "Reactivate this room's room type before restoring the room.")
    room.active = True
    _audit(db, user, "room.restore", "room", room.number, "room returned to active inventory")
    db.commit()
    return RedirectResponse(toast_url("/rooms/inventory", f"Room {room.number} restored.", "success"), status_code=303)


@app.get("/rooms/{room_id}/reserve", response_class=HTMLResponse)
def room_reserve_modal(request: Request, room_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    room = db.get(Room, room_id)
    if (
        not room or not room.active
        or room.housekeeping_status in {"dirty", "needs_cleaning", "maintenance"}
        or room.status in {"dirty", "maintenance"}
    ):
        raise HTTPException(409, "This room is inactive, out of service, or awaiting cleaning.")
    return render(request, "partials/room_reserve_modal.html", db, {
        "room": room, "guests": db.scalars(
            select(Guest).where(Guest.active.is_(True)).order_by(Guest.name)
        ).all(),
        "checkin_local": (local_now() + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M"),
        "checkout_local": (local_now() + timedelta(days=2)).strftime("%Y-%m-%dT%H:%M"),
    })


@app.post("/rooms/{room_id}/reserve")
async def room_reserve(request: Request, room_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    room = db.get(Room, room_id)
    if not room:
        raise HTTPException(404, "Room not found.")
    guests = db.scalars(
        select(Guest).where(Guest.active.is_(True)).order_by(Guest.name)
    ).all()

    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        if (
            not room.active
            or room.status in {"dirty", "maintenance"}
            or room.housekeeping_status in {"dirty", "needs_cleaning", "maintenance"}
        ):
            raise ValueError("This room is inactive, out of service, or awaiting cleaning.")
        guest_id = str(form.get("guest_id", "")).strip()
        guest = _guest_from_form(db, guest_id)
        if guest is not None and not guest.active:
            raise ValueError("Select an active guest.")
        name = str(form.get("guest_name", "")).strip()
        phone = str(form.get("phone", "")).strip()
        email = str(form.get("email", "")).strip()
        address = str(form.get("address", "")).strip()
        if len(name) > 160 or len(phone) > 40 or len(email) > 180 or len(address) > 240:
            raise ValueError("Guest name, phone, email, or address exceeds the allowed length.")
        if guest is None:
            if not name:
                raise ValueError("Enter a guest name to reserve this room.")
            guest = Guest(
                name=name, phone=phone, email=email, address=address,
                created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            )
            db.add(guest)
            db.flush()
        elif email:
            guest.email = email
        check_in = _parse_local_datetime(str(form.get("check_in", "")), "check-in")
        check_out = _parse_local_datetime(str(form.get("check_out", "")), "check-out")
        if check_out <= check_in:
            raise ValueError("Expected check-out must be after check-in.")
        active_checkout = func.coalesce(Stay.expected_checkout, Stay.check_out)
        active_stay = db.scalar(select(Stay.id).where(
            Stay.room_id == room.id, Stay.status == "checked_in",
            Stay.check_in < check_out,
            or_(active_checkout.is_(None), active_checkout > check_in),
        ).limit(1))
        if active_stay is not None:
            raise ValueError("The requested dates overlap an active stay.")
        overlap = db.scalar(select(Reservation.id).where(
            Reservation.room_id == room.id,
            Reservation.status.in_(("confirmed", "checked_in")),
            Reservation.check_in < check_out,
            Reservation.check_out > check_in,
        ).limit(1))
        if overlap is not None:
            raise ValueError("The requested dates overlap another reservation.")
        try:
            rate = Decimal(str(form.get("nightly_rate", "")).strip() or str(room.nightly_rate))
            deposit = Decimal(str(form.get("deposit", "0")).strip() or "0")
        except InvalidOperation as exc:
            raise ValueError("Nightly rate and deposit must be valid amounts.") from exc
        max_amount = Decimal("9999999999.99")
        if (
            not rate.is_finite() or rate < 0 or rate > max_amount
            or not deposit.is_finite() or deposit < 0 or deposit > max_amount
        ):
            raise ValueError("Nightly rate and deposit must be non-negative valid amounts.")
        notes = str(form.get("notes", "")).strip()
        if len(notes) > 1000:
            raise ValueError("Reservation notes must be 1,000 characters or fewer.")
        rate = rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        deposit = deposit.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        reservation = Reservation(
            room_id=room.id,
            guest_id=guest.id,
            guest_name=name or guest.name,
            guest_phone=phone or guest.phone,
            guest_email=email or guest.email,
            guest_address=address or guest.address,
            guest_pan=guest.pan_vat_number,
            check_in=check_in,
            check_out=check_out,
            nightly_rate=rate,
            deposit=deposit,
            source="front desk",
            notes=notes,
            status="confirmed",
            business_date=_business_date(local_now().astimezone(timezone.utc).replace(tzinfo=None)).isoformat(),
            created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            user_id=user.id,
        )
        db.add(reservation)
        room.status = "reserved"
        _audit(
            db, user, "reservation.create", "reservation", room.number,
            f"{guest.name} · {check_in.isoformat()} to {check_out.isoformat()}",
        )
        db.commit()
        return _room_action_response(f'Reservation created for room {room.number}.')
    except (HTTPException, ValueError, InvalidOperation) as exc:
        db.rollback()
        error_msg = exc.detail if isinstance(exc, HTTPException) else str(exc)
        return render(request, "partials/room_reserve_modal.html", db, {
            "room": room, "guests": guests,
            "checkin_local": str(form.get("check_in", "")) or (local_now() + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M"),
            "checkout_local": str(form.get("check_out", "")) or (local_now() + timedelta(days=2)).strftime("%Y-%m-%dT%H:%M"),
            "error": error_msg,
            "guest_name_val": str(form.get("guest_name", "")),
            "phone_val": str(form.get("phone", "")),
            "email_val": str(form.get("email", "")),
            "address_val": str(form.get("address", "")),
            "selected_guest_id": str(form.get("guest_id", "")),
            "rate_val": str(form.get("nightly_rate", "")),
            "deposit_val": str(form.get("deposit", "")),
            "notes_val": str(form.get("notes", "")),
        }, status_code=200)


@app.get("/reservations", response_class=HTMLResponse)
def reservations_page(
    request: Request,
    status: str = "confirmed",
    room: str = "",
    check_in: str = "",
    check_out: str = "",
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    statuses = {"all", "confirmed", "checked_in", "checked_out", "cancelled"}
    if status not in statuses:
        raise HTTPException(422, "Select a valid reservation status.")
    try:
        start_date = date.fromisoformat(check_in) if check_in else None
        end_date = date.fromisoformat(check_out) if check_out else None
    except ValueError as exc:
        raise HTTPException(422, "Enter valid reservation filter dates.") from exc
    if start_date and end_date and end_date < start_date:
        raise HTTPException(422, "The end date must be on or after the start date.")
    room_filter = room.strip()
    if len(room_filter) > 20:
        raise HTTPException(422, "Room search must be 20 characters or fewer.")
    query = select(Reservation).options(
        selectinload(Reservation.room), selectinload(Reservation.guest),
    ).order_by(Reservation.check_in.desc(), Reservation.id.desc())
    if status != "all":
        query = query.where(Reservation.status == status)
    if room_filter:
        query = query.join(Room).where(Room.number.ilike(f"%{room_filter}%"))
    if start_date:
        query = query.where(
            Reservation.check_out >= datetime.combine(start_date, time.min)
        )
    if end_date:
        query = query.where(
            Reservation.check_in < datetime.combine(end_date + timedelta(days=1), time.min)
        )
    return render(request, "reservations.html", db, {
        "reservations": db.scalars(query).all(),
        "status_filter": status,
        "room_filter": room_filter,
        "check_in_filter": check_in,
        "check_out_filter": check_out,
        "cancelled": request.query_params.get("cancelled") == "1",
    })


@app.post("/reservations/{reservation_id}/cancel")
async def reservation_cancel(
    request: Request, reservation_id: int, db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    reservation = db.get(Reservation, reservation_id)
    if reservation is None:
        raise HTTPException(404, "Reservation not found.")
    if reservation.status != "confirmed":
        raise HTTPException(409, "Only confirmed reservations can be cancelled.")
    reason = str(form.get("cancel_reason", "")).strip()
    if len(reason) > 1000:
        raise HTTPException(422, "Cancellation reason must be 1,000 characters or fewer.")
    reservation.status = "cancelled"
    reservation.cancel_reason = reason
    room = db.get(Room, reservation.room_id) if reservation.room_id is not None else None
    if room is not None and room.status == "reserved":
        remaining_reservation = db.scalar(select(Reservation.id).where(
            Reservation.room_id == room.id,
            Reservation.status == "confirmed",
            Reservation.id != reservation.id,
        ).limit(1))
        if remaining_reservation is None:
            room.status = "available"
    _audit(
        db, user, "reservation.cancel", "reservation", str(reservation.id),
        f"{reservation.guest_name} · deposit {money(reservation.deposit)} · "
        f"{reason or 'no reason given'}",
    )
    db.commit()
    return RedirectResponse(toast_url("/reservations?status=all&cancelled=1", "Reservation cancelled.", "warning"), status_code=303)


def _orders_context(
    request: Request, db: Session, status: str = "open", kind: str = "all",
    q: str = "", sim_phone: str = "",
) -> dict:
    status = status if status in {"open", "all", *ORDER_STATUSES} else "open"
    kind = kind if kind in {"all", *KIND_LABELS} else "all"
    q = q.strip()[:80]
    query = select(Order).options(selectinload(Order.items)).order_by(
        Order.created_at.desc(), Order.id.desc(),
    )
    if status == "open":
        query = query.where(Order.status.in_(OPEN_STATUSES))
    elif status != "all":
        query = query.where(Order.status == status)
    if kind != "all":
        query = query.where(Order.kind == kind)
    if q:
        term = f"%{q}%"
        query = query.where(or_(
            Order.order_number.ilike(term), Order.customer_name.ilike(term),
            Order.customer_phone.ilike(term), Order.room_number.ilike(term),
        ))
    orders = db.scalars(query.limit(200)).all()

    status_counts = dict(db.execute(
        select(Order.status, func.count(Order.id)).group_by(Order.status)
    ).all())
    start, end = business_window(local_now().date())
    completed_today = db.scalar(select(func.count(Order.id)).where(
        Order.status == "completed", Order.closed_at >= start, Order.closed_at < end,
    )) or 0
    metrics = {
        "new": int(status_counts.get("new", 0)),
        "active": sum(int(status_counts.get(name, 0)) for name in ("confirmed", "preparing", "ready")),
        "ready": int(status_counts.get("ready", 0)),
        "completed_today": int(completed_today),
    }

    sim_phone = whatsapp.normalize_phone(sim_phone) or whatsapp.DEFAULT_PHONE
    sim_messages = list(db.scalars(
        select(WhatsAppMessage).where(WhatsAppMessage.wa_id == sim_phone)
        .order_by(WhatsAppMessage.id.desc()).limit(60)
    ).all())[::-1]
    sim_conversation = db.scalar(
        select(WhatsAppConversation).where(WhatsAppConversation.wa_id == sim_phone)
    )
    wa_settings = whatsapp.config()
    webhook_url = f"{str(request.base_url).rstrip('/')}/webhooks/whatsapp"
    return {
        "orders": orders,
        "metrics": metrics,
        "status_counts": status_counts,
        "status_filter": status,
        "kind_filter": kind,
        "q": q,
        "sim_phone": sim_phone,
        "sim_messages": sim_messages,
        "sim_conversation": sim_conversation,
        "wa_connected": whatsapp.cloud_configured(),
        "wa_enabled": bool(wa_settings["enabled"]),
        "wa_verify_set": bool(wa_settings["verify_token"]),
        "wa_secret_set": bool(wa_settings["app_secret"]),
        "webhook_url": webhook_url,
        "status_labels": STATUS_LABELS,
        "status_badges": STATUS_BADGES,
        "kind_labels": KIND_LABELS,
        "fulfilment_labels": FULFILMENT_LABELS,
        "open_statuses": OPEN_STATUSES,
        "item_summary": item_summary,
    }


@app.get("/orders", response_class=HTMLResponse)
def orders_page(
    request: Request, status: str = "open", kind: str = "all", q: str = "",
    sim_phone: str = "", db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    return render(request, "coming_soon.html", db)


@app.get("/orders/{order_id}", response_class=HTMLResponse)
def order_detail(request: Request, order_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    order = db.scalar(
        select(Order).options(selectinload(Order.items)).where(Order.id == order_id)
    )
    if order is None:
        raise HTTPException(404, "Order not found.")
    conversation = db.scalar(
        select(WhatsAppConversation).where(WhatsAppConversation.wa_id == order.contact_wa_id)
    ) if order.contact_wa_id else None
    messages = list(db.scalars(
        select(WhatsAppMessage).where(WhatsAppMessage.wa_id == order.contact_wa_id)
        .order_by(WhatsAppMessage.id.desc()).limit(80)
    ).all())[::-1] if order.contact_wa_id else []
    errors = {
        "status": "That status change is not allowed from the order's current state.",
        "phone": "This order has no WhatsApp number to reply to.",
    }
    return render(request, "order_detail.html", db, {
        "order": order,
        "conversation": conversation,
        "messages": messages,
        "available_statuses": [
            name for name in ORDER_STATUSES if can_transition(order.status, name)
        ],
        "primary_status": next_status(order.status),
        "status_labels": STATUS_LABELS,
        "status_badges": STATUS_BADGES,
        "kind_labels": KIND_LABELS,
        "fulfilment_labels": FULFILMENT_LABELS,
        "item_summary": item_summary,
        "notice": {
            "1": "sent", "0": "failed",
        }.get(request.query_params.get("notice", ""), ""),
        "error": errors.get(request.query_params.get("error", ""), ""),
        "reply_state": request.query_params.get("reply", ""),
        "wa_connected": whatsapp.cloud_configured(),
    })


@app.post("/orders/{order_id}/status")
def order_set_status(
    request: Request, order_id: int, status: str = Form(""),
    note: str = Form(""), csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    order = db.get(Order, order_id)
    if order is None:
        raise HTTPException(404, "Order not found.")
    if not can_transition(order.status, status):
        return RedirectResponse(toast_url(
            f"/orders/{order_id}?error=status",
            "That status change is not allowed from the order's current state.",
            "error",
        ), status_code=303)
    note = note.strip()
    if len(note) > 500:
        raise HTTPException(422, "The staff note must be 500 characters or fewer.")
    previous = order.status
    order.status = status
    order.updated_at = models.utcnow()
    order.user_id = user.id
    if status in {"completed", "cancelled"}:
        order.closed_at = models.utcnow()
    if note:
        order.notes = f"{order.notes}\nStaff: {note}".strip()[-2000:]
    _audit(db, user, "order.status", "order", order.order_number,
           f"{previous} → {status}" + (f" · {note}" if note else ""))
    db.commit()
    sent = whatsapp.notify_order_status(db, order, status)
    message = f"Order {order.order_number} marked {STATUS_LABELS.get(status, status).lower()}."
    if not sent:
        message += " The guest could not be messaged on WhatsApp."
    return RedirectResponse(toast_url(
        f"/orders/{order_id}?notice={1 if sent else 0}", message,
        "success" if sent else "warning",
    ), status_code=303)


@app.post("/orders/{order_id}/reply")
def order_reply(
    request: Request, order_id: int, body: str = Form(""),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    order = db.get(Order, order_id)
    if order is None:
        raise HTTPException(404, "Order not found.")
    conversation = whatsapp.conversation_for(db, order.contact_wa_id, order.customer_name)
    if conversation is None:
        return RedirectResponse(toast_url(
            f"/orders/{order_id}?error=phone",
            "This order has no WhatsApp number to reply to.", "error",
        ), status_code=303)
    ok, _detail = whatsapp.send_manual(db, conversation, body)
    _audit(db, user, "order.reply", "order", order.order_number,
           "manual WhatsApp reply" if ok else "manual WhatsApp reply failed")
    db.commit()
    return RedirectResponse(toast_url(
        f"/orders/{order_id}?reply={'sent' if ok else 'failed'}",
        "Message sent to the guest." if ok else "The message could not be sent. Check the WhatsApp connection.",
        "success" if ok else "error",
    ), status_code=303)


@app.post("/orders/conversations/{conversation_id}/bot")
def order_bot_control(
    request: Request, conversation_id: int, action: str = Form("resume"),
    next_url: str = Form("/orders"), csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    conversation = db.get(WhatsAppConversation, conversation_id)
    if conversation is None:
        raise HTTPException(404, "Conversation not found.")
    paused = action != "resume"
    whatsapp.set_bot_paused(db, conversation, paused)
    _audit(db, user, "order.bot", "conversation", str(conversation.id),
           "assistant paused" if paused else "assistant resumed")
    db.commit()
    destination = next_url if next_url.startswith("/orders") and "//" not in next_url else "/orders"
    return RedirectResponse(toast_url(
        destination,
        "Assistant resumed for this chat." if not paused else "Assistant paused for this chat.",
        "success" if not paused else "info",
    ), status_code=303)


@app.post("/orders/simulator", response_class=HTMLResponse)
async def whatsapp_simulator(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    phone = whatsapp.normalize_phone(str(form.get("phone", ""))) or whatsapp.DEFAULT_PHONE
    error = ""
    toast = None
    action = str(form.get("action", "send"))
    if action == "resume":
        conversation = whatsapp.conversation_for(db, phone, "Simulator guest")
        if conversation is not None:
            whatsapp.set_bot_paused(db, conversation, False)
            toast = toast_event("Assistant resumed for this chat.", "success")
    else:
        message = str(form.get("message", "")).strip()
        if not message:
            error = "Type a message to send."
            toast = toast_event(error, "error")
        else:
            try:
                result = whatsapp.process_inbound(
                    db, phone, "Simulator guest", message, transport="simulator",
                )
            except Exception:
                db.rollback()
                _log.exception("Simulator message failed for %s", phone)
                error = "The message could not be processed. Please try again."
                toast = toast_event(error, "error")
            else:
                if result["order"] is not None:
                    _audit(db, user, "order.create", "order",
                           result["order"].order_number, "created in the WhatsApp simulator")
                    db.commit()
                    toast = toast_event(
                        f"Order {result['order'].order_number} created from the conversation.",
                        "success",
                    )
                else:
                    toast = toast_event("Message sent to the assistant.", "success")
    messages = list(db.scalars(
        select(WhatsAppMessage).where(WhatsAppMessage.wa_id == phone)
        .order_by(WhatsAppMessage.id.desc()).limit(60)
    ).all())[::-1]
    conversation = db.scalar(
        select(WhatsAppConversation).where(WhatsAppConversation.wa_id == phone)
    )
    headers = {"HX-Trigger": json.dumps({"showToast": toast})} if toast else None
    return render(request, "partials/whatsapp_simulator.html", db, {
        "sim_phone": phone, "sim_messages": messages,
        "sim_conversation": conversation, "error": error,
    }, headers=headers)


@app.get("/webhooks/whatsapp", include_in_schema=False)
def whatsapp_webhook_verify(request: Request):
    params = request.query_params
    expected = str(whatsapp.config()["verify_token"])
    if (
        params.get("hub.mode") == "subscribe" and expected
        and secrets.compare_digest(params.get("hub.verify_token", ""), expected)
    ):
        return PlainTextResponse(params.get("hub.challenge", ""))
    raise HTTPException(403, "WhatsApp webhook verification failed.")


@app.post("/webhooks/whatsapp", include_in_schema=False)
async def whatsapp_webhook(request: Request, db: Session = Depends(get_db)):
    raw = await request.body()
    if not whatsapp.verify_signature(raw, request.headers.get("x-hub-signature-256")):
        raise HTTPException(403, "Invalid WhatsApp signature.")
    if not whatsapp.config()["enabled"]:
        return {"status": "ignored"}
    try:
        payload = json.loads(raw or b"{}")
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise HTTPException(400, "Malformed WhatsApp payload.") from exc
    if not isinstance(payload, dict):
        raise HTTPException(400, "Malformed WhatsApp payload.")
    for message in whatsapp.parse_webhook(payload):
        try:
            whatsapp.process_inbound(
                db, message["wa_id"], message["name"], message["text"],
                transport="cloud", message_id=message["message_id"],
            )
        except Exception:
            # One garbage inbound must never fail the whole webhook — Meta
            # interpret a non-200 as failure and retries forever.
            _log.exception("WhatsApp inbound message failed from %s", message.get("wa_id"))
            db.rollback()
    return {"status": "accepted"}


def _customer_page_data(db: Session, q: str = "") -> dict:
    invoice_stats = select(
        Invoice.guest_id.label("guest_id"),
        func.count(Invoice.id).label("invoice_count"),
        func.coalesce(func.sum(Invoice.total), 0).label("total_billed"),
    ).group_by(Invoice.guest_id).subquery()
    query = select(
        Guest,
        func.coalesce(invoice_stats.c.invoice_count, 0),
        func.coalesce(invoice_stats.c.total_billed, 0),
    ).outerjoin(invoice_stats, invoice_stats.c.guest_id == Guest.id).where(
        Guest.active.is_(True)
    ).order_by(Guest.name)
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.where(or_(
            Guest.name.ilike(term),
            Guest.phone.ilike(term),
            Guest.email.ilike(term),
            Guest.pan_vat_number.ilike(term),
        ))
    rows = [
        (guest, count, total) for guest, count, total in db.execute(query.limit(1000)).all()
        if not _is_demo_guest(guest)
    ][:300]
    return {
        "customers": [
            {"guest": guest, "invoice_count": int(count), "total_billed": total}
            for guest, count, total in rows
        ],
        "q": q,
    }


@app.get("/customers", response_class=HTMLResponse)
def customers_page(request: Request, q: str = "", db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    return render(request, "customers.html", db, _customer_page_data(db, q))


@app.post("/customers")
def create_customer(
    request: Request, name: str = Form(""), phone: str = Form(""),
    email: str = Form(""), id_type: str = Form(""), id_number: str = Form(""),
    address: str = Form(""), pan_vat_number: str = Form(""),
    customer_type: str = Form("individual"),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    form_values = {
        "name": name, "phone": phone, "email": email, "id_type": id_type,
        "id_number": id_number, "address": address, "pan_vat_number": pan_vat_number,
        "customer_type": customer_type,
    }
    try:
        check_csrf(request, csrf_token)
    except HTTPException:
        data = _customer_page_data(db)
        data.update({"error": "Form expired. Please try again.", "form_values": form_values})
        return render(request, "customers.html", db, {
            **data,
        }, status_code=403)
    name = name.strip()
    if not name:
        data = _customer_page_data(db)
        data.update({"error": "Guest name is required.", "form_values": form_values})
        return render(request, "customers.html", db, data, status_code=422)
    if customer_type not in {"individual", "business"}:
        data = _customer_page_data(db)
        data.update({"error": "Choose an individual or business guest type.", "form_values": form_values})
        return render(request, "customers.html", db, data, status_code=422)

    phone_value = phone.strip()[:40]
    guest = db.scalar(select(Guest).where(
        Guest.active.is_(True), Guest.name == name[:160], Guest.phone == phone_value,
    ))
    if guest is None:
        guest = Guest(name=name[:160], phone=phone_value,
                      created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None))
        db.add(guest)
    guest.email = email.strip()[:180] or guest.email
    guest.id_type = id_type.strip()[:40] or guest.id_type
    guest.id_number = id_number.strip()[:100] or guest.id_number
    guest.address = address.strip()[:240] or guest.address
    guest.pan_vat_number = pan_vat_number.strip()[:80] or guest.pan_vat_number
    guest.customer_type = customer_type
    db.flush()
    _audit(db, user, "customer.save", "guest", str(guest.id), "Guest profile saved.")
    db.commit()
    return RedirectResponse(toast_url("/customers?created=1", "Customer added.", "success"), status_code=303)


@app.get("/customers/{guest_id}/edit", response_class=HTMLResponse)
def edit_customer_page(request: Request, guest_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    guest = db.get(Guest, guest_id)
    if guest is None:
        raise HTTPException(404, "Guest not found.")
    return render(request, "customer_edit.html", db, {"guest": guest})


@app.post("/customers/{guest_id}/edit")
def edit_customer(
    request: Request, guest_id: int, name: str = Form(""), phone: str = Form(""),
    email: str = Form(""), id_type: str = Form(""), id_number: str = Form(""),
    address: str = Form(""), pan_vat_number: str = Form(""),
    customer_type: str = Form("individual"), csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    guest = db.get(Guest, guest_id)
    if guest is None:
        raise HTTPException(404, "Guest not found.")
    name = name.strip()
    if not name:
        raise HTTPException(422, "Guest name is required.")
    if customer_type not in {"individual", "business"}:
        raise HTTPException(422, "Choose an individual or business guest type.")
    guest.name = name[:160]
    guest.phone = phone.strip()[:40]
    guest.email = email.strip()[:180]
    guest.id_type = id_type.strip()[:40]
    guest.id_number = id_number.strip()[:100]
    guest.address = address.strip()[:240]
    guest.pan_vat_number = pan_vat_number.strip()[:80]
    guest.customer_type = customer_type
    _audit(db, user, "customer.update", "guest", str(guest.id), "Guest profile updated.")
    db.commit()
    return RedirectResponse(toast_url("/customers?updated=1", "Customer details updated.", "success"), status_code=303)


@app.post("/customers/{guest_id}/deactivate")
def deactivate_customer(
    request: Request, guest_id: int, csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can deactivate guests.")
    check_csrf(request, csrf_token)
    guest = db.get(Guest, guest_id)
    if guest is None:
        raise HTTPException(404, "Guest not found.")
    guest.active = False
    _audit(db, user, "customer.deactivate", "guest", str(guest.id), "Guest deactivated; invoice history retained.")
    db.commit()
    return RedirectResponse(toast_url("/customers?deactivated=1", "Customer deactivated.", "warning"), status_code=303)


@app.get("/invoices", response_class=HTMLResponse)
def invoices_page(request: Request, q: str = "", status: str = "all", room_id: int | None = None,
                  guest_id: int | None = None,
                  db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    query = select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room)).order_by(Invoice.issue_date.desc())
    selected_guest = None
    selected_room = None
    if guest_id is not None:
        selected_guest = db.get(Guest, guest_id)
        if selected_guest is None:
            raise HTTPException(404, "Guest not found.")
        query = query.where(Invoice.guest_id == guest_id)
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.outerjoin(Guest).outerjoin(Room).where(
            or_(Invoice.invoice_number.ilike(term), Guest.name.ilike(term), Room.number.ilike(term))
        )
    if status in {"issued", "paid", "void", "draft"}:
        query = query.where(Invoice.status == status)
    if status == "outstanding":
        query = query.where(Invoice.status == "issued", Invoice.balance > 0)
    if room_id is not None:
        selected_room = db.get(Room, room_id)
        if selected_room is None:
            raise HTTPException(404, "Room not found.")
        query = query.where(Invoice.room_id == room_id)
    invoices = [
        invoice for invoice in db.scalars(query.limit(1000)).all()
        if not _is_demo_invoice(invoice)
    ][:300]
    room_number_label = selected_room.number if selected_room else (
        db.get(Room, room_id).number if room_id else ""
    )
    return render(request, "invoices.html", db,
                   {"invoices": invoices, "q": q,
                    "status_filter": status, "room_id": room_id,
                    "guest_id": guest_id, "selected_guest": selected_guest,
                    "selected_room": selected_room,
                    "room_number_label": room_number_label})


def _billing_page_data(db: Session) -> dict:
    settings = _billing_settings(db)
    return {
        "rooms": [room for room in db.scalars(
            select(Room).where(Room.active.is_(True)).order_by(Room.number)
        ).all() if not _is_demo_room(room)],
        "printer_logo_src": _printer_logo_source(settings),
        "catalog": db.scalars(select(CatalogItem).where(CatalogItem.active.is_(True)).order_by(CatalogItem.category, CatalogItem.sort_order, CatalogItem.name)).all(),
        "guests": [guest for guest in db.scalars(
            select(Guest).order_by(Guest.name)
        ).all() if not _is_demo_guest(guest)],
        "today": local_now().date().isoformat(),
        "invoice_number": _next_billing_number(db),
        "billing_settings": settings,
        "invoice_language_default": (
            "Nepali" if settings.get("invoice_language", "en").lower() in {"ne", "nepali"}
            else "English"
        ),
        "invoice_copy_default": settings.get("invoice_copy", "Customer Copy"),
        "payment_terms_default": settings.get("payment_terms", "Payment due on presentation."),
        "payment_methods": _payment_methods(db),
        "occupied_stays": _occupied_stay_data(db),
        "draft_payload": None,
        "draft_kots": [],
        "prefill_invoice_number": None,
    }


def _occupied_stay_data(db: Session) -> dict[str, dict]:
    """Folio data for every room with an active stay, keyed by room id (str).

    Lets the billing workspace bill occupied rooms and preload a stay's
    charges (dishes, services, extra nights) until checkout.
    """
    data: dict[str, dict] = {}
    stays = db.scalars(select(Stay).where(Stay.status == "checked_in")).all()
    for stay in stays:
        room = db.get(Room, stay.room_id)
        if room is None:
            continue
        try:
            folio = _stay_folio(db, stay)
        except HTTPException:
            continue
        guest = db.get(Guest, stay.guest_id) if stay.guest_id else None
        items = []
        for index, line in enumerate(folio["lines"]):
            description = str(line.description)
            items.append({
                "key": f"stay_{stay.id}_{index}",
                "room_id": room.id if line.category == "room" else None,
                "description": description,
                "short": (f"{room.number} · {room.type}" if line.category == "room"
                          else description[:30]),
                "quantity": float(line.qty),
                "rate": float(line.unit_price),
                "discount": 0,
                "taxable": bool(line.taxable),
                "category": line.category,
            })
        data[str(room.id)] = {
            "guest_id": stay.guest_id if stay.guest_id else (guest.id if guest else None),
            "guest_name": (stay.guest_name or (guest.name if guest else "")) or "Walk-in guest",
            "guest_phone": (guest.phone if guest else "") or "",
            "guest_email": (guest.email if guest else "") or "",
            "guest_address": (guest.address if guest else "") or "",
            "guest_pan": (guest.pan_vat_number if guest else "") or "",
            "nights": int(folio["nights"]),
            "total": float(folio["totals"].total),
            "invoice_number": stay.invoice_number or "",
            "check_in": (local_datetime(stay.check_in, "%Y-%m-%dT%H:%M") or ""),
            "check_out": (local_datetime(stay.expected_checkout or stay.check_out, "%Y-%m-%dT%H:%M") or ""),
            "items": items,
        }
    return data


def _billing_room_prefill(db: Session, room_id: int) -> dict | None:
    """Initial billing-workspace payload when a room is opened for billing
    directly (e.g. double-clicked on the room board)."""
    room = db.get(Room, room_id)
    if room is None or not room.active:
        return None
    stay = db.scalar(
        select(Stay).where(Stay.room_id == room.id, Stay.status == "checked_in")
    )
    if stay is not None:
        folio = _stay_folio(db, stay)
        guest = db.get(Guest, stay.guest_id) if stay.guest_id else None
        items = []
        for index, line in enumerate(folio["lines"]):
            description = str(line.description)
            items.append({
                "key": f"stay_{stay.id}_{index}",
                "room_id": room.id if line.category == "room" else None,
                "description": description,
                "short": (f"{room.number} · {room.type}" if line.category == "room"
                          else description[:30]),
                "quantity": float(line.qty),
                "rate": float(line.unit_price),
                "discount": 0,
                "taxable": bool(line.taxable),
                "category": line.category,
            })
        return {
            "invoice_id": None,
            "guest_id": stay.guest_id if stay.guest_id else (guest.id if guest else None),
            "guest_name": (stay.guest_name or (guest.name if guest else "")) or "Walk-in guest",
            "guest_phone": (guest.phone if guest else "") or "",
            "guest_email": (guest.email if guest else "") or "",
            "guest_address": (guest.address if guest else "") or "",
            "guest_pan": (guest.pan_vat_number if guest else "") or "",
            "customer_type": "individual",
            "nights": int(folio["nights"]),
            "check_in": local_datetime(stay.check_in, "%Y-%m-%dT%H:%M"),
            "check_out": local_datetime(stay.expected_checkout or stay.check_out, "%Y-%m-%dT%H:%M"),
            "items": items,
        }
    return {
        "invoice_id": None,
        "guest_id": None,
        "guest_name": "",
        "guest_phone": "",
        "guest_email": "",
        "guest_address": "",
        "guest_pan": "",
        "customer_type": "individual",
        "nights": 1,
        "check_in": "",
        "check_out": "",
        "items": [{
            "key": f"room_{room.id}",
            "room_id": room.id,
            "description": f"Room {room.number} - {room.type}",
            "short": f"{room.number} · {room.type}",
            "quantity": 1,
            "rate": float(room.nightly_rate),
            "discount": 0,
            "taxable": True,
            "category": "room",
        }],
    }


def _billing_settings(db: Session) -> dict[str, str]:
    values = dict(DEFAULT_SETTINGS)
    values.update({
        row.key: row.value
        for row in db.scalars(select(HotelSetting)).all()
    })
    return values


def _payment_methods(db: Session) -> list[str]:
    configured = db.get(HotelSetting, "payment_methods")
    if configured is None:
        return list(PAYMENT_METHODS)
    methods = list(dict.fromkeys(
        value.strip()
        for value in configured.value.split(",")
        if value.strip()
    ))
    return methods or list(PAYMENT_METHODS)


def _normalize_payment_methods(value: str) -> str:
    if len(value) > 500:
        raise ValueError("Payment methods must be 500 characters or fewer.")
    methods = [method.strip() for method in value.split(",")]
    if not methods or any(not method for method in methods):
        raise ValueError("Enter payment methods separated by commas without empty entries.")
    if any(len(method) > 40 or method.casefold() == "all" for method in methods):
        raise ValueError("Each payment method must be 40 characters or fewer; 'All' is reserved.")
    normalized = [method.casefold() for method in methods]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Payment method names must be unique.")
    return ", ".join(methods)


def _printer_config(settings: dict[str, str]) -> dict[str, object]:
    paper = settings.get("printer_paper_size", "A4")
    if paper not in {"A4", "80", "58"}:
        paper = "A4"
    margins = {"narrow": 8, "normal": 12, "wide": 18}
    margin = margins.get(settings.get("printer_margins", "normal"), 12)
    try:
        copies = int(settings.get("printer_copies", "1"))
    except ValueError:
        copies = 1
    copies = max(1, min(5, copies))

    def enabled(key: str) -> bool:
        return settings.get(key, "1").strip().lower() not in {"0", "false", "no", "off"}

    workflow = settings.get("printer_default_workflow", "preview")
    if workflow not in {"preview", "direct"}:
        workflow = "preview"
    return {
        "paper": paper,
        "margin_mm": margin,
        "copies": copies,
        "workflow": workflow,
        "show_logo": enabled("printer_show_logo"),
        "show_header": enabled("printer_header"),
        "show_footer": enabled("printer_footer"),
        "show_contact": enabled("printer_contact_details"),
        "show_tax": enabled("printer_tax_details"),
    }


def _receipt_printer_config(settings: dict[str, str]) -> dict[str, object]:
    config = _printer_config(settings)
    if config["paper"] == "A4":
        config["paper"] = "80"
    config["margin_mm"] = 4
    config["copies"] = 1
    return config


def _printer_logo_source(settings: dict[str, str], *, for_pdf: bool = False) -> str:
    logo_path = settings.get("logo_path", "").strip() or "/static/aratilogo.webp"
    if not logo_path.startswith("/static/") or ".." in logo_path.split("/"):
        return ""
    static_root = (ROOT / "static").resolve()
    logo_file = (static_root / logo_path.removeprefix("/static/")).resolve()
    if not logo_file.is_relative_to(static_root) or not logo_file.is_file():
        return ""
    return logo_file.as_uri() if for_pdf else logo_path


def _next_billing_number(db: Session) -> str:
    bill_prefix = _billing_settings(db).get("bill_prefix", "ARH").strip().rstrip("-") or "ARH"
    prefix = f"{bill_prefix}-{local_now():%Y%m%d}-"
    count = db.scalar(select(func.count(Invoice.id)).where(Invoice.invoice_number.like(f"{prefix}%"))) or 0
    while True:
        count += 1
        candidate = f"{prefix}{count:04d}"
        if not db.scalar(select(Invoice.id).where(Invoice.invoice_number == candidate)):
            return candidate


def _next_kot_number(db: Session) -> str:
    prefix = f"KOT-{local_now():%Y%m%d}-"
    count = db.scalar(select(func.count(Kot.id)).where(Kot.kot_number.like(f"{prefix}%"))) or 0
    while True:
        count += 1
        candidate = f"{prefix}{count:04d}"
        if not db.scalar(select(Kot.id).where(Kot.kot_number == candidate)):
            return candidate


KOT_STATUSES = ("new", "preparing", "ready", "served")
KOT_STATUS_LABELS = {
    "new": "New",
    "preparing": "Preparing",
    "ready": "Ready",
    "served": "Served",
    "saved": "New",
    "open": "New",
    "cancelled": "Cancelled",
}
KOT_OPEN_STATUSES = ("new", "preparing", "ready", "saved", "open")
_KOT_NEXT_STATUS = {"new": "preparing", "saved": "preparing", "open": "preparing",
                    "preparing": "ready", "ready": "served", "served": None}


def _kot_status(value: str | None) -> str:
    """Map legacy KOT states onto the kitchen workflow."""
    return "new" if not value or value in {"saved", "open"} else value


def _kot_status_label(value: str | None) -> str:
    return KOT_STATUS_LABELS.get(value or "", (value or "New").replace("_", " ").title())


def _kot_next_status(value: str | None) -> str | None:
    return _KOT_NEXT_STATUS.get(_kot_status(value))


def _kot_can_transition(current: str | None, target: str) -> bool:
    if target not in KOT_STATUSES:
        return False
    return target == _KOT_NEXT_STATUS.get(_kot_status(current))


def _kot_items(kot: Kot) -> list[dict]:
    try:
        data = json.loads(kot.items or "[]")
    except (TypeError, ValueError):
        return []
    return [row for row in data if isinstance(row, dict)] if isinstance(data, list) else []


def _staff_name(user: User | None) -> str:
    return (user.full_name or user.username).strip() if user else ""


def _kot_issuer_name(kot: Kot) -> str:
    return kot.issued_by_name or _staff_name(kot.issued_by) or "Unknown staff"


def _kot_cut_name(kot: Kot) -> str:
    return _staff_name(kot.cut_by)


def _kot_dict(kot: Kot) -> dict:
    next_status = _kot_next_status(kot.status)
    return {
        "id": kot.id,
        "kot_number": kot.kot_number,
        "invoice_id": kot.invoice_id,
        "room_id": kot.room_id,
        "room_number": kot.room_number,
        "table_number": kot.table_number,
        "guest_id": kot.guest_id,
        "guest_name": kot.guest_name or (kot.guest.full_name or kot.guest.name if kot.guest else "") or "Walk-in guest",
        "issued_by_name": _kot_issuer_name(kot),
        "cut_by_name": _kot_cut_name(kot),
        "cut_at": local_datetime(kot.cut_at, "%Y-%m-%d %H:%M") if kot.cut_at else "",
        "items": _kot_items(kot),
        "total": float(kot.total or 0),
        "status": _kot_status(kot.status),
        "status_label": _kot_status_label(kot.status),
        "next_status": next_status,
        "next_status_label": _kot_status_label(next_status) if next_status else "",
        "notes": kot.notes or "",
        "date": local_datetime(kot.created_at, "%Y-%m-%d %H:%M") if kot.created_at else "",
        "business_date": kot.business_date,
    }


def _clean_kot_items(raw) -> list[dict]:
    if not isinstance(raw, list) or not raw:
        raise ValueError("Add at least one restaurant item to this KOT.")
    if len(raw) > 200:
        raise ValueError("A KOT can hold at most 200 items.")
    cleaned = []
    for index, row in enumerate(raw, start=1):
        if not isinstance(row, dict):
            raise ValueError("Invalid KOT item.")
        particular = str(row.get("particular") or row.get("description") or "").strip()[:160]
        if not particular:
            raise ValueError("Each KOT item needs a particular.")
        try:
            qty = Decimal(str(row.get("qty", row.get("quantity", 1))))
            rate = Decimal(str(row.get("rate", 0)))
        except (InvalidOperation, ValueError, TypeError) as exc:
            raise ValueError("KOT quantities and rates must be valid numbers.") from exc
        _finite(qty, "KOT quantity")
        _finite(rate, "KOT rate")
        if qty <= 0 or qty > MAX_QUANTITY_INPUT:
            raise ValueError(f"KOT quantity must be between 0 and {MAX_QUANTITY_INPUT}.")
        if rate < 0 or rate > MAX_MONEY_INPUT:
            raise ValueError(f"KOT rate must be between 0 and {MAX_MONEY_INPUT}.")
        amount = (qty * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        catalog_id = None
        raw_catalog_id = str(row.get("catalog_id") or "").strip()
        if raw_catalog_id.isdigit():
            catalog_id = int(raw_catalog_id)
        category = str(row.get("category") or "food").strip().lower()[:30] or "food"
        cleaned.append({
            "sn": index,
            "particular": particular,
            "qty": float(qty),
            "rate": float(rate),
            "amount": float(amount),
            "key": str(row.get("key", ""))[:80],
            "catalog_id": catalog_id,
            "category": category,
            "menu_group": str(row.get("menu_group") or "").strip()[:80],
            "unit": str(row.get("unit") or "pcs").strip()[:30] or "pcs",
            "taxable": bool(row.get("taxable", True)),
        })
    return cleaned


def _kot_guest(db: Session, guest_id) -> Guest | None:
    raw = str(guest_id or "").strip()
    if not raw.isdigit():
        return None
    return db.get(Guest, int(raw))


@app.post("/billing/create/walkin")
async def billing_create_walkin(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
    except HTTPException:
        pass

    guest = guest_service.get_or_create_guest(
        db,
        name=form.get("guest_name", "") or form.get("name", ""),
        phone=form.get("guest_phone", "") or form.get("phone", ""),
        email=form.get("guest_email", "") or form.get("email", ""),
        address=form.get("guest_address", "") or form.get("address", ""),
        pan_vat=form.get("guest_pan", "") or form.get("pan_vat", ""),
        customer_type=form.get("customer_type", "individual") or "individual",
    )

    settings = _billing_settings(db)
    draft = invoice_service.get_or_create_draft_invoice(
        db,
        guest_id=guest.id,
        stay_id=None,
        room_id=None,
        user_id=user.id,
        settings=settings,
    )
    draft.guest_name = guest.full_name or guest.name or draft.guest_name
    draft.guest_phone = guest.phone or draft.guest_phone
    draft.guest_address = guest.address or draft.guest_address
    draft.guest_pan = guest.pan_vat_number or draft.guest_pan
    draft.customer_type = guest.customer_type or draft.customer_type
    db.commit()
    return RedirectResponse(toast_url(f"/billing/create?draft_id={draft.id}", "Walk-in guest added.", "success"), status_code=303)


@app.get("/billing/create", response_class=HTMLResponse)
@app.get("/billing/new", response_class=HTMLResponse, include_in_schema=False)
def new_invoice_page(
    request: Request, draft_id: int | None = None, room_id: int | None = None,
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    page_data = _billing_page_data(db)
    draft = None
    if draft_id is not None:
        draft = db.scalar(
            select(Invoice)
            .options(selectinload(Invoice.guest), selectinload(Invoice.room),
                     selectinload(Invoice.items))
            .where(Invoice.id == draft_id, Invoice.status == "draft")
        )
        if draft is None:
            raise HTTPException(404, "Draft invoice not found.")
        _draft_access(draft, user)
        page_data["invoice_number"] = draft.invoice_number
        page_data["draft_kots"] = [
            _kot_dict(kot) for kot in db.scalars(
                select(Kot).where(Kot.invoice_id == draft.id).order_by(Kot.id)
            ).all()
        ]
        page_data["draft_payload"] = {
            "id": draft.id,
            "guest_id": draft.guest_id,
            "guest_name": draft.guest_name or draft.guest.name,
            "guest_phone": draft.guest_phone or draft.guest.phone,
            "guest_email": draft.guest.email,
            "guest_address": draft.guest_address or draft.guest.address,
            "guest_pan": draft.guest_pan or draft.guest.pan_vat_number,
            "invoice_type": draft.invoice_type,
            "invoice_date": local_datetime(draft.issue_date, "%Y-%m-%d"),
            "language": draft.language,
            "copy_type": draft.copy_type,
            "payment_terms": draft.payment_terms,
            "check_in": local_datetime(draft.check_in, "%Y-%m-%dT%H:%M") if draft.check_in else "",
            "check_out": local_datetime(draft.check_out, "%Y-%m-%dT%H:%M") if draft.check_out else "",
            "nights": draft.nights,
            "notes": draft.notes,
            "discount_amount": str(draft.discount_amount),
            "paid_amount": str(draft.paid_amount),
            "customer_type": draft.customer_type,
            "payment_method": draft.payment_method,
            "items": [
                {
                    "key": f"draft_{item.id}",
                    "room_id": draft.room_id if item.description.startswith("Room ") else None,
                    "description": item.description,
                    "short": item.description,
                    "quantity": float(item.quantity),
                    "rate": float(item.rate),
                    "discount": float(item.discount_amount),
                    "taxable": item.taxable,
                    "category": item.category,
                    "kot_id": item.kot_id,
                }
                for item in draft.items
            ],
        }
    elif room_id is not None:
        # Room opened straight from the board (double-click): preload its
        # stay folio (room nights + dishes/services) so nothing is lost
        # before checkout.
        prefill = _billing_room_prefill(db, room_id)
        if prefill is not None:
            stay = db.scalar(
                select(Stay).where(Stay.room_id == room_id, Stay.status == "checked_in")
            )
            stay_invoice = None
            if stay and stay.invoice_number:
                # A room with an already-invoiced stay: prefill the billing page
                # with the existing invoice so the user can add charges to it.
                stay_invoice = stay.invoice_number
                page_data["prefill_invoice_number"] = stay_invoice
                page_data["draft_payload"] = None
                # Load the existing invoice items so they can be continued
                existing_invoice = db.scalar(
                    select(Invoice)
                    .options(selectinload(Invoice.items))
                    .where(Invoice.stay_id == stay.id, Invoice.status != "void")
                )
                if existing_invoice and existing_invoice.items:
                    page_data["existing_invoice_items"] = [
                        {
                            "key": f"existing_{item.id}",
                            "room_id": item.room_id if item.description.startswith("Room ") else None,
                            "description": item.description,
                            "short": item.description,
                            "quantity": float(item.quantity),
                            "rate": float(item.rate),
                            "discount": float(item.discount_amount),
                            "taxable": item.taxable,
                            "category": item.category,
                            "kot_id": item.kot_id,
                        }
                        for item in existing_invoice.items
                    ]
                else:
                    page_data["existing_invoice_items"] = []
                page_data["toast_message"] = f"Room already has invoice {stay_invoice}. Adding charges to that invoice."
                page_data["toast_type"] = "info"
            else:
                # No existing invoice for this stay: create a new draft with room prefill
                page_data["prefill_invoice_number"] = ""
                page_data["draft_payload"] = prefill
                page_data["already_invoiced"] = False

    else:
        page_data["draft_payload"] = None
    return render(request, "new_invoice.html", db, page_data)


@app.post("/billing/create/items/add")
async def billing_add_item(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
    except HTTPException:
        pass
    kind = str(form.get("kind", ""))
    key = secrets.token_urlsafe(8)
    try:
        quantity = float(form.get("quantity", 1) or 1)
    except (TypeError, ValueError) as exc:
        raise HTTPException(422, "Enter a valid item quantity.") from exc
    if quantity <= 0 or quantity > 365:
        raise HTTPException(422, "Enter a valid item quantity.")
    if kind == "room":
        try:
            room_id = int(str(form.get("room_id", "0")))
        except ValueError as exc:
            raise HTTPException(422, "Select a valid room.") from exc
        room = db.get(Room, room_id)
        if not room:
            raise HTTPException(404, "Room not found. Refresh the room list.")
        item = {"key": key, "room_id": room.id, "description": f"Room {room.number} - {room.type}",
                "short": f"{room.number} · {room.type}", "quantity": quantity,
                "rate": float(room.nightly_rate), "discount": 0}
    elif kind == "catalog":
        try:
            catalog_id = int(str(form.get("catalog_id", "0")))
        except ValueError as exc:
            raise HTTPException(422, "Select a valid catalog item.") from exc
        catalog_item = db.get(CatalogItem, catalog_id)
        if not catalog_item or not catalog_item.active:
            raise HTTPException(404, "Item not found. Refresh the catalog.")
        item = {"key": key, "catalog_id": catalog_item.id, "description": catalog_item.name,
                "short": catalog_item.name, "quantity": quantity,
                "rate": float(catalog_item.rate), "discount": 0}
    elif kind == "custom":
        item = {"key": key, "description": "Custom item", "short": "Custom item",
                "quantity": 1, "rate": 0, "discount": 0}
    else:
        raise HTTPException(422, "Invalid item type.")
    return Response(status_code=204, headers={"HX-Trigger": json.dumps({"billing-item-added": item})})


@app.post("/billing/create/items/remove")
async def billing_remove_item(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    form = await request.form()
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
    except HTTPException:
        pass
    key = str(form.get("key", ""))
    return Response(status_code=204, headers={"HX-Trigger": json.dumps({"billing-item-removed": {"key": key}})})


@app.post("/billing/create/preview", response_class=HTMLResponse)
async def billing_preview_refresh(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    return Response(status_code=204, headers={"HX-Trigger": "billing-preview-refresh"})


@app.post("/billing/create")
async def create_invoice(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    invoice = None
    try:
        check_csrf(request, str(form.get("csrf_token", "")))
        rows = json.loads(str(form.get("line_items", "[]")))
        if not isinstance(rows, list) or not rows:
            raise ValueError("Add at least one room, restaurant item, or service before creating the invoice.")
        name = str(form.get("guest_name", "Walk-in guest")).strip()[:160] or "Walk-in guest"
        guest_id = str(form.get("guest_id", "")).strip()
        guest = None
        if guest_id:
            if not guest_id.isdigit():
                raise ValueError("Select a valid guest.")
            guest = db.get(Guest, int(guest_id))
            if guest is None:
                raise ValueError("The selected guest no longer exists. Refresh and try again.")

        check_in_raw, check_out_raw = str(form.get("check_in", "")), str(form.get("check_out", ""))
        check_in_local = datetime.fromisoformat(check_in_raw) if check_in_raw else local_now().replace(tzinfo=None)
        check_out_local = datetime.fromisoformat(check_out_raw) if check_out_raw else check_in_local + timedelta(days=1)
        if check_out_local <= check_in_local:
            raise ValueError("Check-out must be after check-in.")
        check_in = check_in_local.replace(tzinfo=BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
        check_out = check_out_local.replace(tzinfo=BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
        nights = max(1, (check_out_local.date() - check_in_local.date()).days)
        room_ids = []
        parsed = []
        lines = []
        row_keys = []
        row_kot_ids = []
        for row in rows:
            if not isinstance(row, dict):
                raise ValueError("Invalid invoice item.")
            description = str(row.get("description", "")).strip()[:240]
            try:
                quantity = Decimal(str(row.get("quantity", 1)))
                rate = Decimal(str(row.get("rate", 0)))
                discount = Decimal(str(row.get("discount", 0)))
            except (InvalidOperation, ValueError, TypeError):
                raise ValueError("Quantities, rates, and discounts must be valid numbers.")
            _finite(quantity, "Quantity")
            _finite(rate, "Rate")
            _finite(discount, "Discount")
            if quantity <= 0 or quantity > MAX_QUANTITY_INPUT:
                raise ValueError(f"Quantity must be between 0 and {MAX_QUANTITY_INPUT}.")
            if rate < 0 or rate > MAX_MONEY_INPUT:
                raise ValueError(f"Rate must be between 0 and {MAX_MONEY_INPUT}.")
            if discount < 0 or discount > MAX_MONEY_INPUT:
                raise ValueError(f"Discount must be between 0 and {MAX_MONEY_INPUT}.")
            if not description:
                raise ValueError("Each item needs a description.")
            base = (quantity * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if discount > base:
                raise ValueError(f"Discount for '{description}' cannot exceed its line amount ({base}).")
            raw_room = row.get("room_id")
            try:
                room_id = int(raw_room) if raw_room else None
            except (TypeError, ValueError) as exc:
                raise ValueError("Select a valid room.") from exc
            if room_id and room_id not in room_ids:
                room_ids.append(room_id)
            taxable = str(row.get("taxable", "true")).lower() not in {"false", "0", "off"}
            category = str(row.get("category", "food"))[:30]
            line = Line(
                description=description, qty=float(quantity), unit_price=float(rate),
                taxable=taxable, category=category, discount=float(discount),
            )
            lines.append(line)
            parsed.append((line, room_id))
            row_keys.append(str(row.get("key", ""))[:80])
            raw_kot_id = str(row.get("kot_id", "") or "").strip()
            row_kot_ids.append(int(raw_kot_id) if raw_kot_id.isdigit() else None)
        for room_id in room_ids:
            room = db.get(Room, room_id)
            if room is None:
                raise ValueError("One of the selected rooms no longer exists. Refresh and try again.")
            if not room.active:
                raise ValueError(f"Room {room.number} is archived and cannot be billed.")
            if room.status not in {"available", "occupied"}:
                raise ValueError(
                    f"Room {room.number} is not available for billing right now "
                    f"(status: {room.status.replace('_', ' ')})."
                )
        invoice_date_val = str(form.get("invoice_date", local_now().date().isoformat()))
        try:
            invoice_date = datetime.fromisoformat(invoice_date_val)
        except ValueError as exc:
            raise ValueError("Enter a valid invoice date.") from exc
        invoice_date_utc = datetime.combine(invoice_date.date(), time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
        action = str(form.get("action", "issue"))
        if action not in {"draft", "issue", "print", "pdf", "preview"}:
            raise ValueError("Choose a valid invoice action.")
        try:
            kot_pages_raw = json.loads(str(form.get("kots_data_json", "[]") or "[]"))
        except (TypeError, ValueError):
            kot_pages_raw = []
        if not isinstance(kot_pages_raw, list):
            raise ValueError("Invalid KOT data.")
        if len(kot_pages_raw) > 50:
            raise ValueError("Too many KOT pages on one invoice.")
        kot_pages = []
        for page in kot_pages_raw:
            if not isinstance(page, dict):
                continue
            raw_kot_items = page.get("items")
            if not isinstance(raw_kot_items, list) or not raw_kot_items:
                continue
            kot_pages.append({
                "table_number": str(page.get("tableNo") or page.get("table_number") or "").strip()[:40],
                "guest_id": str(page.get("guestId") or page.get("guest_id") or "").strip(),
                "guest_name": str(page.get("guestName") or page.get("guest_name") or "").strip()[:160],
                "items": _clean_kot_items(raw_kot_items),
            })
        invoice_type = str(form.get("invoice_type", "Tax Invoice"))
        if invoice_type not in {"Tax Invoice", "Bill of Supply", "Proforma"}:
            raise ValueError("Choose Tax Invoice, Bill of Supply, or Proforma.")
        bill_type = "supply" if invoice_type == "Bill of Supply" else "tax"
        settings = _billing_settings(db)
        language = str(form.get("language", "Nepali" if settings.get("invoice_language") == "ne" else "English"))
        copy_type = str(form.get("copy_type", settings.get("invoice_copy", "Customer Copy")))
        payment_terms = str(form.get("payment_terms", settings.get("payment_terms", "Payment due on presentation.")))
        notes = str(form.get("notes", "")).strip()[:2000]
        try:
            service_charge_rate = Decimal(settings.get("service_charge_rate", "10"))
            vat_rate = Decimal(settings.get("vat_rate", "13"))
            discount_value = Decimal(str(form.get("discount_value", "0") or "0"))
            paid_value = Decimal(str(form.get("amount_paid", "0") or "0"))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("Billing rates and discounts must be valid numbers.") from exc
        _bounded_money(service_charge_rate, "Service charge rate", Decimal("100"))
        _bounded_money(vat_rate, "VAT rate", Decimal("100"))
        _bounded_money(discount_value, "Discount")
        _bounded_money(paid_value, "Amount received")
        discount_kind = str(form.get("discount_kind", "amount"))
        if discount_kind not in {"amount", "percent"}:
            raise ValueError("Choose a valid discount type.")
        payment_method = str(form.get("payment_method", "Cash"))
        if payment_method not in _payment_methods(db):
            raise ValueError("Choose a valid payment method.")
        subtotal = sum((Decimal(str(line.total)) for line in lines), Decimal("0.00"))
        overall_discount = Decimal(str(discount_from_input(
            float(subtotal), discount_kind, float(discount_value)
        )))
        totals = compute_totals(
            lines,
            discount=float(overall_discount),
            service_charge_rate=float(service_charge_rate),
            vat_rate=float(vat_rate),
            paid=float(paid_value),
            bill_type=bill_type,
        )
        if paid_value > Decimal(str(totals.total)):
            raise ValueError("Amount received cannot exceed the invoice total.")
        guest_name = name
        preview_guest = Guest(
            name=guest_name,
            phone=str(form.get("guest_phone", "")).strip()[:40],
            email=str(form.get("guest_email", "")).strip()[:180],
            address=str(form.get("guest_address", "")).strip()[:240],
            pan_vat_number=str(form.get("guest_pan", "")).strip()[:80],
        )
        draft_id = str(form.get("invoice_id", "")).strip()
        existing_draft = None
        if draft_id:
            if not draft_id.isdigit():
                raise ValueError("Invalid draft invoice.")
            existing_draft = db.scalar(
                select(Invoice).where(Invoice.id == int(draft_id), Invoice.status == "draft")
            )
            if existing_draft is None:
                raise ValueError("This invoice is no longer a draft. Refresh and try again.")
            _draft_access(existing_draft, user)

        # In-memory preview mode: renders printable preview without creating duplicate DB records
        if action == "preview":
            preview_inv = Invoice(
                id=0,
                invoice_number=_next_billing_number(db),
                guest=preview_guest,
                room=db.get(Room, room_ids[0]) if room_ids else None,
                issue_date=invoice_date_utc,
                discount_amount=Decimal(str(totals.discount)),
                vat_amount=Decimal(str(totals.vat)),
                total=Decimal(str(totals.total)),
                paid_amount=Decimal(str(totals.paid)),
                subtotal=Decimal(str(totals.subtotal)),
                check_in=check_in,
                check_out=check_out,
                nights=nights,
                balance=Decimal(str(totals.balance)),
                status="draft",
                invoice_type=invoice_type,
                payment_method=payment_method,
                language=language,
                copy_type=copy_type,
                payment_terms=payment_terms,
                notes=notes,
                bill_type=bill_type,
                item_discount=Decimal(str(totals.item_discount)),
                net_subtotal=Decimal(str(totals.net_subtotal)),
                taxable_net=Decimal(str(totals.taxable_net)),
                service_charge_rate=Decimal(str(totals.service_charge_rate)),
                service_charge=Decimal(str(totals.service_charge)),
                vat_rate=Decimal(str(totals.vat_rate)),
            )
            for line, _room_id in parsed:
                line_total = Decimal(str(line.total))
                line_discount = Decimal(str(line.discount))
                vat = Decimal(str(compute_totals(
                    [line], service_charge_rate=0, vat_rate=float(vat_rate),
                    bill_type=bill_type,
                ).vat))
                preview_inv.items.append(InvoiceItem(
                    description=line.description, quantity=line.qty, rate=line.unit_price,
                    discount_amount=line_discount, vat_amount=vat,
                    amount=line_total - line_discount + vat, taxable=line.taxable,
                    category=line.category,
                ))
            return render(request, "print_invoice.html", db, {
                "invoice": preview_inv, "auto_print": False, "is_preview": True,
                "hotel_settings": settings,
                "printer_config": _printer_config(settings),
                "printer_logo_src": _printer_logo_source(settings),
            })

        if guest is None:
            guest = Guest(name=guest_name,
                          created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None))
            db.add(guest)
            db.flush()
        guest.name = guest_name
        guest_phone_value = str(form.get("guest_phone", "")).strip()[:40]
        guest.phone = guest_phone_value or None
        guest.email = str(form.get("guest_email", "")).strip()[:180]
        guest.address = str(form.get("guest_address", "")).strip()[:240]
        guest.pan_vat_number = str(form.get("guest_pan", "")).strip()[:80]

        stay_rows = []
        if action not in {"draft"}:
            for room_id in room_ids:
                room = db.get(Room, room_id)
                if room:
                    if room.status == "available":
                        stay = Stay(room_id=room.id, guest_id=guest.id,
                                    check_in=check_in, check_out=check_out, status="checked_in",
                                    business_date=_business_date(check_in).isoformat())
                        room.status = "occupied"
                        db.add(stay)
                        db.flush()
                        stay_rows.append(stay)
                    elif room.status == "occupied":
                        existing_stay = db.scalar(select(Stay).where(Stay.room_id == room.id, Stay.status == "checked_in"))
                        if existing_stay and existing_stay.invoice_number:
                            raise ValueError(
                                f"Room {room.number} already has an open invoice "
                                f"({existing_stay.invoice_number}). Complete or settle the "
                                "checkout for that stay instead."
                            )
                        if existing_stay and existing_stay.guest_id != guest.id:
                            raise ValueError(
                                f"Room {room.number} is occupied by another guest. "
                                "Check that guest out or start a new invoice for that stay."
                            )
                        if existing_stay:
                            stay_rows.append(existing_stay)

        invoice = existing_draft or Invoice(invoice_number=_next_billing_number(db))
        invoice.guest = guest
        invoice.room_id = room_ids[0] if room_ids else None
        invoice.stay_id = stay_rows[0].id if stay_rows else None
        invoice.issue_date = invoice_date_utc
        invoice.discount_amount = Decimal(str(totals.discount))
        invoice.item_discount = Decimal(str(totals.item_discount))
        invoice.net_subtotal = Decimal(str(totals.net_subtotal))
        invoice.taxable_net = Decimal(str(totals.taxable_net))
        invoice.service_charge_rate = Decimal(str(totals.service_charge_rate))
        invoice.service_charge = Decimal(str(totals.service_charge))
        invoice.vat_rate = Decimal(str(totals.vat_rate))
        invoice.vat_amount = Decimal(str(totals.vat))
        invoice.total = Decimal(str(totals.total))
        invoice.paid_amount = Decimal(str(totals.paid))
        invoice.balance = Decimal(str(totals.balance))
        invoice.subtotal = Decimal(str(totals.subtotal))
        invoice.check_in = check_in
        invoice.check_out = check_out
        invoice.nights = nights
        invoice.bill_date = invoice_date.date().isoformat()
        invoice.guest_name = guest.name
        invoice.guest_phone = guest.phone or ""
        invoice.guest_address = guest.address
        invoice.guest_pan = guest.pan_vat_number
        invoice.room_number = db.get(Room, room_ids[0]).number if room_ids else ""
        invoice.status = (
            "draft" if action == "draft"
            else "paid" if totals.balance == 0
            else "issued"
        )
        invoice.is_draft = action == "draft"
        invoice.invoice_type = invoice_type
        invoice.bill_type = bill_type
        invoice.language = language
        invoice.copy_type = copy_type
        invoice.payment_terms = payment_terms
        invoice.payment_method = payment_method
        invoice.notes = notes
        invoice.customer_type = str(form.get("customer_type", "individual"))
        invoice.prepared_by = user.username
        invoice.received_by = str(form.get("received_by", "")).strip()[:80]
        invoice.user_id = user.id
        invoice.items.clear()
        created_items = []
        for (line, _room_id), row_key, row_kot_id in zip(parsed, row_keys, row_kot_ids):
            line_total = Decimal(str(line.total))
            line_discount = Decimal(str(line.discount))
            line_totals = compute_totals(
                [line], service_charge_rate=0, vat_rate=float(vat_rate),
                bill_type=bill_type,
            )
            line_vat = Decimal(str(line_totals.vat))
            item_row = InvoiceItem(
                description=line.description, quantity=line.qty, rate=line.unit_price,
                discount_amount=line_discount, vat_amount=line_vat,
                amount=line_total - line_discount + line_vat,
                taxable=line.taxable, category=line.category,
                kot_id=row_kot_id,
            )
            invoice.items.append(item_row)
            created_items.append((item_row, row_key, row_kot_id))
        db.add(invoice)
        db.flush()
        for linked_kot_id in {kid for kid in row_kot_ids if kid}:
            linked_kot = db.get(Kot, linked_kot_id)
            if linked_kot is not None and linked_kot.invoice_id is None:
                linked_kot.invoice_id = invoice.id
        for page in kot_pages:
            kot_guest = _kot_guest(db, page["guest_id"]) or guest
            kot = Kot(
                kot_number=_next_kot_number(db),
                invoice_id=invoice.id,
                table_number=page["table_number"],
                guest_id=kot_guest.id if kot_guest else None,
                guest_name=(
                    (kot_guest.full_name or kot_guest.name) if kot_guest
                    else (page["guest_name"] or guest.name)
                ),
                items=json.dumps(page["items"]),
                total=sum((Decimal(str(row["amount"])) for row in page["items"]), Decimal("0.00")),
                status="saved",
                business_date=invoice_date.date().isoformat(),
                created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
                user_id=user.id,
            )
            db.add(kot)
            db.flush()
            for kot_row in page["items"]:
                key = kot_row.get("key")
                target = next(
                    (item_row for item_row, item_key, _kid in created_items
                     if key and item_key == key),
                    None,
                )
                if target is None:
                    target = next(
                        (item_row for item_row, _item_key, kid in created_items
                         if kid is None and item_row.category in {"restaurant", "food"}
                         and item_row.description == kot_row["particular"]
                         and float(item_row.quantity) == kot_row["qty"]
                         and float(item_row.rate) == kot_row["rate"]),
                        None,
                    )
                if target is not None and target.kot_id is None:
                    target.kot_id = kot.id
            _audit(db, user, "kot.create", "kot", kot.kot_number,
                   f"{page['table_number'] or 'Walk-in'} · {len(page['items'])} items")
        if action != "draft":
            # The invoice now carries the stay total; keep the room occupied
            # and let checkout settle this invoice instead of issuing a new one.
            for stay_row in stay_rows:
                stay_row.invoice_number = invoice.invoice_number
                stay_row.room_number = stay_row.room_number or (
                    db.get(Room, stay_row.room_id).number if db.get(Room, stay_row.room_id) else ""
                )
        _audit(db, user, f"invoice.{action}", "invoice", invoice.invoice_number,
               f"{invoice.total} {invoice.invoice_type}, rooms={room_ids or 'none'}")
        if action != "draft" and totals.paid > 0:
            db.add(Payment(
                invoice=invoice,
                amount=Decimal(str(totals.paid)),
                method=payment_method,
                received_by=user.username,
                note="Payment received at invoice issue",
                invoice_number=invoice.invoice_number,
                business_date=invoice_date.date().isoformat(),
                paid_at=invoice_date_utc,
                user_id=user.id,
            ))
        db.commit()
    except HTTPException:
        db.rollback()
        raise
    except (ValueError, TypeError, InvalidOperation, IntegrityError) as exc:
        db.rollback()
        page_data = _billing_page_data(db)
        page_data["error"] = str(exc)
        page_data["server_toast"] = toast_event(str(exc), "error")
        return render(request, "new_invoice.html", db, page_data, status_code=422)

    if action == "draft":
        return RedirectResponse(toast_url(
            f"/invoices?q={invoice.invoice_number}",
            f"Invoice {invoice.invoice_number} saved as a draft.",
        ), status_code=303)
    if action == "print":
        return RedirectResponse(toast_url(
            f"/billing/invoices/{invoice.id}/receipt",
            f"Receipt {invoice.invoice_number} ready to print.",
            "success",
        ), status_code=303)
    if action == "pdf":
        return RedirectResponse(toast_url(
            f"/billing/invoices/{invoice.id}/pdf",
            f"Invoice {invoice.invoice_number} downloaded.",
            "success",
        ), status_code=303)
    return RedirectResponse(toast_url(
        f"/invoices?q={invoice.invoice_number}",
        f"Invoice {invoice.invoice_number} created for {invoice.guest_name}.",
    ), status_code=303)


@app.get("/kots", response_class=HTMLResponse)
def kots_page(
    request: Request, invoice_id: int | None = None, status: str = "all",
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    query = select(Kot).options(
        selectinload(Kot.guest), selectinload(Kot.issued_by), selectinload(Kot.cut_by),
    ).order_by(Kot.created_at.desc(), Kot.id.desc())
    if invoice_id is not None:
        query = query.where(Kot.invoice_id == invoice_id)
    if status == "open":
        query = query.where(Kot.status.in_(KOT_OPEN_STATUSES))
    elif status == "new":
        query = query.where(Kot.status.in_(("new", "saved", "open")))
    elif status in KOT_STATUSES:
        query = query.where(Kot.status == status)
    else:
        status = "all"
    kots = [_kot_dict(kot) for kot in db.scalars(query).all()]
    return render(request, "kots.html", db, {
        "kots": kots,
        "invoice_id": invoice_id,
        "status_filter": status,
        "status_options": [
            ("all", "All"), ("open", "Open"), ("new", "New"),
            ("preparing", "Preparing"), ("ready", "Ready"), ("served", "Served"),
        ],
        "total": round(sum(row["total"] for row in kots), 2),
        "active_page": "kots",
        "header_title": "Kitchen Order Tickets",
    })


@app.get("/kots/by-table")
def kots_by_table(
    request: Request, table: str = "", invoice_id: int | None = None,
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    query = select(Kot).options(selectinload(Kot.guest)).order_by(Kot.id)
    if invoice_id is not None:
        query = query.where(Kot.invoice_id == invoice_id)
    elif table.strip() and table.strip().lower() not in {"all", "walk-in"}:
        query = query.where(func.lower(Kot.table_number) == table.strip().lower())
    return JSONResponse({"kots": [_kot_dict(kot) for kot in db.scalars(query).all()]})


@app.post("/kots/api/save")
async def kots_api_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    if not isinstance(payload, dict):
        raise HTTPException(422, "Invalid KOT payload.")
    check_csrf(request, str(payload.get("csrf_token", "")))
    items = _clean_kot_items(payload.get("items"))
    total = sum((Decimal(str(row["amount"])) for row in items), Decimal("0.00"))

    kot_id = str(payload.get("kot_id") or "").strip()
    kot = None
    if kot_id:
        if not kot_id.isdigit():
            raise HTTPException(422, "Invalid KOT.")
        kot = db.get(Kot, int(kot_id))
        if kot is None:
            raise HTTPException(404, "KOT not found.")
    invoice_id = str(payload.get("invoice_id") or "").strip()
    invoice = None
    if invoice_id:
        if not invoice_id.isdigit():
            raise HTTPException(422, "Invalid invoice for this KOT.")
        invoice = db.get(Invoice, int(invoice_id))
        if invoice is None or invoice.status == "void":
            raise HTTPException(404, "Invoice not found for this KOT.")
    guest = _kot_guest(db, payload.get("guest_id"))
    guest_name = (guest.full_name or guest.name) if guest else str(payload.get("guest_name") or "").strip()[:160]
    if kot is None:
        kot = Kot(
            kot_number=_next_kot_number(db),
            issued_by_name=_staff_name(user),
            created_at=local_now().astimezone(timezone.utc).replace(tzinfo=None),
            user_id=user.id,
        )
        db.add(kot)
    kot.table_number = str(payload.get("table_number") or "").strip()[:40]
    kot.guest_id = guest.id if guest else None
    kot.guest_name = guest_name
    if not kot.issued_by_name:
        kot.issued_by_name = _staff_name(user)
    kot.items = json.dumps(items)
    kot.total = total
    kot.status = "saved"
    kot.business_date = _business_date(local_now()).isoformat()
    kot.updated_at = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    if invoice is not None:
        kot.invoice_id = invoice.id
    db.commit()
    return JSONResponse({
        "ok": True, "id": kot.id, "kot_number": kot.kot_number,
        "total": float(kot.total or 0), "invoice_id": kot.invoice_id,
    })


def _kot_menu_items(db: Session) -> list[CatalogItem]:
    return list(db.scalars(
        select(CatalogItem)
        .where(
            CatalogItem.active.is_(True),
            CatalogItem.category.in_(("restaurant", "food", "beverage")),
        )
        .order_by(CatalogItem.menu_group, CatalogItem.sort_order, CatalogItem.name)
    ).all())


def _kot_menu_dict(item: CatalogItem) -> dict:
    return {
        "id": item.id,
        "name": item.name,
        "rate": float(item.rate or 0),
        "category": item.category or "food",
        "group": item.menu_group or "Menu",
        "unit": item.unit or "pcs",
        "taxable": bool(item.taxable),
    }


def _occupied_rooms(db: Session) -> list[dict]:
    """Rooms with a checked-in guest — the room picker for both KOT screens."""
    stays = db.scalars(
        select(Stay)
        .options(selectinload(Stay.guest), selectinload(Stay.room))
        .where(Stay.status == "checked_in")
        .order_by(Stay.room_number, Stay.id)
    ).all()
    rooms = []
    for stay in stays:
        room = stay.room or db.get(Room, stay.room_id)
        if room is None or not room.active:
            continue
        guest = stay.guest or (db.get(Guest, stay.guest_id) if stay.guest_id else None)
        rooms.append({
            "id": room.id,
            "number": room.number,
            "type": room.type,
            "stay_id": stay.id,
            "guest_id": stay.guest_id,
            "guest_name": stay.guest_name or (guest.full_name or guest.name if guest else "") or "Guest",
            "check_in": local_datetime(stay.check_in, "%d %b · %H:%M") if stay.check_in else "",
        })
    return rooms


def _api_user(request: Request, db: Session) -> User:
    """Resolve the signed-in user for a JSON client.

    The page routes answer anonymous callers with a 303 to ``/login``, which a
    phone app cannot act on, so these report 401 instead. ``/api`` is not in
    ``STAFF_ALLOWED_SECTIONS`` (that list is written for page paths) and kitchen
    tickets are front-desk work, so both roles are allowed here.
    """
    try:
        user = auth.current_user(request, db)
    except HTTPException as exc:
        if exc.status_code in (301, 302, 303, 307, 308):
            raise HTTPException(401, "Sign in to continue.") from exc
        raise
    if user.must_change_password:
        raise HTTPException(403, "Set a new password in the web app before using the mobile app.")
    return user


@app.get("/api/rooms")
def api_rooms(request: Request, db: Session = Depends(get_db)) -> JSONResponse:
    """Occupied rooms for the mobile KOT app.

    ``csrf`` rides along because the client must post it back on
    ``/mobile/kot/save``, and signing in rotates the token.
    """
    _api_user(request, db)
    return JSONResponse({"rooms": _occupied_rooms(db), "csrf": csrf(request)})


@app.get("/api/menu")
def api_menu(request: Request, db: Session = Depends(get_db)) -> JSONResponse:
    """Active restaurant menu items for the mobile KOT app."""
    _api_user(request, db)
    return JSONResponse({
        "menu": [_kot_menu_dict(item) for item in _kot_menu_items(db)],
        "csrf": csrf(request),
    })


def _open_kots(db: Session) -> list[dict]:
    return [
        _kot_dict(kot) for kot in db.scalars(
            select(Kot).options(
                selectinload(Kot.guest), selectinload(Kot.issued_by),
                selectinload(Kot.cut_by),
            )
            .where(Kot.status.in_(KOT_OPEN_STATUSES))
            .order_by(Kot.created_at.asc(), Kot.id.asc())
        ).all()
    ]


def _kot_delete_blocker(db: Session, kot: Kot) -> str | None:
    """Explain why a ticket cannot be deleted, or return None when it is safe.

    Deleting reverses the room-folio charges the KOT posted, so it is refused
    once any of those charges (or the ticket itself) has reached an invoice.
    """
    if kot.invoice_id is not None:
        return "This KOT is already on an invoice."
    charge_on_invoice = db.scalar(
        select(func.count(RoomCharge.id)).where(
            RoomCharge.kot_id == kot.id, RoomCharge.invoice_number != "",
        )
    )
    if charge_on_invoice:
        return "This KOT is already on an invoice."
    billed_items = db.scalar(
        select(func.count(InvoiceItem.id)).where(InvoiceItem.kot_id == kot.id)
    )
    if billed_items:
        return "This KOT is already on an invoice."
    return None


def _delete_kot(db: Session, kot: Kot) -> int:
    """Delete the ticket and the room charges it created; return charge count."""
    charges = db.scalars(select(RoomCharge).where(RoomCharge.kot_id == kot.id)).all()
    for charge in charges:
        db.delete(charge)
    db.delete(kot)
    db.flush()
    return len(charges)


@app.get("/api/kots")
def api_kots(request: Request, db: Session = Depends(get_db)) -> JSONResponse:
    """Open kitchen tickets for the mobile app, newest last.

    Any signed-in device can see and cut them; the ticket lives on the server,
    so it does not matter which phone created it.
    """
    _api_user(request, db)
    return JSONResponse({"kots": _open_kots(db), "csrf": csrf(request)})


@app.post("/api/kots/{kot_id}/cut")
async def api_kot_cut(
    request: Request, kot_id: int, db: Session = Depends(get_db),
) -> JSONResponse:
    """Cut a KOT from the mobile app and record the signed-in user as the cook.

    When ``kot_double_confirm`` is on, the first call answers ``428`` so the
    client can ask again; the second call must send ``confirm: true``.
    """
    user = _api_user(request, db)
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    if not isinstance(payload, dict):
        raise HTTPException(422, "Invalid KOT payload.")
    check_csrf(request, str(payload.get("csrf_token", "")))

    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")

    settings = _billing_settings(db)
    double_confirm = settings.get("kot_double_confirm", "1").strip().lower() not in {
        "0", "false", "no", "off",
    }
    confirmed = payload.get("confirm") in {True, 1, "1", "true", "yes", "on"}
    if double_confirm and not confirmed:
        raise HTTPException(428, "Confirm the KOT again before cutting it.")

    current = _kot_status(kot.status)
    if current not in {"new", "preparing"}:
        raise HTTPException(
            409, f"KOT {kot.kot_number} is already {_kot_status_label(current).lower()}.",
        )

    now_utc = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    result = db.execute(
        update(Kot)
        .where(Kot.id == kot.id, Kot.status.in_(("new", "preparing", "saved", "open")))
        .values(status="ready", cut_at=now_utc, cut_by_id=user.id, updated_at=now_utc)
    )
    if result.rowcount == 0:
        db.rollback()
        raise HTTPException(
            409, f"KOT {kot.kot_number} was already cut on another device.",
        )
    _audit(
        db, user, "kot.cut", "kot", kot.kot_number,
        f"Ready by {_staff_name(user)} · room {kot.room_number or kot.table_number or 'walk-in'}",
    )
    db.commit()
    db.refresh(kot)
    return JSONResponse({"ok": True, "kot": _kot_dict(kot)})


@app.post("/api/kots/{kot_id}/status")
async def api_kot_status(
    request: Request, kot_id: int, db: Session = Depends(get_db),
) -> JSONResponse:
    """Move a KOT to preparing/ready/served from the mobile app.

    Mirrors the web forms: forward-only transitions, and ``ready`` records the
    cook and may need the same double confirmation as ``/api/kots/{id}/cut``.
    """
    user = _api_user(request, db)
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    if not isinstance(payload, dict):
        raise HTTPException(422, "Invalid KOT payload.")
    check_csrf(request, str(payload.get("csrf_token", "")))

    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")

    target = str(payload.get("status") or "").strip()
    current = _kot_status(kot.status)
    if not _kot_can_transition(current, target):
        raise HTTPException(409, "That kitchen status change is not allowed.")

    if target == "ready":
        settings = _billing_settings(db)
        double_confirm = settings.get("kot_double_confirm", "1").strip().lower() not in {
            "0", "false", "no", "off",
        }
        confirmed = payload.get("confirm") in {True, 1, "1", "true", "yes", "on"}
        if double_confirm and not confirmed:
            raise HTTPException(428, "Confirm the KOT again before cutting it.")

    now_utc = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    values = {"status": target, "updated_at": now_utc}
    if target == "ready" and kot.cut_at is None:
        values["cut_at"] = now_utc
        values["cut_by_id"] = user.id
    result = db.execute(
        update(Kot)
        .where(Kot.id == kot.id, Kot.status == kot.status)
        .values(**values)
    )
    if result.rowcount == 0:
        db.rollback()
        raise HTTPException(409, "This KOT changed on another device. Refresh and retry.")
    _audit(db, user, "kot.status", "kot", kot.kot_number, f"{current} → {target}")
    db.commit()
    db.refresh(kot)
    return JSONResponse({"ok": True, "kot": _kot_dict(kot)})


@app.post("/api/kots/{kot_id}/delete")
async def api_kot_delete(
    request: Request, kot_id: int, db: Session = Depends(get_db),
) -> JSONResponse:
    """Delete a ticket from the kitchen board and reverse its room charges.

    The client must send ``confirm: true`` so an accidental tap cannot wipe a
    ticket. Tickets already billed to an invoice are refused.
    """
    user = _api_user(request, db)
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    if not isinstance(payload, dict):
        raise HTTPException(422, "Invalid KOT payload.")
    check_csrf(request, str(payload.get("csrf_token", "")))

    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    confirmed = payload.get("confirm") in {True, 1, "1", "true", "yes", "on"}
    if not confirmed:
        raise HTTPException(428, "Confirm deleting the KOT.")

    blocker = _kot_delete_blocker(db, kot)
    if blocker:
        raise HTTPException(409, blocker)

    kot_number = kot.kot_number
    charges = _delete_kot(db, kot)
    _audit(
        db, user, "kot.delete", "kot", kot_number,
        f"Deleted · {charges} room charge(s) reversed",
    )
    db.commit()
    return JSONResponse({
        "ok": True,
        "kot_number": kot_number,
        "reversed_charges": charges,
    })


@app.get("/api/kots/stream")
async def api_kots_stream(request: Request):
    """Server-sent events so the kitchen display updates without a refresh.

    The server polls its own table every couple of seconds and only pushes when
    the open-ticket set (count or newest update) changes.
    """
    session = SessionLocal()
    try:
        _api_user(request, session)
    finally:
        session.close()

    async def events():
        last = None
        yield "retry: 3000\n\n"
        while not await request.is_disconnected():
            session = SessionLocal()
            try:
                count, newest = session.execute(
                    select(func.count(Kot.id), func.max(Kot.updated_at))
                    .where(Kot.status.in_(KOT_OPEN_STATUSES))
                ).one()
            finally:
                session.close()
            version = f"{count}:{newest.isoformat() if newest else ''}"
            if version != last:
                last = version
                yield f"data: {json.dumps({'version': version})}\n\n"
            await asyncio.sleep(2)

    return StreamingResponse(
        events(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/mobile/kot", response_class=HTMLResponse)
def mobile_kot_page(
    request: Request, room_id: int | None = None, db: Session = Depends(get_db),
):
    """Phone-first KOT entry for waiters: pick a room, tap menu items, send.

    Saving here creates the kitchen ticket and posts the same items to the
    occupied room's folio, so checkout already has the food charges.
    """
    user_or_redirect(request, db)
    rooms = _occupied_rooms(db)
    menu = [_kot_menu_dict(item) for item in _kot_menu_items(db)]
    open_kots = [
        _kot_dict(kot) for kot in db.scalars(
            select(Kot).options(
                selectinload(Kot.guest), selectinload(Kot.issued_by),
                selectinload(Kot.cut_by),
            )
            .where(Kot.status.in_(KOT_OPEN_STATUSES))
            .order_by(Kot.created_at.asc(), Kot.id.asc())
        ).all()
    ]
    today = _business_date(local_now()).isoformat()
    served_kots = [
        _kot_dict(kot) for kot in db.scalars(
            select(Kot).options(
                selectinload(Kot.guest), selectinload(Kot.issued_by),
                selectinload(Kot.cut_by),
            )
            .where(Kot.status == "served", Kot.business_date == today)
            .order_by(Kot.updated_at.desc(), Kot.id.desc())
            .limit(20)
        ).all()
    ]
    return render(request, "mobile_kot.html", db, {
        "rooms_data": rooms,
        "menu_items": menu,
        "open_kots": open_kots,
        "served_kots": served_kots,
        "selected_room_id": room_id,
        "active_page": "kots",
        "header_title": "Mobile KOT",
    })


def _kot_room_charge_category(item: dict) -> str:
    category = str(item.get("category") or "").lower()
    group = str(item.get("menu_group") or "").lower()
    if category in {"beverage", "beverages", "drink", "drinks", "bar"}:
        return "beverage"
    if any(word in group for word in ("beverage", "drink", "tea", "coffee", "bar", "juice")):
        return "beverage"
    return "food"


@app.post("/mobile/kot/save")
async def mobile_kot_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    try:
        payload = await request.json()
    except Exception:
        payload = {}
    if not isinstance(payload, dict):
        raise HTTPException(422, "Invalid KOT payload.")
    check_csrf(request, str(payload.get("csrf_token", "")))
    try:
        items = _clean_kot_items(payload.get("items"))
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc

    raw_room_id = str(payload.get("room_id") or "").strip()
    mode = str(payload.get("mode") or "").strip().lower()
    walk_in = mode == "walkin" or (
        not mode and not raw_room_id.isdigit() and bool(
            str(payload.get("guest_name") or payload.get("table_number") or "").strip()
        )
    )
    room = None
    stay = None
    if walk_in:
        guest_name = str(payload.get("guest_name") or "").strip()[:160] or "Walk-in guest"
        table_number = str(payload.get("table_number") or "").strip()[:40] or "Walk-in"
        guest_id = None
    else:
        if not raw_room_id.isdigit():
            raise HTTPException(422, "Choose the room for this KOT, or switch to Walk-in.")
        room = db.get(Room, int(raw_room_id))
        if room is None or not room.active:
            raise HTTPException(404, "Room not found.")
        stay = db.scalar(
            select(Stay).where(Stay.room_id == room.id, Stay.status == "checked_in")
        )
        if stay is None:
            raise HTTPException(422, f"Room {room.number} has no checked-in guest.")

    for item in items:
        if not item.get("catalog_id"):
            raise HTTPException(422, "Choose every item from the menu list.")
        catalog_item = db.get(CatalogItem, int(item["catalog_id"]))
        if catalog_item is None or not catalog_item.active:
            raise HTTPException(422, "One of the selected menu items is no longer available.")
        item.update({
            "particular": catalog_item.name,
            "rate": float(catalog_item.rate or 0),
            "category": catalog_item.category or "food",
            "menu_group": catalog_item.menu_group or "",
            "unit": catalog_item.unit or "pcs",
            "taxable": bool(catalog_item.taxable),
        })
        item["amount"] = float((
            Decimal(str(item["qty"])) * Decimal(str(item["rate"]))
        ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    total = sum((Decimal(str(row["amount"])) for row in items), Decimal("0.00"))

    guest = db.get(Guest, stay.guest_id) if stay and stay.guest_id else None
    if not walk_in:
        guest_name = stay.guest_name or (guest.full_name or guest.name if guest else "") or "Guest"
    now_utc = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    kot = Kot(
        kot_number=_next_kot_number(db),
        room_id=room.id if room else None,
        room_number=room.number if room else "",
        table_number=table_number if walk_in else room.number,
        guest_id=(stay.guest_id or None) if stay else None,
        guest_name=guest_name,
        issued_by_name=(
            str(payload.get("issued_by_name") or "").strip()[:160] or _staff_name(user)
        ),
        items=json.dumps(items),
        total=total,
        status="new",
        notes=str(payload.get("notes") or "").strip()[:500],
        business_date=_business_date(local_now()).isoformat(),
        created_at=now_utc,
        updated_at=now_utc,
        user_id=user.id,
    )
    db.add(kot)
    db.flush()

    if not walk_in:
        charge_date = _business_date(local_now()).isoformat()
        for item in items:
            amount = Decimal(str(item["amount"]))
            if amount <= 0:
                continue
            db.add(RoomCharge(
                stay_id=stay.id,
                room_number=room.number,
                description=str(item["particular"])[:240],
                category=_kot_room_charge_category(item),
                quantity=Decimal(str(item["qty"])),
                unit_price=Decimal(str(item["rate"])),
                amount=amount,
                charge_date=charge_date,
                created_at=now_utc,
                user_id=user.id,
                kot_id=kot.id,
            ))
        _audit(
            db, user, "kot.create_room", "kot", kot.kot_number,
            f"Room {room.number} · {len(items)} items · {total:.2f} posted to folio",
        )
    else:
        _audit(
            db, user, "kot.create_walkin", "kot", kot.kot_number,
            f"{table_number} · {guest_name} · {len(items)} items · {total:.2f} (pay at counter)",
        )
    db.commit()
    return JSONResponse({
        "ok": True,
        "id": kot.id,
        "kot_number": kot.kot_number,
        "mode": "walkin" if walk_in else "room",
        "room_number": room.number if room else "",
        "table_number": kot.table_number,
        "guest_name": kot.guest_name,
        "total": float(kot.total or 0),
    })


@app.get("/kitchen", response_class=HTMLResponse)
def kitchen_page(request: Request, db: Session = Depends(get_db)):
    """Kitchen display: pending room KOTs, cut confirmation, and price setup."""
    user = user_or_redirect(request, db)
    settings = _billing_settings(db)
    double_confirm = settings.get("kot_double_confirm", "1").strip().lower() not in {
        "0", "false", "no", "off",
    }
    open_kots = [
        _kot_dict(kot) for kot in db.scalars(
            select(Kot).options(
                selectinload(Kot.guest), selectinload(Kot.issued_by),
                selectinload(Kot.cut_by),
            )
            .where(Kot.status.in_(KOT_OPEN_STATUSES))
            .order_by(Kot.created_at.asc(), Kot.id.asc())
        ).all()
    ]
    groups = {"new": [], "preparing": [], "ready": []}
    for kot in open_kots:
        groups.setdefault(kot["status"], []).append(kot)
    menu_items = [_kot_menu_dict(item) for item in _kot_menu_items(db)] if user.role == "admin" else []
    return render(request, "kitchen.html", db, {
        "new_kots": groups.get("new", []),
        "preparing_kots": groups.get("preparing", []),
        "ready_kots": groups.get("ready", []),
        "menu_items": menu_items,
        "kot_double_confirm": double_confirm,
        "active_page": "kitchen",
        "header_title": "Kitchen Display",
    })


@app.post("/kots/{kot_id}/cut")
async def kot_cut(request: Request, kot_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    next_url = str(form.get("next") or "").strip()
    if not (next_url.startswith("/kitchen") or next_url.startswith("/kots")
            or next_url.startswith("/mobile/kot")):
        next_url = "/kitchen"
    settings = _billing_settings(db)
    double_confirm = settings.get("kot_double_confirm", "1").strip().lower() not in {
        "0", "false", "no", "off",
    }
    confirmed = str(form.get("cut_confirmed", "")).strip().lower() in {"1", "true", "yes", "on"}
    if double_confirm and not confirmed:
        return RedirectResponse(
            toast_url(next_url, "Confirm the KOT again before cutting it.", "warning"),
            status_code=303,
        )
    current = _kot_status(kot.status)
    if current not in {"new", "preparing"}:
        return RedirectResponse(
            toast_url(next_url, f"KOT {kot.kot_number} is already {_kot_status_label(current).lower()}.",
                       "warning"),
            status_code=303,
        )
    now_utc = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    kot.status = "ready"
    kot.cut_at = now_utc
    kot.cut_by_id = user.id
    kot.updated_at = now_utc
    _audit(
        db, user, "kot.cut", "kot", kot.kot_number,
        f"Ready by {_staff_name(user)} · room {kot.room_number or kot.table_number or 'walk-in'}",
    )
    db.commit()
    return RedirectResponse(
        toast_url(next_url, f"KOT {kot.kot_number} cut and marked ready.", "success"),
        status_code=303,
    )


@app.post("/kots/{kot_id}/delete")
async def kot_delete(request: Request, kot_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    next_url = str(form.get("next") or "").strip()
    if not (next_url.startswith("/kitchen") or next_url.startswith("/kots")
            or next_url.startswith("/mobile/kot")):
        next_url = "/kitchen"
    blocker = _kot_delete_blocker(db, kot)
    if blocker:
        return RedirectResponse(
            toast_url(next_url, blocker, "warning"),
            status_code=303,
        )
    kot_number = kot.kot_number
    charges = _delete_kot(db, kot)
    _audit(
        db, user, "kot.delete", "kot", kot_number,
        f"Deleted · {charges} room charge(s) reversed",
    )
    db.commit()
    return RedirectResponse(
        toast_url(
            next_url,
            f"KOT {kot_number} deleted; {charges} room charge(s) reversed.",
            "success",
        ),
        status_code=303,
    )


@app.post("/kitchen/settings")
async def kitchen_settings_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change kitchen settings.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    enabled = str(form.get("kot_double_confirm", "")).strip().lower() in {"1", "on", "true", "yes"}
    db.merge(HotelSetting(
        key="kot_double_confirm",
        value="1" if enabled else "0",
    ))
    _audit(
        db, user, "settings.kot_double_confirm", "setting", "kot_double_confirm",
        "enabled" if enabled else "disabled",
    )
    db.commit()
    return RedirectResponse(
        toast_url("/kitchen", "Double confirmation enabled." if enabled else "Double confirmation disabled.",
                  "success" if enabled else "warning"),
        status_code=303,
    )


@app.post("/kitchen/items/{item_id}/price")
async def kitchen_item_price(
    request: Request, item_id: int, db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change menu prices.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    item = db.get(CatalogItem, item_id)
    if item is None or item.category not in {"restaurant", "food", "beverage"}:
        raise HTTPException(404, "Menu item not found.")
    try:
        price = _bounded_money(Decimal(str(form.get("price", "")).strip()), "Price")
    except (InvalidOperation, ValueError):
        return RedirectResponse(
            toast_url("/kitchen", "Enter a valid menu price between 0 and 10 crore.", "error"),
            status_code=303,
        )
    old_price = item.rate or Decimal("0")
    item.rate = price.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    _audit(
        db, user, "catalog.price", "catalog_item", str(item.id),
        f"{item.name}: {old_price:.2f} → {item.rate:.2f}",
    )
    db.commit()
    return RedirectResponse(
        toast_url("/kitchen", f"{item.name} price updated to {item.rate:.2f}.", "success"),
        status_code=303,
    )


@app.post("/kots/{kot_id}/status")
async def kot_set_status(request: Request, kot_id: int, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    target = str(form.get("status", "")).strip()
    current = _kot_status(kot.status)
    next_url = str(form.get("next") or "").strip()
    if not (next_url.startswith("/kitchen") or next_url.startswith("/kots")
            or next_url.startswith("/mobile/kot")):
        next_url = "/kots"
    if not _kot_can_transition(current, target):
        return RedirectResponse(
            toast_url(next_url, "That kitchen status change is not allowed.", "warning"),
            status_code=303,
        )
    now_utc = local_now().astimezone(timezone.utc).replace(tzinfo=None)
    kot.status = target
    kot.updated_at = now_utc
    if target == "ready" and kot.cut_at is None:
        kot.cut_at = now_utc
        kot.cut_by_id = user.id
    _audit(db, user, "kot.status", "kot", kot.kot_number, f"{current} → {target}")
    db.commit()
    return RedirectResponse(
        toast_url(next_url, f"KOT {kot.kot_number} marked {_kot_status_label(target).lower()}.", "success"),
        status_code=303,
    )


@app.get("/kots/{kot_id}/view", response_class=HTMLResponse)
def kot_view(request: Request, kot_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    items = _kot_items(kot)
    return render(request, "kot_view.html", db, {
        "kot": kot, "items": items, "total": float(kot.total or 0),
        "status_label": _kot_status_label(kot.status),
        "next_status": _kot_next_status(kot.status),
        "next_status_label": _kot_status_label(_kot_next_status(kot.status))
        if _kot_next_status(kot.status) else "",
        "active_page": "kots", "header_title": f"KOT {kot.kot_number}",
    })


@app.get("/kots/{kot_id}/print", response_class=HTMLResponse)
def kot_print(request: Request, kot_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    kot = db.get(Kot, kot_id)
    if kot is None:
        raise HTTPException(404, "KOT not found.")
    settings = _billing_settings(db)
    return render(request, "print_kot.html", db, {
        "kot": kot,
        "items": _kot_items(kot),
        "total": float(kot.total or 0),
        "kot_guest_name": kot.guest_name or (kot.guest.full_name or kot.guest.name if kot.guest else ""),
        "hotel_settings": settings,
        "today": local_now().strftime("%Y-%m-%d %H:%M"),
    })


@app.post("/billing/invoices/{invoice_id}/delete-draft")
async def delete_invoice_draft(
    request: Request, invoice_id: int, csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can delete draft invoices.")
    check_csrf(request, csrf_token)
    invoice = db.get(Invoice, invoice_id)
    if invoice is None or invoice.status != "draft":
        raise HTTPException(404, "Draft invoice not found.")
    _draft_access(invoice, user)
    number = invoice.invoice_number
    db.delete(invoice)
    _audit(db, user, "invoice.delete_draft", "invoice", number, "Draft discarded.")
    db.commit()
    return RedirectResponse(toast_url("/invoices?status=draft", "Draft invoice deleted.", "warning"), status_code=303)


@app.post("/billing/invoices/{invoice_id}/void")
async def void_invoice(
    request: Request, invoice_id: int, csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can void invoices.")
    check_csrf(request, csrf_token)
    invoice = db.get(Invoice, invoice_id)
    if invoice is None or invoice.status in {"draft", "void"}:
        raise HTTPException(404, "Issued invoice not found.")
    if db.scalar(select(func.count(Payment.id)).where(Payment.invoice_id == invoice.id)):
        raise HTTPException(
            409,
            "This invoice has recorded payments. Refund the payment instead of voiding the invoice.",
        )
    invoice.status = "void"
    _audit(db, user, "invoice.void", "invoice", invoice.invoice_number,
           f"Voided {invoice.total} {invoice.invoice_type}")
    db.commit()
    return RedirectResponse(toast_url("/invoices?status=void", f"Invoice {invoice.invoice_number} voided.", "warning"), status_code=303)


def _load_invoice_for_print(db: Session, invoice_id: int) -> Invoice:
    invoice = db.scalar(select(Invoice).options(
        selectinload(Invoice.guest), selectinload(Invoice.room),
        selectinload(Invoice.items),
    ).where(Invoice.id == invoice_id))
    if not invoice:
        raise HTTPException(404, "Invoice not found.")
    return invoice


def _receipt_context(
    invoice: Invoice, settings: dict[str, str], *,
    paper: str | None = None, auto_print: bool | None = None,
) -> dict[str, object]:
    """What the thermal receipt prints.

    Every value is read from the stored invoice, never recomputed from the
    catalogue, so a receipt reprinted after the bill was paid, voided or
    corrected still shows exactly what was issued.
    """
    config = _receipt_printer_config(settings)
    if paper in {"58", "80"}:
        config["paper"] = paper

    date_format = settings.get("date_format", "DD/MM/YYYY")
    date_pattern = "%Y-%m-%d" if date_format == "YYYY-MM-DD" else "%d/%m/%Y"
    bill_type = (invoice.bill_type or "tax").strip().lower()
    total = float(invoice.total or 0)
    paid = float(invoice.paid_amount or 0)
    return {
        "invoice": invoice,
        "hotel_settings": settings,
        "printer_config": config,
        "printer_logo_src": _printer_logo_source(settings),
        "receipt_title": BILL_TYPE_LABELS.get(bill_type, "Tax Invoice").upper(),
        "copy_label": invoice.copy_label or settings.get("invoice_copy") or "Customer Copy",
        "receipt_date": local_datetime(invoice.issue_date, date_pattern),
        "receipt_time": local_datetime(invoice.issue_date, "%H:%M"),
        "status_text": payment_status(total, paid),
        "amount_words": amount_in_words(total),
        "auto_print": (config["workflow"] == "direct") if auto_print is None else auto_print,
    }


@app.get("/billing/invoices/{invoice_id}/print", response_class=HTMLResponse)
def print_invoice_page(
    request: Request, invoice_id: int, preview: bool = False,
    receipt: bool = False, db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    invoice = _load_invoice_for_print(db, invoice_id)
    settings = _billing_settings(db)
    return render(request, "receipt.html", db, _receipt_context(
        invoice, settings, auto_print=False if preview else None,
    ))


@app.get("/billing/invoices/{invoice_id}/pdf")
def invoice_pdf(request: Request, invoice_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    invoice = db.scalar(select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room), selectinload(Invoice.items)).where(Invoice.id == invoice_id))
    if not invoice:
        raise HTTPException(404, "Invoice not found.")
    settings = _billing_settings(db)
    try:
        from weasyprint import HTML
        html = templates.get_template("print_invoice.html").render(
            request=request, invoice=invoice, auto_print=False,
            is_preview=False,
            hotel_settings=settings,
            printer_config=_printer_config(settings),
            printer_logo_src=_printer_logo_source(settings, for_pdf=True),
        )
        pdf = HTML(string=html, base_url=str(ROOT)).write_pdf()
    except Exception:
        # Browser-print fallback keeps the invoice layout available when PDF
        # generation fails, without changing the receipt-only /print page.
        return RedirectResponse(
            f"/billing/invoices/{invoice_id}/invoice?fallback=1", status_code=303,
        )
    _save_pdf_copy(pdf, "invoices", f"{invoice.invoice_number}.pdf")
    return Response(pdf, media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="{invoice.invoice_number}.pdf"',
    })


@app.get("/billing/invoices/{invoice_id}/receipt", response_class=HTMLResponse)
def invoice_receipt(
    request: Request, invoice_id: int, paper: str | None = None,
    preview: bool = False, db: Session = Depends(get_db),
):
    """The thermal receipt — printable for every status, including paid and void."""
    user_or_redirect(request, db)
    invoice = _load_invoice_for_print(db, invoice_id)
    settings = _billing_settings(db)
    return render(request, "receipt.html", db, _receipt_context(
        invoice, settings, paper=paper, auto_print=False if preview else None,
    ))


@app.get("/billing/invoices/{invoice_id}/invoice", response_class=HTMLResponse)
def invoice_page(
    request: Request, invoice_id: int, db: Session = Depends(get_db),
):
    """The full invoice layout, opened from the receipt's View Invoice action."""
    user_or_redirect(request, db)
    invoice = _load_invoice_for_print(db, invoice_id)
    settings = _billing_settings(db)
    return render(request, "print_invoice.html", db, {
        "invoice": invoice,
        "auto_print": False,
        "is_preview": False,
        "receipt_mode": False,
        "hotel_settings": settings,
        "printer_config": _printer_config(settings),
        "printer_logo_src": _printer_logo_source(settings),
    })


@app.get("/billing/invoices/{invoice_id}/receipt.pdf")
def invoice_receipt_pdf(
    request: Request, invoice_id: int, paper: str | None = None,
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    invoice = _load_invoice_for_print(db, invoice_id)
    settings = _billing_settings(db)
    context = _receipt_context(invoice, settings, paper=paper, auto_print=False)
    try:
        from weasyprint import HTML

        context["printer_logo_src"] = _printer_logo_source(settings, for_pdf=True)
        html = templates.get_template("receipt.html").render(request=request, **context)
        pdf = HTML(string=html, base_url=str(ROOT)).write_pdf()
    except Exception:
        suffix = f"?paper={paper}" if paper in {"58", "80"} else ""
        return RedirectResponse(
            f"/billing/invoices/{invoice_id}/receipt{suffix}", status_code=303,
        )
    _save_pdf_copy(pdf, "receipts", f"{invoice.invoice_number}-receipt.pdf")
    return Response(pdf, media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="{invoice.invoice_number}-receipt.pdf"',
    })


@app.get("/invoices/{invoice_id}/collect", response_class=HTMLResponse)
def collect_modal(request: Request, invoice_id: int, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    invoice = db.scalar(select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room)).where(Invoice.id == invoice_id))
    if not invoice or invoice.status not in {"issued", "paid"} or invoice.balance <= 0:
        return HTMLResponse('<div class="fixed inset-0 z-50 grid place-items-center" style="background:rgba(27,58,42,.45);padding:1rem"><div class="rounded-2xl p-6 max-w-sm w-full text-center" style="background:var(--color-surface,#fff);box-shadow:0 24px 60px rgba(27,58,42,.25)"><p class="font-bold mb-2" style="color:var(--color-primary,#2E7D3E)">Invoice Settled</p><p class="text-sm mb-4" style="color:#6B6B60">This invoice has no outstanding balance.</p><button onclick="this.closest(\'.fixed\').remove()" class="rounded-lg border-0 px-4 py-2 text-sm font-bold" style="background:var(--color-primary,#1B3A2A);color:#F7F1DE;cursor:pointer">Close</button></div></div>', status_code=200)
    return render(request, "partials/payment_modal.html", db, {
        "invoice": invoice, "methods": _payment_methods(db),
    }, 200)


@app.post("/invoices/{invoice_id}/collect")
def collect_invoice(
    request: Request, invoice_id: int, amount: str = Form("0"), method: str = Form("Cash"),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    invoice = db.scalar(select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room)).where(Invoice.id == invoice_id))
    if not invoice:
        raise HTTPException(404, "Invoice not found.")
    if invoice.status not in {"issued", "paid"}:
        raise HTTPException(409, "Only issued invoices can receive payments.")
    if invoice.balance <= 0:
        raise HTTPException(409, "This invoice has no outstanding balance.")
    if method not in _payment_methods(db):
        raise HTTPException(422, "Choose a valid payment method.")
    try:
        value = Decimal(amount)
        collect_payment(db, invoice, value, method, user.username, user.id)
        _audit(db, user, "payment.receive", "invoice", invoice.invoice_number,
               f"{value:.2f} received by {method}")
        db.commit()
        if request.headers.get("HX-Request") == "true":
            return Response(status_code=204, headers={
                "HX-Trigger": json.dumps({
                    "payment-saved": {},
                    "showToast": toast_event("Payment recorded successfully.", "success"),
                }),
                "HX-Redirect": toast_url("/payments", "Payment recorded successfully.", "success"),
            })
        return RedirectResponse(toast_url("/payments", "Payment recorded successfully.", "success"), status_code=303)
    except (InvalidOperation, ValueError) as exc:
        db.rollback()
        invoice = db.scalar(select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room)).where(Invoice.id == invoice_id))
        return render(request, "partials/payment_modal.html", db, {
            "invoice": invoice, "methods": _payment_methods(db), "error": str(exc),
            "amount_val": amount, "method_val": method,
        }, status_code=200)
    except Exception:
        db.rollback()
        _log.exception("Payment collection failed for invoice %s", invoice_id)
        if request.headers.get("HX-Request") == "true":
            return Response(
                f'<div role="alert" class="rounded-xl border border-red-200 bg-red-50 p-3 text-sm font-semibold text-red-700">The payment could not be recorded. Please try again.</div>',
                status_code=200,
                headers={"HX-Trigger": json.dumps({"showToast": toast_event("The payment could not be recorded. Please try again.", "error")})},
            )
        return RedirectResponse(toast_url("/payments", "The payment could not be recorded. Please try again.", "error"), status_code=303)


@app.get("/payments", response_class=HTMLResponse)
def payments_page(
    request: Request, start: str = "", end: str = "", method: str = "All",
    invoice: str = "", db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    for label, value in (("start", start), ("end", end)):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(422, f"Enter a valid {label} date.") from exc
    if start and end and start > end:
        raise HTTPException(422, "Start date must not be after end date.")
    payment_methods = _payment_methods(db)
    if method != "All" and method not in payment_methods:
        raise HTTPException(422, "Choose a valid payment method.")

    payment_statement = select(Payment).outerjoin(
        Invoice, Payment.invoice_id == Invoice.id
    ).options(
        selectinload(Payment.invoice).selectinload(Invoice.guest),
        selectinload(Payment.invoice).selectinload(Invoice.room),
    )
    refund_statement = select(Refund, Invoice).outerjoin(
        Invoice, Refund.invoice_id == Invoice.id
    )
    if start:
        payment_statement = payment_statement.where(Payment.business_date >= start)
        refund_statement = refund_statement.where(Refund.business_date >= start)
    if end:
        payment_statement = payment_statement.where(Payment.business_date <= end)
        refund_statement = refund_statement.where(Refund.business_date <= end)
    if method != "All":
        payment_statement = payment_statement.where(Payment.method == method)
        refund_statement = refund_statement.where(Refund.method == method)
    if invoice.strip():
        term = f"%{invoice.strip()}%"
        match_invoice = or_(
            Payment.invoice_number.ilike(term), Invoice.invoice_number.ilike(term),
            Invoice.guest_name.ilike(term), Invoice.room_number.ilike(term),
        )
        payment_statement = payment_statement.where(match_invoice)
        refund_statement = refund_statement.where(or_(
            Refund.invoice_number.ilike(term), Invoice.invoice_number.ilike(term),
            Invoice.guest_name.ilike(term), Invoice.room_number.ilike(term),
        ))
    payments = [
        payment for payment in db.scalars(
            payment_statement.order_by(Payment.paid_at.desc(), Payment.id.desc()).limit(1000)
        ).all()
        if not _is_demo_invoice(payment.invoice)
    ][:300]
    refunds = [
        {"refund": row[0], "invoice": row[1]}
        for row in db.execute(
            refund_statement.order_by(Refund.refunded_at.desc(), Refund.id.desc()).limit(1000)
        ).all()
        if not _is_demo_invoice(row[1])
    ][:300]
    outstanding = [
        invoice for invoice in db.scalars(select(Invoice).options(
            selectinload(Invoice.guest), selectinload(Invoice.room)
        ).where(Invoice.balance > 0, Invoice.status == "issued")
            .order_by(Invoice.balance.desc()).limit(200)).all()
        if not _is_demo_invoice(invoice)
    ][:50]
    today = local_now().date().isoformat()
    month_start = (local_now().date() - timedelta(days=29)).isoformat()
    collected_today = sum(
        float(payment.amount) for payment in payments if payment.business_date == today
    )
    cash_today = sum(
        float(payment.amount) for payment in payments
        if payment.business_date == today and payment.method == "Cash"
    )
    refunds_today = sum(
        float(row["refund"].amount) for row in refunds
        if getattr(row["refund"], "business_date", None) == today
        and row["refund"].kind == "refund"
    )
    outstanding_total = sum(float(invoice.balance) for invoice in outstanding)
    entries_today = sum(1 for payment in payments if payment.business_date == today)
    method_rows: dict[str, list[float]] = {}
    for payment in payments:
        if month_start <= (payment.business_date or "") <= today:
            bucket = method_rows.setdefault(payment.method, [0, 0.0])
            bucket[0] += 1
            bucket[1] += float(payment.amount)
    method_totals = [
        (name, int(values[0]), values[1])
        for name, values in sorted(method_rows.items(), key=lambda item: -item[1][1])
    ]
    return render(request, "payments.html", db,
                  {
                      "payments": payments, "refunds": refunds, "outstanding": outstanding,
                      "start": start, "end": end, "method": method, "invoice_filter": invoice,
                      "payment_methods": payment_methods,
                      "collected_today": collected_today, "cash_today": cash_today,
                      "refunds_today": refunds_today, "outstanding_total": outstanding_total,
                      "entries_today": entries_today,
                      "method_totals": method_totals,
                  })


@app.get("/payments/actions/{action}", response_class=HTMLResponse)
def payment_action_modal(request: Request, action: str, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can issue refunds or adjustments.")
    if action not in {"refund", "adjustment"}:
        raise HTTPException(404, "Payment action not found.")
    return render(request, "partials/payment_ledger_modal.html", db, {
        "action": action,
        "methods": _payment_methods(db),
        "reasons": REFUND_REASONS if action == "refund" else ADJUSTMENT_REASONS,
    })


@app.post("/payments/actions/{action}")
async def payment_action(request: Request, action: str, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can issue refunds or adjustments.")
    if action not in {"refund", "adjustment"}:
        raise HTTPException(404, "Payment action not found.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    methods = _payment_methods(db)
    reasons = REFUND_REASONS if action == "refund" else ADJUSTMENT_REASONS
    invoice_number = str(form.get("invoice_number", "")).strip()
    amount_text = str(form.get("amount", "")).strip()
    method = str(form.get("method", "Cash"))
    reason = str(form.get("reason", "")).strip()
    selected_reason = reason
    reason_note = str(form.get("reason_note", "")).strip()
    try:
        try:
            raw_amount = Decimal(amount_text)
        except InvalidOperation as exc:
            raise ValueError("Enter a valid amount.") from exc
        if not raw_amount.is_finite() or abs(raw_amount) > Decimal("9999999999.99"):
            raise ValueError("Enter a valid amount within the supported range.")
        try:
            amount = raw_amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except InvalidOperation as exc:
            raise ValueError("Enter a valid amount with no more than two decimal places.") from exc
        if amount == 0:
            raise ValueError("Enter a non-zero amount.")
        if len(invoice_number) > 40:
            raise ValueError("Enter a valid invoice number.")
        if reason not in reasons:
            raise ValueError("Choose a reason for this ledger entry.")
        if reason == "Other (see note)":
            if not reason_note:
                raise ValueError("Add a short note for the selected reason.")
            reason = f"{reason}: {reason_note}"
        if len(reason) > 1000:
            raise ValueError("Reason must be 1,000 characters or fewer.")
        invoice_row = db.scalar(select(Invoice).where(
            Invoice.invoice_number == invoice_number,
            Invoice.status.in_(("issued", "paid")),
        ))
        if invoice_row is None:
            raise ValueError("The invoice was not found or is not eligible for ledger changes.")
        if action == "refund":
            if amount <= 0:
                raise ValueError("A refund amount must be greater than zero.")
            if method not in methods:
                raise ValueError("Choose a valid refund method.")
            paid = Decimal(str(invoice_row.paid_amount)).quantize(Decimal("0.01"))
            balance = Decimal(str(invoice_row.balance)).quantize(Decimal("0.01"))
            if amount > paid:
                raise ValueError("Refund cannot exceed the amount collected on this invoice.")
            invoice_row.paid_amount = paid - amount
            invoice_row.balance = balance + amount
            ledger_method = method
            audit_action = "payment.refund"
        else:
            if amount > Decimal(str(invoice_row.balance)).quantize(Decimal("0.01")):
                raise ValueError("A credit adjustment cannot exceed the outstanding balance.")
            new_balance = Decimal(str(invoice_row.balance)).quantize(Decimal("0.01")) - amount
            if new_balance < 0 or new_balance > Decimal("9999999999.99"):
                raise ValueError("The adjustment would create an invalid invoice balance.")
            invoice_row.balance = new_balance
            ledger_method = "Adjustment"
            audit_action = "payment.adjustment"
        invoice_row.status = "paid" if invoice_row.balance == 0 else "issued"
        now = local_now()
        db.add(Refund(
            invoice_id=invoice_row.id, invoice_number=invoice_row.invoice_number,
            kind=action, method=ledger_method, amount=amount,
            reason=reason, refunded_at=now.astimezone(timezone.utc).replace(tzinfo=None),
            business_date=now.date().isoformat(),
            created_at=now.astimezone(timezone.utc).replace(tzinfo=None),
            user_id=user.id,
        ))
        _audit(db, user, audit_action, "invoice", invoice_row.invoice_number,
               f"{amount:.2f} {action}: {reason}")
        db.commit()
        if request.headers.get("HX-Request") == "true":
            return Response(status_code=204, headers={
                "HX-Trigger": json.dumps({
                    "payment-saved": {},
                    "showToast": toast_event("Payment recorded successfully.", "success"),
                }),
                "HX-Redirect": toast_url("/payments", "Payment recorded successfully.", "success"),
            })
        return RedirectResponse(toast_url("/payments", "Payment recorded successfully.", "success"), status_code=303)
    except ValueError as exc:
        db.rollback()
        return render(request, "partials/payment_ledger_modal.html", db, {
            "action": action, "methods": methods, "reasons": reasons,
            "error": str(exc), "invoice_number_val": invoice_number,
            "amount_val": amount_text, "method_val": method,
            "reason_val": selected_reason, "reason_note_val": reason_note,
        }, status_code=200)
    except Exception:
        db.rollback()
        _log.exception("Ledger payment action failed for invoice %s", invoice_number)
        if request.headers.get("HX-Request") == "true":
            return Response(
                f'<div role="alert" class="rounded-xl border border-red-200 bg-red-50 p-3 text-sm font-semibold text-red-700">The transaction could not be completed. Please try again.</div>',
                status_code=200,
                headers={"HX-Trigger": json.dumps({"showToast": toast_event("The transaction could not be completed. Please try again.", "error")})},
            )
        return RedirectResponse(toast_url("/payments", "The transaction could not be completed. Please try again.", "error"), status_code=303)


def _report_dates(preset: str, start: str, end: str) -> tuple[str, str]:
    today = local_now().date()
    if preset == "today":
        first = last = today
    elif preset == "yesterday":
        first = last = today - timedelta(days=1)
    elif preset == "week":
        first, last = today - timedelta(days=today.weekday()), today
    elif preset == "month":
        first, last = today.replace(day=1), today
    elif preset == "last30":
        first, last = today - timedelta(days=29), today
    else:
        first = date.fromisoformat(start) if start else today
        last = date.fromisoformat(end) if end else today
    reporting.validate_period(first.isoformat(), last.isoformat())
    return first.isoformat(), last.isoformat()


def _report_options(db: Session) -> dict:
    return {
        "employees": db.scalars(select(Invoice.prepared_by).where(
            Invoice.prepared_by != ""
        ).distinct().order_by(Invoice.prepared_by)).all(),
        "methods": _payment_methods(db),
        "rooms": db.scalars(select(Room.number).where(
            Room.active.is_(True)
        ).order_by(Room.number)).all(),
    }


def _report_query_data(
    db: Session,
    report_key: str,
    start: str,
    end: str,
    employee: str,
    payment_method: str,
    transaction_type: str,
    room_no: str,
    invoice_no: str,
) -> tuple[dict, dict]:
    first, last = reporting.validate_period(start, end)
    filters = {
        "employee": employee,
        "payment_method": payment_method,
        "transaction_type": transaction_type,
        "room_no": room_no.strip(),
        "invoice_no": invoice_no.strip(),
    }
    report = reporting.build_report(db, report_key, first.isoformat(), last.isoformat(), filters)
    return report, filters


@app.get("/reports", response_class=HTMLResponse)
def reports_page(
    request: Request, report: str = "daily_sales", preset: str = "custom",
    start: str = "", end: str = "", employee: str = "All",
    payment_method: str = "All", transaction_type: str = "All",
    room_no: str = "", invoice_no: str = "", db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    try:
        start, end = _report_dates(preset, start, end)
        report_data, filters = _report_query_data(
            db, report, start, end, employee, payment_method,
            transaction_type, room_no, invoice_no,
        )
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    today = local_now().date().isoformat()
    yesterday = (local_now().date() - timedelta(days=1)).isoformat()
    daily_reports = daily_closing.list_reports(db)
    for entry in daily_reports:
        try:
            entry.totals = json.loads(entry.totals_json or "{}")
        except (TypeError, ValueError):
            entry.totals = {}
    return render(request, "reports.html", db, {
        "active_page": "reports",
        "report": report_data,
        "report_key": report,
        "start": start,
        "end": end,
        "preset": preset,
        "filters": filters,
        "cards": reporting.dashboard_cards(db, today),
        "options": _report_options(db),
        "report_options": reporting.REPORTS,
        "daily_reports": daily_reports,
        "scheduler_runs": daily_closing.list_runs(db),
        "daily_report_date": yesterday,
        "daily_preview": [
            {"label": "Yesterday", "date": yesterday, "sales": reporting.dashboard_cards(db, yesterday).get("gross_sales", 0)},
            {"label": "Today", "date": today, "sales": reporting.dashboard_cards(db, today).get("gross_sales", 0)},
        ],
        "daily_status": request.query_params.get("daily_status", ""),
    })


@app.post("/reports/excel", response_class=HTMLResponse)
async def analyze_excel_report(
    request: Request,
    file: UploadFile = File(...),
    csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    check_csrf(request, csrf_token)
    filename = (file.filename or "").strip()
    content = bytearray()
    while chunk := await file.read(1024 * 1024):
        content.extend(chunk)
        if len(content) > excel_analysis.MAX_FILE_BYTES:
            return render(request, "excel_analysis.html", db, {
                "active_page": "reports",
                "filename": filename,
                "error": "That file is larger than 10 MB. Upload a smaller workbook.",
                "server_toast": toast_event("That workbook is larger than 10 MB.", "error"),
            }, status_code=413)
    try:
        analysis = excel_analysis.analyze(bytes(content), filename)
    except ValueError as exc:
        return render(request, "excel_analysis.html", db, {
            "active_page": "reports",
            "filename": filename,
            "error": str(exc),
            "server_toast": toast_event(str(exc), "error"),
        }, status_code=422)
    return render(request, "excel_analysis.html", db, {
        "active_page": "reports",
        "analysis": analysis,
        "filename": filename,
        "server_toast": toast_event(
            f"Workbook analysed — {analysis['row_count']} rows.", "success"
        ),
    })


@app.post("/reports/daily/generate")
def generate_daily_close(
    request: Request,
    business_date: str = Form(""),
    csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can generate daily-close reports.")
    check_csrf(request, csrf_token)
    day = business_date.strip() or (local_now().date() - timedelta(days=1)).isoformat()
    try:
        report = daily_closing.generate_report(
            db,
            day,
            generated_by=user.full_name or user.username,
            user_id=user.id,
        )
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc)) from exc
    return RedirectResponse(
        toast_url(f"/reports?daily_status={report.status}", "Daily closing report generated.", "success"),
        status_code=303,
    )


@app.post("/reports/daily/{report_id}/regenerate")
def regenerate_daily_close(
    request: Request,
    report_id: int,
    csrf_token: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can regenerate daily-close reports.")
    check_csrf(request, csrf_token)
    selected = db.get(DailyReport, report_id)
    if selected is None:
        raise HTTPException(404, "Daily-close report not found.")
    previous = db.scalar(
        select(DailyReport)
        .where(DailyReport.business_date == selected.business_date)
        .order_by(DailyReport.version.desc())
        .limit(1)
    )
    if previous is None:
        raise HTTPException(404, "Daily-close report not found.")
    try:
        report = daily_closing.regenerate_report(db, previous, user=user)
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc)) from exc
    return RedirectResponse(
        toast_url(f"/reports?daily_status={report.status}", "Daily closing report regenerated.", "success"),
        status_code=303,
    )


@app.get("/reports/daily/{report_id}/download")
def download_daily_close(
    request: Request, report_id: int, db: Session = Depends(get_db)
):
    user_or_redirect(request, db)
    report = db.get(DailyReport, report_id)
    if report is None or report.status not in {"generated", "superseded"}:
        raise HTTPException(404, "Generated daily-close report not found.")
    try:
        path = daily_closing.report_file_path(report)
    except ValueError as exc:
        raise HTTPException(404, "Daily-close PDF is not available.") from exc
    if not path.is_file():
        raise HTTPException(404, "Daily-close PDF is not available on the server.")
    return FileResponse(
        path,
        media_type="application/pdf",
        filename=path.name,
    )


def _report_export_response(
    request: Request,
    report: str,
    start: str,
    end: str,
    employee: str,
    payment_method: str,
    transaction_type: str,
    room_no: str,
    invoice_no: str,
    db: Session,
) -> dict:
    user_or_redirect(request, db)
    try:
        report_data, _ = _report_query_data(
            db, report, start, end, employee, payment_method,
            transaction_type, room_no, invoice_no,
        )
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return report_data


@app.get("/reports/export.csv")
def reports_export_csv(
    request: Request, report: str = "daily_sales", start: str = "",
    end: str = "", employee: str = "All", payment_method: str = "All",
    transaction_type: str = "All", room_no: str = "", invoice_no: str = "",
    db: Session = Depends(get_db),
):
    report_data = _report_export_response(
        request, report, start, end, employee, payment_method,
        transaction_type, room_no, invoice_no, db,
    )
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([label for _key, label in report_data["columns"]])
    for row in report_data["rows"]:
        writer.writerow([
            "'" + str(row.get(key, "")) if isinstance(row.get(key), str)
            and row.get(key, "").startswith(("=", "+", "-", "@"))
            else row.get(key, "")
            for key, _label in report_data["columns"]
        ])
    filename = f"arati-{report}-{start}_{end}.csv"
    return Response(
        output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/reports/export.xlsx")
def reports_export_excel(
    request: Request, report: str = "daily_sales", start: str = "",
    end: str = "", employee: str = "All", payment_method: str = "All",
    transaction_type: str = "All", room_no: str = "", invoice_no: str = "",
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    report_data = _report_export_response(
        request, report, start, end, employee, payment_method,
        transaction_type, room_no, invoice_no, db,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report_data["title"][:31] or "Report"

    title_font = Font(bold=True, size=13)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B3A2A", end_color="1B3A2A", fill_type="solid")
    total_font = Font(bold=True)

    hotel_name = _billing_settings(db).get("hotel_name", "Arati Hotel")
    sheet.append([hotel_name])
    sheet.append([report_data["title"]])
    sheet.append([f"Business dates {start} to {end}"])
    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(
        [label for _key, label in report_data["columns"]]
    )
    for cell in sheet[header_row]:
        cell.font = header_font
        cell.fill = header_fill

    money_keys = {
        "gross_sales", "net_sales", "discount", "tax", "service_charge",
        "collected", "outstanding", "subtotal", "total", "paid", "balance",
        "amount", "revenue", "rate",
    }
    count_keys = {"invoices", "entries", "transactions", "nights"}

    for row in report_data["rows"]:
        values = []
        for key, _label in report_data["columns"]:
            value = row.get(key, "")
            if value is None:
                value = ""
            elif key in money_keys and isinstance(value, (int, float)):
                value = round(float(value), 2)
            elif key in count_keys and isinstance(value, (int, float)):
                value = int(value)
            values.append(value)
        sheet.append(values)

    last_data_row = sheet.max_row
    total_values = []
    for index, (key, _label) in enumerate(report_data["columns"], start=1):
        value = report_data["totals"].get(key)
        if index == 1 and value is None:
            value = "TOTAL"
        elif key in money_keys and isinstance(value, (int, float)):
            value = round(float(value), 2)
        total_values.append("" if value is None else value)
    if any(v not in ("", None) for v in total_values[1:]):
        sheet.append(total_values)
        for cell in sheet[sheet.max_row]:
            cell.font = total_font

    by_method = report_data["totals"].get("by_method")
    if isinstance(by_method, dict) and by_method:
        sheet.append([])
        method_header_row = sheet.max_row + 1
        sheet.append(["Payment method breakdown"])
        for cell in sheet[method_header_row]:
            cell.font = total_font
        for method, amount in sorted(by_method.items()):
            sheet.append([method, round(float(amount), 2)])

    formulas = report_data.get("formulas") or []
    if formulas:
        sheet.append([])
        formulas_header_row = sheet.max_row + 1
        sheet.append(["Formulas / notes"])
        for cell in sheet[formulas_header_row]:
            cell.font = total_font
        for formula in formulas:
            sheet.append([formula])

    for index, (_key, _label) in enumerate(report_data["columns"], start=1):
        width = 12
        column_cells = sheet[get_column_letter(index)]
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            width = max(width, len(text) + 2)
            if width > 40:
                width = 40
        column_letter = get_column_letter(index)
        sheet.column_dimensions[column_letter].width = width

    out = io.BytesIO()
    workbook.save(out)
    data = out.getvalue()
    filename = f"arati-{report}-{start}_{end}.xlsx"
    return Response(
        data,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/reports/export.pdf")
def reports_export_pdf(
    request: Request, report: str = "daily_sales", start: str = "",
    end: str = "", employee: str = "All", payment_method: str = "All",
    transaction_type: str = "All", room_no: str = "", invoice_no: str = "",
    db: Session = Depends(get_db),
):
    report_data = _report_export_response(
        request, report, start, end, employee, payment_method,
        transaction_type, room_no, invoice_no, db,
    )
    html = templates.get_template("report_print.html").render(
        request=request, report=report_data, hotel_settings=_billing_settings(db),
        money=money,
    )
    from weasyprint import HTML

    pdf = HTML(string=html, base_url=str(ROOT)).write_pdf()
    _save_pdf_copy(pdf, "daily-reports", f"arati-{report}-{start}_{end}.pdf")
    return Response(pdf, media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="arati-{report}-{start}_{end}.pdf"',
    })


def _short_bucket(label: str) -> str:
    """Shorten an ISO day/week/month bucket label for a chart axis."""
    try:
        if len(label) == 10 and label[4] == "-" and label[7] == "-":
            return date.fromisoformat(label).strftime("%d %b")
        if len(label) == 7 and label[4] == "-":
            return date.fromisoformat(f"{label}-01").strftime("%b")
    except ValueError:
        return label
    if "-W" in label:
        return "W" + label.split("-W", 1)[1]
    return label


def _short_points(points: list[dict]) -> list[dict]:
    return [{"label": _short_bucket(point["label"]), "value": point["value"]} for point in points]


def _trend_geometry(points: list[dict], width: int = 600, height: int = 180) -> dict:
    """Lay trend points out as an SVG area chart with an axis row below it."""
    if not points:
        return {"line": "", "area": "", "marks": [], "labels": [], "height": height, "count": 0}
    peak = max(point["value"] for point in points) or 1.0
    step = width / (len(points) - 1) if len(points) > 1 else 0.0
    marks = []
    for index, point in enumerate(points):
        x = width / 2 if len(points) == 1 else index * step
        marks.append({
            "x": round(x, 1),
            # Keep the tallest point clear of the top edge so its dot is not clipped.
            "y": round(height - (point["value"] / peak) * (height - 12) - 6, 1),
            "label": point["label"],
            "value": point["value"],
            "show": len(points) <= 12 or index == len(points) - 1,
        })
    line = "M" + " L".join(f"{mark['x']},{mark['y']}" for mark in marks)
    area = f"{line} L{marks[-1]['x']},{height} L{marks[0]['x']},{height} Z"
    every = max(1, -(-len(marks) // 8))
    labels = [
        {
            "x": mark["x"],
            "label": mark["label"],
            "anchor": (
                "start" if index == 0 and len(marks) > 1
                else "end" if index == len(marks) - 1 and len(marks) > 1
                else "middle"
            ),
        }
        for index, mark in enumerate(marks)
        if index % every == 0 or index == len(marks) - 1
    ]
    return {
        "line": line, "area": area, "marks": marks, "labels": labels,
        "height": height, "count": len(points),
    }


_MIX_COLORS = ("#1B3A2A", "#7D9B7A", "#D4A017", "#A3B89F", "#D4C08A", "#5A6B58", "#8A6A14")


def _mix_geometry(points: list[dict], limit: int = 6) -> dict:
    """Turn chart points into donut segments plus a legend."""
    rows = sorted((point for point in points if point["value"] > 0), key=lambda item: (-item["value"], item["label"]))
    if not rows:
        return {"segments": [], "legend": [], "total": 0.0}
    total = sum(point["value"] for point in rows)
    if len(rows) > limit:
        rows = rows[:limit - 1] + [{
            "label": "Other methods",
            "value": round(sum(point["value"] for point in rows[limit - 1:]), 2),
        }]
    segments, legend = [], []
    offset = 0.0
    for index, row in enumerate(rows):
        share = row["value"] / total * 100
        color = _MIX_COLORS[index % len(_MIX_COLORS)]
        segments.append({"label": row["label"], "color": color, "share": round(share, 2), "offset": round(-offset, 2)})
        legend.append({"label": row["label"], "color": color, "value": row["value"], "share": share})
        offset += share
    return {"segments": segments, "legend": legend, "total": total}


def _paired_bars(sales_points: list[dict], collection_points: list[dict], limit: int = 12) -> list[dict]:
    """Pair sales and collections per bucket for the grouped bar chart."""
    sales = {point["label"]: point["value"] for point in sales_points}
    collections = {point["label"]: point["value"] for point in collection_points}
    labels = sorted(set(sales) | set(collections))[-limit:]
    if not labels:
        return []
    peak = max(
        [sales.get(label, 0.0) for label in labels] + [collections.get(label, 0.0) for label in labels]
    ) or 1.0
    return [{
        "label": label,
        "sales": sales.get(label, 0.0),
        "collections": collections.get(label, 0.0),
        "sales_height": round(sales.get(label, 0.0) / peak * 100, 1),
        "collection_height": round(collections.get(label, 0.0) / peak * 100, 1),
    } for label in labels]


@app.get("/analysis", response_class=HTMLResponse)
def analysis_page(
    request: Request, start: str = "", end: str = "", bucket: str = "day",
    compare: str = "", compare_start: str = "", compare_end: str = "",
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can view Sales Analysis.")
    today = local_now().date()
    start = start or today.isoformat()
    end = end or today.isoformat()
    try:
        if compare == "1" and not (compare_start or compare_end):
            first, last = reporting.validate_period(start, end)
            period_days = (last - first).days + 1
            compare_end = (first - timedelta(days=1)).isoformat()
            compare_start = (first - timedelta(days=period_days)).isoformat()
        result = reporting.analysis_data(
            db, start, end, bucket, compare_start, compare_end
        )
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    charts = {chart["key"]: chart for chart in result["charts"]}
    sales_points = _short_points(charts["sales"]["points"])
    collection_points = _short_points(charts["collections"]["points"])
    return render(request, "analysis.html", db, {
        "active_page": "analysis",
        "analysis": result,
        "metrics": {metric["key"]: metric for metric in result["metrics"]},
        "changes": {
            row["key"]: row
            for row in (result.get("comparison") or {}).get("rows", [])
        },
        "charts": charts,
        "trend": _trend_geometry(sales_points),
        "mix": _mix_geometry(charts["methods"]["points"]),
        "bars": _paired_bars(sales_points, collection_points),
        "start": start,
        "end": end,
        "bucket": bucket,
        "compare": compare == "1",
        "compare_start": compare_start,
        "compare_end": compare_end,
    })


# --- day-wise calendar ------------------------------------------------------


def _calendar_bins(db: Session, range_start_iso: str, range_end_iso: str, month_start: date, month_end: date) -> dict:
    """Per-business-day totals for a calendar range (covers adjacent days)."""
    empty_value = {"invoices": 0, "sales": 0.0, "payments": 0.0, "payment_count": 0,
                   "checkins": 0, "checkouts": 0, "bookings": 0, "orders": 0, "guests": 0,
                   "rooms": 0}
    bins = {}
    cur = month_start
    while cur <= month_end:
        bins[cur.isoformat()] = dict(empty_value)
        cur += timedelta(days=1)

    for day_iso, count, total in db.execute(
        select(Invoice.bill_date, func.count(Invoice.id), func.coalesce(func.sum(Invoice.total), 0.0))
        .where(Invoice.bill_date >= range_start_iso, Invoice.bill_date <= range_end_iso,
               Invoice.status != "void")
        .group_by(Invoice.bill_date)
    ).all():
        if day_iso in bins:
            bins[day_iso]["invoices"] = int(count)
            bins[day_iso]["sales"] = float(total)
    for day_iso, count, total in db.execute(
        select(Payment.business_date, func.count(Payment.id), func.coalesce(func.sum(Payment.amount), 0.0))
        .where(Payment.business_date >= range_start_iso, Payment.business_date <= range_end_iso)
        .group_by(Payment.business_date)
    ).all():
        if day_iso in bins:
            bins[day_iso]["payments"] = float(total)
            bins[day_iso]["payment_count"] = int(count)
    for day_iso, count in db.execute(
        select(Reservation.business_date, func.count(Reservation.id))
        .where(Reservation.business_date >= range_start_iso, Reservation.business_date <= range_end_iso)
        .group_by(Reservation.business_date)
    ).all():
        if day_iso in bins:
            bins[day_iso]["bookings"] = int(count)
    for day_iso, count in db.execute(
        select(Stay.business_date, func.count(Stay.id))
        .where(Stay.business_date >= range_start_iso, Stay.business_date <= range_end_iso)
        .group_by(Stay.business_date)
    ).all():
        if day_iso in bins:
            bins[day_iso]["checkins"] = int(count)
    range_start_dt = business_window(month_start)[0]
    range_end_dt = business_window(month_end + timedelta(days=1))[0]
    for stay in db.scalars(
        select(Stay.check_out).where(Stay.check_out >= range_start_dt, Stay.check_out < range_end_dt)
    ).all():
        day_iso = _business_date(stay).isoformat()
        if day_iso in bins:
            bins[day_iso]["checkouts"] += 1
    for day_iso, count in db.execute(
        select(Order.business_date, func.count(Order.id))
        .where(Order.business_date >= range_start_iso, Order.business_date <= range_end_iso)
        .group_by(Order.business_date)
    ).all():
        if day_iso in bins:
            bins[day_iso]["orders"] = int(count)
    guest_start_dt = range_start_dt
    guest_end_dt = range_end_dt
    for created_at in db.scalars(
        select(Guest.created_at).where(Guest.created_at >= guest_start_dt, Guest.created_at < guest_end_dt)
    ).all():
        if created_at is None:
            continue
        day_iso = _business_date(created_at).isoformat()
        if day_iso in bins:
            bins[day_iso]["guests"] += 1

    room_sets: dict[str, set[int]] = {iso: set() for iso in bins}
    for stay in db.scalars(
        select(Stay).where(
            Stay.check_in < range_end_dt,
            or_(Stay.check_out.is_(None), Stay.check_out >= range_start_dt),
            Stay.status.in_(("checked_in", "checked_out")),
        )
    ).all():
        first = max(_business_date(stay.check_in), month_start)
        last = min(_business_date(stay.check_out) if stay.check_out else month_end, month_end)
        day = first
        while day <= last:
            key = day.isoformat()
            if key in room_sets:
                room_sets[key].add(stay.room_id)
            day += timedelta(days=1)
    for key, room_ids in room_sets.items():
        bins[key]["rooms"] = len(room_ids)
    return bins


def month_start_utc(day: date) -> datetime:
    return business_window(day)[0]


def month_end_utc(day: date) -> datetime:
    next_month = _first_of_next(day)
    return business_window(next_month)[0]


def _first_of_next(day: date) -> date:
    if day.month == 12:
        return date(day.year + 1, 1, 1)
    return date(day.year, day.month + 1, 1)


@app.get("/calendar", response_class=HTMLResponse)
def calendar_page(
    request: Request, month: str = Query(""), day: str = Query("", alias="date"),
    db: Session = Depends(get_db),
):
    """The hotel's day book: a month grid where every day carries its own
    sales, collections, arrivals, departures, bookings, orders and new
    customers. Clicking any day opens every record stored on that date."""
    user_or_redirect(request, db)
    today = local_now().date()
    if month:
        try:
            month_date = datetime.fromisoformat(month + "-01").date()
        except ValueError as exc:
            raise HTTPException(422, "Choose a valid month (YYYY-MM).") from exc
    else:
        month_date = today.replace(day=1)
    day_param = day
    selected = None
    if day_param:
        try:
            candidate = datetime.fromisoformat(day_param).date()
        except ValueError as exc:
            raise HTTPException(422, "Choose a valid day (YYYY-MM-DD).") from exc
        if candidate.year == month_date.year and candidate.month == month_date.month:
            selected = candidate
        else:
            selected = None
    if selected is None:
        selected = today if (today.year == month_date.year and today.month == month_date.month) else month_date

    import calendar as _cal
    _cal.setfirstweekday(_cal.MONDAY)
    last_day = _cal.monthrange(month_date.year, month_date.month)[1]
    month_first = month_date.replace(day=1)
    month_last = month_date.replace(day=last_day)
    grid_start = month_first - timedelta(days=(month_first.weekday()))
    grid_end = month_last + timedelta(days=(6 - month_last.weekday()))
    bins = _calendar_bins(db, grid_start.strftime("%Y-%m-%d"), grid_end.strftime("%Y-%m-%d"), grid_start, grid_end)

    day_iso = selected.isoformat()
    day_invoices = db.scalars(
        select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room))
        .where(Invoice.bill_date == day_iso, Invoice.status != "void")
        .order_by(Invoice.issue_date, Invoice.id)
    ).all()
    day_payments = db.scalars(
        select(Payment).where(Payment.business_date == day_iso).order_by(Payment.paid_at, Payment.id)
    ).all()
    day_refunds = db.scalars(
        select(Refund).where(Refund.business_date == day_iso).order_by(Refund.refunded_at)
    ).all()
    day_checkins = db.scalars(
        select(Stay).options(selectinload(Stay.guest), selectinload(Stay.room))
        .where(Stay.business_date == day_iso).order_by(Stay.check_in, Stay.id)
    ).all()
    checked_out_this_month = db.scalars(
        select(Stay).options(selectinload(Stay.guest), selectinload(Stay.room))
        .where(Stay.check_out >= month_start_utc(month_date), Stay.check_out < month_end_utc(month_date))
    ).all()
    day_checkouts = [s for s in checked_out_this_month if _business_date(s.check_out) == selected]
    day_reservations = db.scalars(
        select(Reservation).options(selectinload(Reservation.room), selectinload(Reservation.guest))
        .where(Reservation.business_date == day_iso)
        .order_by(Reservation.check_in, Reservation.id)
    ).all()
    day_orders = db.scalars(
        select(Order).options(selectinload(Order.items))
        .where(Order.business_date == day_iso).order_by(Order.created_at, Order.id)
    ).all()
    start_d, end_d = business_window(selected)
    day_guests = db.scalars(
        select(Guest).where(Guest.created_at >= start_d, Guest.created_at < end_d)
    ).all()
    day_charges = db.execute(
        select(func.count(RoomCharge.id), func.coalesce(func.sum(RoomCharge.amount), 0.0))
        .where(RoomCharge.charge_date == day_iso)
    ).one()

    prev_month = (month_date - timedelta(days=1)).replace(day=1)
    today_iso = local_now().date().isoformat()
    return render(request, "calendar.html", db, {
        "month": month_date, "grid": bins_to_grid(month_date, bins),
        "selected": selected, "selected_iso": day_iso,
        "prev_month": prev_month.strftime("%Y-%m"),
        "next_month": _first_of_next(month_date).strftime("%Y-%m"),
        "today_iso": today_iso,
        "day_invoices": day_invoices, "day_payments": day_payments,
        "day_refunds": day_refunds, "day_checkins": day_checkins,
        "day_checkouts": day_checkouts, "day_reservations": day_reservations,
        "day_orders": day_orders, "day_guests": day_guests,
        "day_charges_count": int(day_charges[0]), "day_charges_total": float(day_charges[1]),
        "day_sales": sum(float(inv.total) for inv in day_invoices),
        "day_collected": sum(float(p.amount) for p in day_payments),
    })


def bins_to_grid(month_date: date, bins: dict) -> list[list[dict]]:
    """Weeks of the month as cells carrying that day's aggregated figures (Mon start)."""
    import calendar as _calendar
    _calendar.setfirstweekday(_calendar.MONDAY)
    last_day = _calendar.monthrange(month_date.year, month_date.month)[1]
    month_first = month_date.replace(day=1)
    month_last = month_date.replace(day=last_day)
    grid_start = month_first - timedelta(days=month_first.weekday())
    grid_end = month_last + timedelta(days=6 - month_last.weekday())
    rows = []
    cur = grid_start
    empty = {"invoices": 0, "sales": 0.0, "payments": 0.0, "payment_count": 0,
             "checkins": 0, "checkouts": 0, "bookings": 0, "orders": 0, "guests": 0,
             "rooms": 0}
    while cur <= grid_end:
        week_row = []
        for _ in range(7):
            cell_date = cur
            week_row.append({"date": cell_date, "iso": cell_date.isoformat(),
                             "other_month": cell_date.month != month_date.month,
                             **bins.get(cell_date.isoformat(), dict(empty))})
            cur += timedelta(days=1)
        rows.append(week_row)
    return rows


MENU_NOTES = {
    ("Breakfast Set", "Simple Breakfast"): "Egg Any style, Breads Tost, Butter/Jam, Fruits & Tea",
    ("Breakfast Set", "Organic Breakfast"): "Buckwheat Pancake with Aalu Matar, Boiled Egg, Fruits, Tea",
    ("Breakfast Set", "Indian Breakfast"): "Puri Sabji / Aalu Paratha, Curd, Egg, Fruits and Tea",
}


def _menu_groups() -> list[dict]:
    groups: list[dict] = []
    index: dict[str, dict] = {}
    for group, name, rate, unit in MENU_ITEMS:
        entry = index.get(group)
        if entry is None:
            entry = {"no": len(groups) + 1, "name": group, "items": []}
            index[group] = entry
            groups.append(entry)
        entry["items"].append({
            "name": name,
            "rate": rate,
            "price": f"{rate:,}",
            "unit": unit,
            "note": MENU_NOTES.get((group, name), ""),
        })
    return groups


@app.get("/menu", response_class=HTMLResponse)
def menu_page(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    settings = _billing_settings(db)
    return render(request, "menu.html", db, {
        "menu_groups": _menu_groups(),
        "menu_item_count": len(MENU_ITEMS),
        "hotel_settings": settings,
    })


@app.get("/catalog", response_class=HTMLResponse)
def catalog_page(
    request: Request, q: str = "", category: str = "all",
    status: str = "active", edit_id: int | None = None, saved: int = 0,
    db: Session = Depends(get_db),
):
    user_or_redirect(request, db)
    if category not in {"all", "restaurant", "food", "service", "room", "other"}:
        raise HTTPException(422, "Choose a valid catalog category.")
    if status not in {"active", "inactive", "all"}:
        raise HTTPException(422, "Choose a valid catalog status.")
    statement = select(CatalogItem)
    if status in {"active", "inactive"}:
        statement = statement.where(CatalogItem.active.is_(status == "active"))
    if category in {"restaurant", "food", "service", "room", "other"}:
        statement = statement.where(CatalogItem.category == category)
    if q.strip():
        term = f"%{q.strip()}%"
        statement = statement.where(
            or_(CatalogItem.name.ilike(term), CatalogItem.item_code.ilike(term))
        )
    items = db.scalars(
        statement.order_by(CatalogItem.category, CatalogItem.sort_order, CatalogItem.name)
        .limit(500)
    ).all()
    editing = db.get(CatalogItem, edit_id) if edit_id is not None else None
    if edit_id is not None and editing is None:
        raise HTTPException(404, "Catalog item not found.")
    return render(request, "catalog.html", db, {
        "items": items,
        "editing": editing,
        "q": q,
        "category_filter": category,
        "status_filter": status,
        "saved": saved == 1,
        "form_values": {},
    })


@app.post("/catalog/items")
async def catalog_item_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can edit the catalog.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    item_id = str(form.get("item_id", "")).strip()
    name = str(form.get("name", "")).strip()
    category = str(form.get("category", "")).strip()
    unit = str(form.get("unit", "")).strip() or "pcs"
    item_code = str(form.get("item_code", "")).strip()
    try:
        try:
            rate = Decimal(str(form.get("rate", "")))
            sort_order = int(str(form.get("sort_order", "0") or "0"))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("Enter a valid price and sort order.") from exc
        if not name or len(name) > 160:
            raise ValueError("Item name is required and must be 160 characters or fewer.")
        if category not in {"restaurant", "food", "service", "room", "other"}:
            raise ValueError("Choose a valid item category.")
        if category == "room" and len(name) > 50:
            raise ValueError("Room type names must be 50 characters or fewer.")
        if not rate.is_finite() or rate < 0 or rate > Decimal("9999999999.99"):
            raise ValueError("Price must be a valid non-negative amount.")
        if not unit or len(unit) > 30 or len(item_code) > 80:
            raise ValueError("Unit must be 30 characters or fewer and item code 80 characters or fewer.")
        if sort_order < 0 or sort_order > 1_000_000:
            raise ValueError("Sort order must be between 0 and 1,000,000.")
        item = None
        if item_id:
            if not item_id.isdigit():
                raise ValueError("Select a valid catalog item.")
            item = db.get(CatalogItem, int(item_id))
            if item is None:
                raise HTTPException(404, "Catalog item not found.")
        else:
            item = CatalogItem(name=name, category=category)
            db.add(item)
        linked_rooms = []
        if item.id is not None and item.category == "room":
            linked_rooms = db.scalars(
                select(Room).where(Room.room_type_item_id == item.id)
            ).all()
            if linked_rooms and category != "room":
                raise ValueError("A room type assigned to rooms cannot be changed to another category.")
            if linked_rooms and len(name) > 50:
                raise ValueError("Room type names assigned to rooms must be 50 characters or fewer.")
        item.name = name
        item.category = category
        item.rate = rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        item.unit = unit
        item.item_code = item_code
        item.taxable = str(form.get("taxable", "")) in {"1", "on", "true"}
        item.sort_order = sort_order
        db.flush()
        for room in linked_rooms:
            room.type = item.name
            if not room.nightly_rate_override:
                room.nightly_rate = item.rate
        _audit(
            db, user, "catalog.save", "catalog_item", str(item.id),
            f"{item.name} · {item.category} · {item.rate:.2f}",
        )
        db.commit()
        return RedirectResponse(toast_url("/catalog?saved=1", "Catalog item saved.", "success"), status_code=303)
    except HTTPException:
        db.rollback()
        raise
    except (InvalidOperation, ValueError) as exc:
        db.rollback()
        items = db.scalars(
            select(CatalogItem).order_by(CatalogItem.category, CatalogItem.sort_order, CatalogItem.name)
            .limit(500)
        ).all()
        return render(request, "catalog.html", db, {
            "items": items, "editing": None, "q": "", "category_filter": "all",
            "status_filter": "all", "saved": False,
            "error": str(exc),
            "server_toast": toast_event(str(exc), "error"),
            "form_values": {
                "item_id": item_id, "name": name, "category": category,
                "rate": str(form.get("rate", "")), "unit": unit, "item_code": item_code,
                "taxable": str(form.get("taxable", "")) in {"1", "on", "true"},
                "sort_order": str(form.get("sort_order", "0")),
            },
        }, status_code=422)


@app.post("/catalog/items/{item_id}/active")
def catalog_item_set_active(
    request: Request, item_id: int, active: bool = Form(...),
    csrf_token: str = Form(""), db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can edit the catalog.")
    check_csrf(request, csrf_token)
    item = db.get(CatalogItem, item_id)
    if item is None:
        raise HTTPException(404, "Catalog item not found.")
    if item.category == "room" and not active and db.scalar(
        select(Room.id).where(Room.room_type_item_id == item.id, Room.active.is_(True))
    ) is not None:
        raise HTTPException(409, "A room type assigned to active rooms cannot be deactivated.")
    item.active = active
    _audit(db, user, "catalog.activate" if active else "catalog.deactivate",
           "catalog_item", str(item.id), item.name)
    db.commit()
    return RedirectResponse(toast_url("/catalog", "Catalog item reactivated." if active else "Catalog item deactivated.", "success" if active else "info"), status_code=303)


@app.get("/printer", response_class=HTMLResponse)
def printer_page(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change printer settings.")
    return render(request, "printer.html", db, {
        "settings": _billing_settings(db),
        "saved": request.query_params.get("saved") == "1",
    })


@app.post("/printer")
async def printer_settings_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change printer settings.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    keys = (
        "printer_paper_size", "printer_margins", "printer_show_logo",
        "printer_header", "printer_footer", "printer_contact_details",
        "printer_tax_details", "printer_copies", "printer_default_workflow",
        "printer_default_device",
    )
    values = {key: str(form.get(key, "")).strip() for key in keys}
    for key in (
        "printer_show_logo", "printer_header", "printer_footer",
        "printer_contact_details", "printer_tax_details",
    ):
        values[key] = "1" if str(form.get(key, "")).strip().lower() in {"1", "on", "true"} else "0"
    try:
        if values["printer_paper_size"] not in {"A4", "80", "58"}:
            raise ValueError("Choose A4, 80 mm, or 58 mm paper.")
        if values["printer_margins"] not in {"narrow", "normal", "wide"}:
            raise ValueError("Choose narrow, normal, or wide margins.")
        if values["printer_default_workflow"] not in {"preview", "direct"}:
            raise ValueError("Choose preview or direct print workflow.")
        if values["printer_copies"] not in {"1", "2", "3", "4", "5"}:
            raise ValueError("Copies must be between 1 and 5.")
        if len(values["printer_default_device"]) > 120:
            raise ValueError("Default printer name must be 120 characters or fewer.")
        for key in (
            "printer_show_logo", "printer_header", "printer_footer",
            "printer_contact_details", "printer_tax_details",
        ):
            if values[key] not in {"0", "1"}:
                raise ValueError("Printer display options must be enabled or disabled.")
        for key, value in values.items():
            db.merge(HotelSetting(key=key, value=value))
        _audit(db, user, "settings.printer.update", "printer", "configuration",
               "Printer layout and output settings updated")
        db.commit()
        return RedirectResponse(toast_url("/printer?saved=1", "Printer settings saved.", "success"), status_code=303)
    except ValueError as exc:
        db.rollback()
        values.update({key: str(form.get(key, "")) for key in keys})
        return render(request, "printer.html", db, {
            "settings": {**_billing_settings(db), **values},
            "error": str(exc),
            "saved": False,
        }, status_code=422)


@app.get("/printer/test", response_class=HTMLResponse)
def printer_test_page(request: Request, preview: bool = False, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can test printer settings.")
    settings = _billing_settings(db)
    guest = SimpleNamespace(
        name="Printer Test Guest", phone="9800000000", address="Khandbari",
        pan_vat_number="",
    )
    item = SimpleNamespace(
        description="Sample accommodation", quantity=Decimal("1"),
        rate=Decimal("2000.00"), discount_amount=Decimal("0.00"),
        vat_amount=Decimal("260.00"), amount=Decimal("2260.00"),
    )
    invoice = SimpleNamespace(
        id=0, invoice_number="PRINTER-TEST", issue_date=local_now(),
        invoice_type="Tax Invoice",
        language="Nepali" if settings.get("invoice_language") == "ne" else "English",
        bill_type="tax", guest_name=guest.name, guest_phone=guest.phone,
        guest_address=guest.address, guest_pan="", guest=guest, room_number="",
        room=None, payment_terms=settings.get("payment_terms", "Payment due on presentation."),
        items=[item], subtotal=Decimal("2000.00"), item_discount=Decimal("0.00"),
        discount_amount=Decimal("0.00"), net_subtotal=Decimal("2000.00"),
        service_charge_rate=Decimal("10.00"), service_charge=Decimal("200.00"),
        vat_rate=Decimal("13.00"), vat_amount=Decimal("260.00"),
        total=Decimal("2460.00"), paid_amount=Decimal("0.00"),
        balance=Decimal("2460.00"), copy_type=settings.get("invoice_copy", "Customer Copy"),
        payment_method="Cash", prepared_by="Printer setup", received_by="",
        notes="This is a sample only; it is not saved as an invoice.",
    )
    config = _printer_config(settings)
    return render(request, "print_invoice.html", db, {
        "invoice": invoice,
        "auto_print": not preview and config["workflow"] == "direct",
        "is_preview": preview,
        "hotel_settings": settings,
        "printer_config": config,
        "printer_logo_src": _printer_logo_source(settings),
        "is_test_print": True,
    })


@app.get("/staff", response_class=HTMLResponse)
def staff_page(
    request: Request, created: int = 0, saved: str = "", edit_id: int | None = None,
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    staff = db.scalars(select(User).order_by(User.role, User.full_name, User.username)).all()
    editing = db.get(User, edit_id) if edit_id is not None else None
    if edit_id is not None and editing is None:
        raise HTTPException(404, "Staff account not found.")
    return render(request, "staff.html", db, {
        "staff": staff,
        "created": created == 1,
        "saved": saved,
        "editing": editing,
        "form_values": {},
        "active_sessions_count": _active_sessions_by_user(db),
    })


def _active_sessions_by_user(db: Session) -> dict[str, int]:
    rows = db.execute(
        select(User.id, func.count(UserSession.id))
        .outerjoin(
            UserSession,
            (UserSession.user_id == User.id) & (UserSession.revoked_at.is_(None)),
        )
        .where(User.active.is_(True))
        .group_by(User.id)
    ).all()
    return {str(user_id): int(count) for user_id, count in rows}


def _staff_tracking_range(start: str, end: str) -> tuple[date, date, datetime, datetime]:
    today = local_now().date()
    try:
        start_date = date.fromisoformat(start.strip()) if start.strip() else today.replace(day=1)
        end_date = date.fromisoformat(end.strip()) if end.strip() else today
    except ValueError as exc:
        raise HTTPException(422, "Enter a valid tracking date range.") from exc
    if start_date > end_date:
        raise HTTPException(422, "The tracking start date must be on or before the end date.")
    start_dt = datetime.combine(start_date, time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    end_dt = datetime.combine(end_date + timedelta(days=1), time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    return start_date, end_date, start_dt, end_dt


def _staff_duration_label(seconds: int) -> str:
    minutes = max(0, int(seconds)) // 60
    hours, minutes = divmod(minutes, 60)
    return f"{hours}h {minutes}m" if hours else f"{minutes}m"


def _staff_tracking_data(db: Session, start_dt: datetime, end_dt: datetime) -> dict:
    users = db.scalars(select(User).order_by(User.role, User.full_name, User.username)).all()
    buckets: dict[int, dict] = {}
    for person in users:
        buckets[person.id] = {
            "sessions": [], "stay_rows": [], "invoice_rows": [], "refunds": [],
            "deleted_drafts": [], "audit_rows": [], "guest_ids": set(),
            "sales": Decimal("0"), "discount": Decimal("0"), "collected": Decimal("0"),
            "void_count": 0, "void_amount": Decimal("0"), "error_count": 0,
            "error_amount": Decimal("0"), "session_seconds": 0,
            "active_sessions": 0, "last_signin": None, "error_rows": [],
        }

    for record in db.scalars(
        select(UserSession).where(
            UserSession.created_at >= start_dt, UserSession.created_at < end_dt
        ).order_by(UserSession.created_at.desc())
    ).all():
        bucket = buckets.get(record.user_id)
        if bucket is None:
            continue
        ended_at = record.revoked_at or record.last_seen_at or record.created_at
        seconds = max(0, int((ended_at - record.created_at).total_seconds()))
        bucket["sessions"].append({
            "session": record, "active": record.revoked_at is None,
            "ended_at": ended_at, "duration_label": _staff_duration_label(seconds),
        })
        bucket["session_seconds"] += seconds
        if record.revoked_at is None:
            bucket["active_sessions"] += 1
        if bucket["last_signin"] is None or record.created_at > bucket["last_signin"]:
            bucket["last_signin"] = record.created_at

    for stay in db.scalars(
        select(Stay).options(selectinload(Stay.guest), selectinload(Stay.room)).where(
            Stay.check_in >= start_dt, Stay.check_in < end_dt
        ).order_by(Stay.check_in.desc())
    ).all():
        bucket = buckets.get(stay.user_id)
        if bucket is None or _is_demo_guest(stay.guest):
            continue
        bucket["stay_rows"].append(stay)

    for invoice in db.scalars(
        select(Invoice).options(selectinload(Invoice.guest), selectinload(Invoice.room)).where(
            Invoice.issue_date >= start_dt, Invoice.issue_date < end_dt
        ).order_by(Invoice.issue_date.desc())
    ).all():
        if _is_demo_invoice(invoice):
            continue
        bucket = buckets.get(invoice.user_id)
        if bucket is None or invoice.status == "draft":
            continue
        bucket["invoice_rows"].append(invoice)
        if invoice.status in {"issued", "paid"}:
            bucket["sales"] += Decimal(str(invoice.total or 0))
            bucket["discount"] += (
                Decimal(str(invoice.discount_amount or 0))
                + Decimal(str(invoice.item_discount or 0))
            )
            if invoice.guest_id is not None:
                bucket["guest_ids"].add(invoice.guest_id)
        elif invoice.status == "void":
            bucket["void_count"] += 1
            bucket["void_amount"] += Decimal(str(invoice.total or 0))

    for payment in db.scalars(
        select(Payment).where(
            Payment.paid_at >= start_dt, Payment.paid_at < end_dt
        )
    ).all():
        bucket = buckets.get(payment.user_id)
        if bucket is not None:
            bucket["collected"] += Decimal(str(payment.amount or 0))

    for refund in db.scalars(
        select(Refund).where(
            Refund.refunded_at >= start_dt, Refund.refunded_at < end_dt
        ).order_by(Refund.refunded_at.desc())
    ).all():
        bucket = buckets.get(refund.user_id)
        if bucket is not None:
            bucket["refunds"].append(refund)

    draft_deletions = []
    for audit in db.scalars(
        select(AuditLog).where(
            AuditLog.created_at >= start_dt, AuditLog.created_at < end_dt
        ).order_by(AuditLog.created_at.desc())
    ).all():
        bucket = buckets.get(audit.user_id)
        if bucket is None:
            continue
        bucket["audit_rows"].append(audit)
        if audit.action == "invoice.delete_draft":
            draft_deletions.append(audit)

    draft_creators: dict[str, int] = {}
    draft_numbers = {audit.entity_id for audit in draft_deletions if audit.entity_id}
    if draft_numbers:
        rows = db.execute(
            select(AuditLog.entity_id, AuditLog.user_id, AuditLog.created_at).where(
                AuditLog.action == "invoice.draft",
                AuditLog.entity_id.in_(draft_numbers),
            ).order_by(AuditLog.created_at)
        ).all()
        for entity_id, creator_id, _created_at in rows:
            if creator_id is not None and entity_id not in draft_creators:
                draft_creators[entity_id] = creator_id
    for audit in draft_deletions:
        bucket = buckets.get(draft_creators.get(audit.entity_id, audit.user_id))
        if bucket is not None:
            bucket["deleted_drafts"].append(audit)

    totals = {
        "staff": len(users), "signins": 0, "active_sessions": 0,
        "checkins": 0, "customers": 0, "sales": Decimal("0"),
        "discount": Decimal("0"), "collected": Decimal("0"), "errors": 0,
        "error_amount": Decimal("0"), "session_seconds": 0,
    }
    for bucket in buckets.values():
        bucket["customers"] = len(bucket["guest_ids"])
        bucket["error_count"] = (
            bucket["void_count"] + len(bucket["deleted_drafts"]) + len(bucket["refunds"])
        )
        bucket["error_amount"] = bucket["void_amount"] + sum(
            (Decimal(str(refund.amount or 0)) for refund in bucket["refunds"]),
            Decimal("0"),
        )
        bucket["session_time"] = _staff_duration_label(bucket["session_seconds"])
        bucket["audit_rows"] = bucket["audit_rows"][:500]
        error_rows = []
        for invoice in bucket["invoice_rows"]:
            if invoice.status != "void":
                continue
            error_rows.append({
                "when": invoice.issue_date, "type": "Voided invoice",
                "reference": invoice.invoice_number,
                "amount": Decimal(str(invoice.total or 0)),
                "detail": invoice.guest_name or "",
                "href": f"/billing/invoices/{invoice.id}/invoice",
            })
        for audit in bucket["deleted_drafts"]:
            error_rows.append({
                "when": audit.created_at, "type": "Deleted draft",
                "reference": audit.entity_id or "Draft", "amount": None,
                "detail": audit.detail, "href": "",
            })
        for refund in bucket["refunds"]:
            error_rows.append({
                "when": refund.refunded_at,
                "type": "Refund" if refund.kind == "refund" else "Adjustment",
                "reference": refund.invoice_number,
                "amount": Decimal(str(refund.amount or 0)),
                "detail": f"{refund.method} · {refund.reason}".strip(" ·"),
                "href": f"/billing/invoices/{refund.invoice_id}/invoice" if refund.invoice_id else "",
            })
        error_rows.sort(key=lambda item: item["when"], reverse=True)
        bucket["error_rows"] = error_rows

        totals["signins"] += len(bucket["sessions"])
        totals["active_sessions"] += bucket["active_sessions"]
        totals["checkins"] += len(bucket["stay_rows"])
        totals["customers"] += bucket["customers"]
        totals["sales"] += bucket["sales"]
        totals["discount"] += bucket["discount"]
        totals["collected"] += bucket["collected"]
        totals["errors"] += bucket["error_count"]
        totals["error_amount"] += bucket["error_amount"]
        totals["session_seconds"] += bucket["session_seconds"]
    totals["session_time"] = _staff_duration_label(totals["session_seconds"])
    return {"users": users, "buckets": buckets, "totals": totals}


@app.get("/staff/tracking", response_class=HTMLResponse)
def staff_tracking_page(
    request: Request, start: str = "", end: str = "", db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    start_date, end_date, start_dt, end_dt = _staff_tracking_range(start, end)
    data = _staff_tracking_data(db, start_dt, end_dt)
    return render(request, "staff_tracking.html", db, {
        **data, "start": start_date.isoformat(), "end": end_date.isoformat(),
    })


@app.get("/staff/tracking/{staff_id}", response_class=HTMLResponse)
def staff_tracking_detail(
    request: Request, staff_id: int, start: str = "", end: str = "",
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    target = db.get(User, staff_id)
    if target is None:
        raise HTTPException(404, "Staff account not found.")
    start_date, end_date, start_dt, end_dt = _staff_tracking_range(start, end)
    data = _staff_tracking_data(db, start_dt, end_dt)
    return render(request, "staff_tracking_detail.html", db, {
        **data, "target": target, "bucket": data["buckets"][target.id],
        "start": start_date.isoformat(), "end": end_date.isoformat(),
    })


@app.get("/staff/sessions", response_class=HTMLResponse)
def staff_sessions_page(
    request: Request, user_id: int | None = None, db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    if user_id is None:
        users = db.scalars(
            select(User).where(User.active.is_(True)).order_by(User.role, User.full_name, User.username)
        ).all()
    else:
        target = db.get(User, user_id)
        if target is None or not target.active:
            raise HTTPException(404, "Staff account not found.")
        users = [target]
    sessions_by_user: dict[int, list[dict]] = {}
    for u in users:
        rows = db.execute(
            select(UserSession).where(
                UserSession.user_id == u.id,
            ).order_by(UserSession.last_seen_at.desc())
        ).scalars().all()
        sessions_by_user[u.id] = [
            {
                "id": row.id,
                "session_id": row.session_id,
                "ip_address": row.ip_address,
                "user_agent": row.user_agent,
                "created_at": row.created_at,
                "last_seen_at": row.last_seen_at,
                "revoked_at": row.revoked_at,
                "revoked_by": row.revoked_by,
                "revoked_reason": row.revoked_reason,
                "active": row.revoked_at is None,
            }
            for row in rows
        ]
    return render(request, "partials/staff_sessions.html", db, {
        "users": users,
        "sessions_by_user": sessions_by_user,
        "current_user_id": user.id,
    })


@app.post("/staff/sessions/{user_id:int}/{session_id}/revoke")
def revoke_staff_session(
    request: Request,
    user_id: int,
    session_id: str,
    csrf_token: str = Form(""),
    reason: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = user_or_redirect(request, db)
    if actor.role != "admin":
        raise HTTPException(403, "Admin access required.")
    check_csrf(request, csrf_token)
    target = db.get(User, user_id)
    if target is None or not target.active:
        raise HTTPException(404, "Staff account not found.")
    session_record = db.scalar(
        select(UserSession).where(
            UserSession.id == session_id if session_id.isdigit()
            else UserSession.session_id == session_id,
            UserSession.user_id == user_id,
        )
    )
    if session_record is None:
        raise HTTPException(404, "Session not found.")
    if session_record.revoked_at is not None:
        raise HTTPException(409, "That session has already been revoked.")
    now = models.utcnow()
    session_record.revoked_at = now
    session_record.revoked_by_id = actor.id
    session_record.revoked_reason = reason.strip() or "Revoked by administrator"
    db.commit()
    _audit(
        db, actor, "session.revoke", "user", str(user_id),
        f"Revoked session for {target.username}: {session_record.session_id[:24]}... ({session_record.ip_address or 'unknown'}, {session_record.last_seen_at:%Y-%m-%d %H:%M}) — {session_record.revoked_reason}",
    )
    return RedirectResponse(
        toast_url(f"/staff/sessions?user_id={user_id}&revoked=1", "Session revoked.", "warning"),
        status_code=303,
    )


@app.post("/staff")
def create_staff_account(
    request: Request,
    username: str = Form(""),
    full_name: str = Form(""),
    password: str = Form(""),
    user_id: str = Form(""),
    role: str = Form("staff"),
    active: str | None = Form(None),
    active_present: str = Form(""),
    csrf_token: str = Form(...),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    check_csrf(request, csrf_token)

    username = username.strip().lower()
    full_name = full_name.strip()
    password = password.strip()
    target = None
    if user_id.strip():
        try:
            target = db.get(User, int(user_id))
        except ValueError as exc:
            raise HTTPException(422, "Choose a valid staff account.") from exc
        if target is None:
            raise HTTPException(404, "Staff account not found.")
    is_active = (
        active == "on"
        if active is not None
        else target is None and active_present != "1"
    )
    form_values = {
        "username": username, "full_name": full_name, "role": role,
        "active": is_active,
    }
    error = None
    if len(username) < 3 or len(username) > 80:
        error = "Username must be between 3 and 80 characters."
    elif not username.isascii() or not all(char.isalnum() or char in "._-" for char in username):
        error = "Username may contain letters, numbers, dots, underscores, and hyphens."
    elif len(full_name) > 160:
        error = "Full name must be 160 characters or fewer."
    elif role not in {"admin", "staff"}:
        error = "Choose a valid staff role."
    elif target is None and len(password) < 8:
        error = "Password must be at least 8 characters."
    elif password and len(password) < 8:
        error = "Password must be at least 8 characters."
    elif len(password) > 256:
        error = "Password must be 256 characters or fewer."
    elif target is not None and target.id == user.id and (
        role != "admin" or not is_active
    ):
        error = "You cannot remove administrator access or disable your own account."
    elif target is not None and target.role == "admin" and target.active and (
        role != "admin" or not is_active
    ) and (db.scalar(select(func.count(User.id)).where(
        User.role == "admin", User.active.is_(True)
    )) or 0) <= 1:
        error = "At least one active administrator must remain."

    if error:
        staff = db.scalars(select(User).order_by(User.role, User.full_name, User.username)).all()
        return render(request, "staff.html", db, {
            "staff": staff, "error": error, "form_values": form_values,
            "editing": target, "created": False, "saved": "",
        }, status_code=422)

    if target is None:
        db.add(User(
            username=username,
            full_name=full_name or username,
            password_hash=auth.hash_password(password),
            role=role,
            active=is_active,
            created_at=datetime.now(timezone.utc).replace(tzinfo=None),
        ))
        action = "staff.create"
        detail = f"Created {username} ({role})"
    else:
        credentials_changed = (
            password != ""
            or target.role != role
            or target.active != is_active
        )
        target.username = username
        target.full_name = full_name or username
        target.role = role
        target.active = is_active
        if password:
            target.password_hash = auth.hash_password(password)
            target.must_change_password = False
        if credentials_changed:
            target.session_version = (target.session_version or 0) + 1
        action = "staff.update"
        detail = f"Updated {username} ({role}, {'active' if target.active else 'disabled'})"
    try:
        _audit(db, user, action, "user", user_id.strip() or username, detail)
        db.commit()
        if target is not None and target is user:
            request.session["sv"] = user.session_version
    except IntegrityError:
        db.rollback()
        staff = db.scalars(select(User).order_by(User.role, User.full_name, User.username)).all()
        return render(request, "staff.html", db, {
            "staff": staff, "error": "That username is already in use.",
            "form_values": form_values, "editing": target,
            "created": False, "saved": "",
        }, status_code=409)
    return RedirectResponse(
        toast_url("/staff?created=1", "Staff account created.", "success")
        if target is None
        else toast_url("/staff?saved=updated", "Staff account updated.", "success"),
        status_code=303,
    )


@app.post("/staff/{staff_id}/delete")
def delete_staff_account(
    request: Request,
    staff_id: int,
    csrf_token: str = Form(...),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    check_csrf(request, csrf_token)
    target = db.get(User, staff_id)
    if target is None:
        raise HTTPException(404, "Staff account not found.")
    if target.id == user.id:
        raise HTTPException(422, "You cannot delete the signed-in account.")
    if target.role == "admin" and target.active and (db.scalar(select(func.count(User.id)).where(
        User.role == "admin", User.active.is_(True), User.id != target.id
    )) or 0) == 0:
        raise HTTPException(422, "At least one active administrator must remain.")

    references = (
        (Invoice, Invoice.user_id),
        (Payment, Payment.user_id),
        (Refund, Refund.user_id),
        (Reservation, Reservation.user_id),
        (Stay, Stay.user_id),
        (RoomCharge, RoomCharge.user_id),
        (AuditLog, AuditLog.user_id),
    )
    if any(db.scalar(select(func.count()).select_from(model).where(column == target.id)) for model, column in references):
        raise HTTPException(
            409,
            "This account is linked to transaction or audit history. Disable it instead of deleting it.",
        )
    username = target.username
    db.delete(target)
    _audit(db, user, "staff.delete", "user", str(staff_id), f"Deleted {username}")
    db.commit()
    return RedirectResponse(toast_url("/staff?saved=deleted", "Staff account removed.", "warning"), status_code=303)


# --- Attendance routes --------------------------------------------------------

def _attendance_range(start: str, end: str) -> tuple[date, date, datetime, datetime]:
    today = local_now().date()
    try:
        start_date = date.fromisoformat(start.strip()) if start.strip() else today.replace(day=1)
        end_date = date.fromisoformat(end.strip()) if end.strip() else today
    except ValueError as exc:
        raise HTTPException(422, "Enter a valid attendance date range.") from exc
    if start_date > end_date:
        raise HTTPException(422, "The attendance start date must be on or before the end date.")
    start_dt = datetime.combine(start_date, time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    end_dt = datetime.combine(end_date + timedelta(days=1), time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    return start_date, end_date, start_dt, end_dt


def _attendance_data(db: Session, start_dt: datetime, end_dt: datetime) -> dict:
    """Get attendance records and compute summaries for the given period."""
    # Get all active staff
    users = db.scalars(select(User).where(User.active.is_(True)).order_by(User.role, User.full_name, User.username)).all()
    users_by_id = {u.id: u for u in users}

    # Get attendance records in the period
    records = db.scalars(
        select(StaffAttendance).options(selectinload(StaffAttendance.user)).where(
            StaffAttendance.created_at >= start_dt, StaffAttendance.created_at < end_dt
        ).order_by(StaffAttendance.business_date.desc(), StaffAttendance.check_in.desc())
    ).all()

    # Build today's records (current business day)
    today_str = local_now().date().isoformat()
    today_records = [r for r in records if r.business_date == today_str]

    # Current time in UTC (naive) for ongoing session calculations
    now_utc = models.utcnow()

    # Add ongoing_hours to records with check_in but no check_out
    for r in records:
        if r.check_in and not r.check_out:
            r.ongoing_hours = (now_utc - r.check_in).total_seconds() / 3600

    # Compute totals
    total_records = len(records)
    present_today = len([r for r in today_records if r.check_in is not None])
    on_duty = len([r for r in today_records if r.check_in is not None and r.check_out is None])

    # Late count (check-in after 9:30 AM)
    late_count = 0
    total_hours = Decimal("0")
    for r in records:
        if r.check_in and r.check_out:
            hours = (r.check_out - r.check_in).total_seconds() / 3600
            total_hours += Decimal(str(hours))
            check_in_local = r.check_in.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ)
            if check_in_local.time() > time(9, 30):
                late_count += 1
        elif r.check_in and not r.check_out:
            # Ongoing session - both are naive UTC
            hours = (now_utc - r.check_in).total_seconds() / 3600
            total_hours += Decimal(str(hours))
            check_in_local = r.check_in.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ)
            if check_in_local.time() > time(9, 30):
                late_count += 1

    avg_hours = float(total_hours / max(1, len([r for r in records if r.check_in and r.check_out]))) if records else 0

    # Staff summaries
    staff_summaries = {}
    for user in users:
        user_records = [r for r in records if r.user_id == user.id]
        days_present = len([r for r in user_records if r.check_in is not None])
        completed_records = [r for r in user_records if r.check_in is not None and r.check_out is not None]
        total_user_hours = sum(
            Decimal(str((r.check_out - r.check_in).total_seconds() / 3600))
            for r in completed_records
        )
        ongoing_records = [r for r in user_records if r.check_in is not None and r.check_out is None]
        for r in ongoing_records:
            total_user_hours += Decimal(str((now_utc - r.check_in).total_seconds() / 3600))

        avg_user_hours = float(total_user_hours / max(1, len(completed_records))) if completed_records else 0

        late_user = sum(
            1 for r in user_records
            if r.check_in and r.check_in.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ).time() > time(9, 30)
        )
        early_departures = sum(
            1 for r in user_records
            if r.check_out and r.check_out.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ).time() < time(17, 0)
        )

        staff_summaries[user.id] = {
            "days_present": days_present,
            "total_hours": float(total_user_hours),
            "avg_hours": avg_user_hours,
            "late_count": late_user,
            "early_departures": early_departures,
        }

    return {
        "users": users,
        "users_by_id": users_by_id,
        "records": records,
        "today_records": today_records,
        "staff_summaries": staff_summaries,
        "totals": {
            "present_today": present_today,
            "on_duty": on_duty,
            "total_records": total_records,
            "avg_hours": round(avg_hours, 2),
            "late_count": late_count,
            "total_hours": float(total_hours),
        },
    }


@app.get("/my-attendance", response_class=HTMLResponse)
def my_attendance_page(
    request: Request, db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    
    # Get current month range
    today = local_now().date()
    start_date = today.replace(day=1)
    if start_date.month == 12:
        end_date = start_date.replace(year=start_date.year + 1, month=1, day=1) - timedelta(days=1)
    else:
        end_date = start_date.replace(month=start_date.month + 1, day=1) - timedelta(days=1)
    
    start_dt = datetime.combine(start_date, time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    end_dt = datetime.combine(end_date + timedelta(days=1), time.min, BUSINESS_TZ).astimezone(timezone.utc).replace(tzinfo=None)
    
    # Get attendance records for this user this month
    records = db.scalars(
        select(StaffAttendance).where(
            StaffAttendance.user_id == user.id,
            StaffAttendance.created_at >= start_dt, StaffAttendance.created_at < end_dt
        ).order_by(StaffAttendance.business_date.desc())
    ).all()
    
    # Add ongoing_hours
    now_utc = models.utcnow()
    for r in records:
        if r.check_in and not r.check_out:
            r.ongoing_hours = (now_utc - r.check_in).total_seconds() / 3600
    
    # Compute month summary
    today_str = today.isoformat()
    today_record = next((r for r in records if r.business_date == today_str), None)
    is_checked_in = today_record and today_record.check_in is not None
    is_checked_out = today_record and today_record.check_out is not None
    
    days_present = len([r for r in records if r.check_in is not None])
    completed = [r for r in records if r.check_in and r.check_out]
    ongoing = [r for r in records if r.check_in and not r.check_out]
    
    total_hours = sum(Decimal(str((r.check_out - r.check_in).total_seconds() / 3600)) for r in completed)
    for r in ongoing:
        total_hours += Decimal(str((now_utc - r.check_in).total_seconds() / 3600))
    
    avg_hours = float(total_hours / max(1, len(completed))) if completed else 0
    late_count = sum(
        1 for r in records
        if r.check_in and r.check_in.replace(tzinfo=timezone.utc).astimezone(BUSINESS_TZ).time() > time(9, 30)
    )
    
    month_summary = {
        "days_present": days_present,
        "total_hours": float(total_hours),
        "avg_hours": avg_hours,
        "late_count": late_count,
    }
    
    return render(request, "my_attendance.html", db, {
        "attendance_records": records,
        "today_record": today_record,
        "is_checked_in": is_checked_in,
        "is_checked_out": is_checked_out,
        "month_summary": month_summary,
    })


@app.get("/attendance", response_class=HTMLResponse)
def attendance_page(
    request: Request, start: str = "", end: str = "", db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Admin access required.")
    start_date, end_date, start_dt, end_dt = _attendance_range(start, end)
    data = _attendance_data(db, start_dt, end_dt)
    return render(request, "attendance.html", db, {
        **data, "start": start_date.isoformat(), "end": end_date.isoformat(),
    })


@app.post("/attendance/check-in")
def attendance_check_in(
    request: Request,
    csrf_token: str = Form(""),
    notes: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)

    today_str = local_now().date().isoformat()
    now = models.utcnow()

    # Check if already checked in today
    existing = db.scalar(
        select(StaffAttendance).where(
            StaffAttendance.user_id == user.id,
            StaffAttendance.business_date == today_str,
        )
    )
    if existing and existing.check_in is not None:
        return RedirectResponse(
            toast_url("/attendance", "You are already checked in today.", "warning"),
            status_code=303,
        )

    if existing:
        existing.check_in = now
        existing.check_in_ip = request.client.host if request.client else ""
        existing.check_in_user_agent = request.headers.get("user-agent", "")[:512]
        existing.notes = notes.strip() or existing.notes
        existing.updated_at = now
    else:
        db.add(StaffAttendance(
            user_id=user.id,
            business_date=today_str,
            check_in=now,
            check_in_ip=request.client.host if request.client else "",
            check_in_user_agent=request.headers.get("user-agent", "")[:512],
            notes=notes.strip(),
        ))

    _audit(db, user, "attendance.check_in", "attendance", today_str, f"Checked in at {now:%H:%M}")
    db.commit()
    return RedirectResponse(toast_url("/attendance", "Checked in successfully.", "success"), status_code=303)


@app.post("/attendance/check-out")
def attendance_check_out(
    request: Request,
    csrf_token: str = Form(""),
    notes: str = Form(""),
    db: Session = Depends(get_db),
):
    user = user_or_redirect(request, db)
    check_csrf(request, csrf_token)

    today_str = local_now().date().isoformat()
    now = models.utcnow()

    record = db.scalar(
        select(StaffAttendance).where(
            StaffAttendance.user_id == user.id,
            StaffAttendance.business_date == today_str,
        )
    )
    if not record or record.check_in is None:
        return RedirectResponse(
            toast_url("/attendance", "You must check in first.", "error"),
            status_code=303,
        )
    if record.check_out is not None:
        return RedirectResponse(
            toast_url("/attendance", "You are already checked out today.", "warning"),
            status_code=303,
        )

    record.check_out = now
    record.check_out_ip = request.client.host if request.client else ""
    record.check_out_user_agent = request.headers.get("user-agent", "")[:512]
    record.notes = (record.notes + " | " if record.notes else "") + f"Check-out: {notes.strip()}" if notes.strip() else record.notes
    record.updated_at = now

    _audit(db, user, "attendance.check_out", "attendance", today_str, f"Checked out at {now:%H:%M}")
    db.commit()
    return RedirectResponse(toast_url("/attendance", "Checked out successfully.", "success"), status_code=303)


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change hotel settings.")
    settings = _billing_settings(db)
    settings["payment_methods"] = ", ".join(_payment_methods(db))
    return render(request, "settings.html", db, {
        "settings": settings,
        "saved": request.query_params.get("saved") == "1",
    })


@app.post("/settings")
async def settings_save(request: Request, db: Session = Depends(get_db)):
    user = user_or_redirect(request, db)
    if user.role != "admin":
        raise HTTPException(403, "Only administrators can change hotel settings.")
    form = await request.form()
    check_csrf(request, str(form.get("csrf_token", "")))
    keys = (
        "hotel_name", "hotel_name_ne", "hotel_address1", "hotel_address2",
        "hotel_phone", "hotel_email", "hotel_website", "hotel_pan", "hotel_vat",
        "bill_prefix", "vat_rate", "service_charge_rate", "invoice_copy",
        "payment_terms", "payment_methods", "return_policy", "logo_path", "date_format",
        "invoice_language", "footer_note", "close_time",
    )
    values = {key: str(form.get(key, "")).strip() for key in keys}
    if "payment_methods" not in form:
        values["payment_methods"] = ", ".join(_payment_methods(db))
    if not values["close_time"]:
        values["close_time"] = _billing_settings(db).get("close_time", "00:05")
    try:
        if not values["hotel_name"] or len(values["hotel_name"]) > 160:
            raise ValueError("Hotel name is required and must be 160 characters or fewer.")
        limits = {
            "hotel_name_ne": 160, "hotel_address1": 240, "hotel_address2": 240,
            "hotel_phone": 40, "hotel_email": 180, "hotel_website": 180,
            "hotel_pan": 80, "hotel_vat": 80, "invoice_copy": 80,
            "payment_terms": 120, "return_policy": 1000, "logo_path": 500,
            "footer_note": 1000,
        }
        if any(len(values[key]) > limit for key, limit in limits.items()):
            raise ValueError("One or more settings exceed their allowed length.")
        values["payment_methods"] = _normalize_payment_methods(values["payment_methods"])
        if values["hotel_email"] and (
            "@" not in values["hotel_email"] or values["hotel_email"].startswith("@")
            or values["hotel_email"].endswith("@")
        ):
            raise ValueError("Enter a valid hotel email address.")
        logo_path = values["logo_path"]
        if logo_path and (
            not logo_path.startswith("/static/")
            or ".." in logo_path.split("/")
        ):
            raise ValueError("Logo path must be under /static/.")
        prefix = values["bill_prefix"].upper()
        if not prefix or len(prefix) > 12 or not prefix.isalnum():
            raise ValueError("Invoice prefix must contain only letters and digits (maximum 12).")
        values["bill_prefix"] = prefix
        try:
            vat = Decimal(values["vat_rate"])
            service = Decimal(values["service_charge_rate"])
        except InvalidOperation as exc:
            raise ValueError("VAT and service charge rates must be valid numbers.") from exc
        if (
            not vat.is_finite() or not service.is_finite()
            or not Decimal("0") <= vat <= Decimal("100")
            or not Decimal("0") <= service <= Decimal("100")
        ):
            raise ValueError("VAT and service charge must each be between 0 and 100.")
        values["vat_rate"] = format(vat.normalize(), "f")
        values["service_charge_rate"] = format(service.normalize(), "f")
        if values["date_format"] not in {"DD/MM/YYYY", "YYYY-MM-DD"}:
            raise ValueError("Date format must be DD/MM/YYYY or YYYY-MM-DD.")
        if values["invoice_language"] not in {"en", "ne"}:
            raise ValueError("Choose English or Nepali as the default invoice language.")
        try:
            close_time = time.fromisoformat(values["close_time"])
        except ValueError as exc:
            raise ValueError("Daily close time must be a valid 24-hour time.") from exc
        if close_time.second or close_time.microsecond:
            raise ValueError("Daily close time must use hours and minutes only.")
        values["close_time"] = close_time.strftime("%H:%M")
        for key, value in values.items():
            db.merge(HotelSetting(key=key, value=value))
        _audit(db, user, "settings.hotel.update", "hotel", "profile",
               "Hotel billing, invoice, and daily-close settings updated")
        db.commit()
        return RedirectResponse(toast_url("/settings?saved=1", "Hotel settings saved.", "success"), status_code=303)
    except ValueError as exc:
        db.rollback()
        values.update({
            key: str(form.get(key, values[key]))
            for key in keys
        })
        return render(request, "settings.html", db, {
            "settings": {**_billing_settings(db), **values},
            "error": str(exc),
            "saved": False,
        }, status_code=422)


# Website admin proxy
_WEBSITE_ADMIN_TOKEN = None
_WEBSITE_ADMIN_TOKEN_LOCK = threading.Lock()


async def _get_website_admin_token():
    global _WEBSITE_ADMIN_TOKEN
    with _WEBSITE_ADMIN_TOKEN_LOCK:
        if _WEBSITE_ADMIN_TOKEN:
            return _WEBSITE_ADMIN_TOKEN
        email = os.getenv("ARATI_WEBSITE_ADMIN_EMAIL", "aratiuser1@gmail.com").strip()
        password = os.getenv("ARATI_WEBSITE_ADMIN_PASSWORD", "SecurePass$#20255#4@%9")
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                resp = await client.post("https://aratihotel.com.np/api/login", json={"email": email, "password": password})
                if resp.status_code == 200:
                    data = resp.json()
                    _WEBSITE_ADMIN_TOKEN = data.get("token")
                    return _WEBSITE_ADMIN_TOKEN
        except Exception as e:
            _log.exception("Failed to login to website admin")
        return None


@app.get("/website-admin", response_class=HTMLResponse)
def website_admin(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    user = db.get(User, request.session.get("user_id"))
    if not user or user.role != "admin":
        raise HTTPException(status_code=403, detail="Admin access required.")
    return RedirectResponse(url="/website-admin/enquiries", status_code=303)


@app.get("/website-admin/enquiries", response_class=HTMLResponse)
def website_admin_enquiries(request: Request, db: Session = Depends(get_db)):
    user_or_redirect(request, db)
    user = db.get(User, request.session.get("user_id"))
    if not user or user.role != "admin":
        raise HTTPException(status_code=403, detail="Admin access required.")
    return render(request, "website_admin.html", db)


@app.api_route("/website-admin-proxy/{path:path}", methods=["GET", "POST", "PUT", "DELETE", "PATCH", "OPTIONS"])
@app.api_route("/website-admin-proxy", methods=["GET", "POST", "PUT", "DELETE", "PATCH", "OPTIONS"])
async def website_admin_proxy(request: Request, path: str = "", db: Session = Depends(get_db)):
    if request.method == "OPTIONS":
        return Response(status_code=200)
    user_or_redirect(request, db)
    user = db.get(User, request.session.get("user_id"))
    if not user or user.role != "admin":
        raise HTTPException(status_code=403, detail="Admin access required.")
    
    token = await _get_website_admin_token()
    if not token:
        raise HTTPException(status_code=502, detail="Failed to authenticate with website admin")
    
    # Build target URL
    if not path:
        target_url = "https://aratihotel.com.np/admin-dashboard"
    elif path.startswith("http"):
        target_url = path
    else:
        target_url = f"https://aratihotel.com.np/{path}"
    if request.url.query:
        target_url += "?" + request.url.query
    
    headers = {}
    for k, v in request.headers.items():
        if k.lower() in ('host', 'content-length', 'connection', 'accept-encoding'):
            continue
        headers[k] = v
    headers['Authorization'] = f'Bearer {token}'
    headers['Host'] = 'aratihotel.com.np'
    
    body = await request.body()
    
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=False) as client:
        try:
            resp = await client.request(
                method=request.method,
                url=target_url,
                headers=headers,
                content=body,
            )
            response_headers = {}
            for k, v in resp.headers.items():
                kl = k.lower()
                if kl in ('content-encoding', 'content-length', 'transfer-encoding', 'connection'):
                    continue
                if kl in ('x-frame-options'):
                    continue
                if kl == 'location':
                    # Rewrite redirects to go through proxy if they go to aratihotel.com.np
                    loc = v
                    if loc.startswith('https://aratihotel.com.np/'):
                        loc = '/website-admin-proxy/' + loc.replace('https://aratihotel.com.np/', '')
                    elif loc.startswith('/'):
                        loc = '/website-admin-proxy' + loc
                    response_headers[k] = loc
                    continue
                response_headers[k] = v
            return Response(content=resp.content, status_code=resp.status_code, headers=response_headers, media_type=resp.headers.get('content-type'))
        except Exception as e:
            _log.exception("Proxy error")
            raise HTTPException(status_code=502, detail="Proxy error")


@app.get("/healthz")
def healthz():
    return {"status": "ok"}
